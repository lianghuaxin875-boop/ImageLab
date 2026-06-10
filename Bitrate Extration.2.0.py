#!/usr/bin/env python3
"""
视频码率与综合质量检测工具（交互式选择拍摄模式）
功能：码率分析、清晰度、曝光、稳定性、掉帧、黑屏/白屏、画面冻结、音频基础
依赖：ffprobe, openpyxl, numpy, opencv-python, ffmpeg
"""

import subprocess
import json
import sys
import os
import re
import cv2
import numpy as np
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl 库。请执行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# ==================== 配置区域 ====================
FFPROBE_PATH = r"D:\码率工具FFepeg\ffmpeg-8.1.1-full_build\bin\ffprobe.exe"
INPUT_FOLDER = r"D:\Codex\码率测试\input"                     # 视频文件夹
OUTPUT_EXCEL_NAME = "video_report"           # 报告基础名（自动加时间戳）
INCLUDE_FRAMES = True                        # 是否包含逐帧大小数据
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".mpg", ".mpeg"}

# 抽帧检测间隔（秒）
SHARPNESS_SAMPLE_INTERVAL = 2.0
STABILITY_SAMPLE_INTERVAL = 0.5
FREEZE_THRESHOLD_FRAMES = 5
FREEZE_MOTION_THRESH = 0.5
OVEREXPOSURE_THRESH = 230
UNDEREXPOSURE_THRESH = 30
BLACK_THRESH = 5
WHITE_THRESH = 250
STABILITY_SHAKE_THRESH = 3.0
# =================================================

# ---------- 四种模式的码率标准（完整版）----------
STANDARDS_CONFIG = {
    1: {
        "name": "普通手机拍摄模式",
        "standards": [
            (921600, 30, 4.0, None),
            (2073600, 30, 8.0, 12.0),
            (8294400, 30, 35.0, 50.0),
            (float("inf"), 30, 100.0, 150.0),
        ],
        "table": [
            ["720P 30fps", "1280×720", "30", "4 - 6 Mbps", "低于 4 Mbps 偏低"],
            ["1080P 30fps", "1920×1080", "30", "8 - 12 Mbps", "低于 8 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 30fps", "3840×2160", "30", "35 - 50 Mbps", "低于 35 Mbps 偏低"],
            ["4K 60fps", "3840×2160", "60", "50 - 80 Mbps", "低于 50 Mbps 偏低"],
        ]
    },
    2: {
        "name": "短视频平台拍摄/上传模式",
        "standards": [
            (921600, 30, 3.0, None),
            (2073600, 30, 6.0, 10.0),
            (3686400, 30, 12.0, 16.0),
            (8294400, 30, 30.0, 45.0),
            (float("inf"), 30, 30.0, 45.0),
        ],
        "table": [
            ["720P 竖屏", "720×1280", "30fps", "3 - 6 Mbps", "低于 3 Mbps 偏低"],
            ["1080P 竖屏", "1080×1920", "30fps", "6 - 10 Mbps", "低于 6 Mbps 偏低"],
            ["1080P 竖屏高帧率", "1080×1920", "60fps", "10 - 16 Mbps", "低于 10 Mbps 偏低"],
            ["2K 竖屏", "1440×2560", "30fps", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 竖屏", "2160×3840", "30fps", "30 - 45 Mbps", "低于 30 Mbps 偏低"],
        ]
    },
    3: {
        "name": "监控/低码率存储模式",
        "standards": [
            (921600, 30, 1.0, None),
            (2073600, 30, 2.0, None),
            (3686400, 30, 4.0, None),
            (8294400, 30, 8.0, None),
            (float("inf"), 30, 8.0, None),
        ],
        "table": [
            ["720P 监控", "1280×720", "15/25fps", "1 - 3 Mbps", "低于 1 Mbps 偏低"],
            ["1080P 监控", "1920×1080", "15/25fps", "2 - 6 Mbps", "低于 2 Mbps 偏低"],
            ["2K 监控", "2560×1440", "15/25fps", "4 - 10 Mbps", "低于 4 Mbps 偏低"],
            ["4K 监控", "3840×2160", "15/25fps", "8 - 20 Mbps", "低于 8 Mbps 偏低"],
        ]
    },
    4: {
        "name": "高质量拍摄/相机素材模式",
        "standards": [
            (2073600, 30, 20.0, 30.0),
            (8294400, 30, 60.0, 100.0),
            (18000000, 30, 200.0, 200.0),
            (float("inf"), 30, 200.0, 200.0),
        ],
        "table": [
            ["1080P 30fps", "1920×1080", "30", "20 - 50 Mbps", "低于 20 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "30 - 80 Mbps", "低于 30 Mbps 偏低"],
            ["4K 24/25/30fps", "3840×2160", "24/25/30", "60 - 100 Mbps", "低于 60 Mbps 偏低"],
            ["4K 50/60fps", "3840×2160", "50/60", "100 - 200 Mbps", "低于 100 Mbps 偏低"],
            ["4K 10bit/LOG", "3840×2160", "24/25/30", "100 - 400 Mbps", "视编码而定"],
            ["6K/8K 高质量", "6K/8K", "24/30/60", "200 Mbps 以上", "低于 200 Mbps 偏低"],
        ]
    }
}

# ---------- 辅助函数 ----------
def select_mode():
    print("请选择拍摄模式：")
    print("1 - 普通手机拍摄模式")
    print("2 - 短视频平台模式")
    print("3 - 监控/低码率存储模式")
    print("4 - 高质量拍摄/相机素材模式")
    while True:
        try:
            choice = input("请输入数字 (1-4): ").strip()
            mode = int(choice)
            if mode in STANDARDS_CONFIG:
                return mode
            else:
                print("输入无效，请输入 1、2、3 或 4。")
        except ValueError:
            print("输入无效，请输入数字。")

def format_time(seconds):
    return str(timedelta(seconds=seconds))[:10]

def get_bitrate_threshold_mbps(standards, pixel_count, fps):
    if pixel_count is None or fps is None:
        return None
    for max_pixels, fps_limit, normal_min, high_min in standards:
        if pixel_count <= max_pixels:
            if high_min is not None and fps > fps_limit:
                return high_min
            else:
                return normal_min
    return None

def judge_bitrate(video_bitrate_bps, threshold_mbps):
    if video_bitrate_bps is None or threshold_mbps is None:
        return "N/A"
    return "OK" if video_bitrate_bps >= threshold_mbps * 1_000_000 else "NG"

# ---------- ffprobe 提取 ----------
def run_ffprobe(args):
    cmd = [FFPROBE_PATH, "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams"] + args
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return json.loads(result.stdout)
    except subprocess.CalledProcessError as e:
        print(f"ffprobe 执行失败: {e.stderr}", file=sys.stderr)
        return {}
    except FileNotFoundError:
        print(f"错误: 未找到 ffprobe，请检查路径: {FFPROBE_PATH}", file=sys.stderr)
        sys.exit(1)

def extract_video_info(file_path):
    info = run_ffprobe([file_path])
    if not info:
        return {}
    fmt = info.get("format", {})
    total_bitrate = int(fmt["bit_rate"]) if fmt.get("bit_rate") else None
    video_bitrate = None
    audio_bitrate = None
    width = height = None
    fps = None
    video_total_frames = None
    audio_info = {}
    for stream in info.get("streams", []):
        codec_type = stream.get("codec_type")
        if codec_type == "video":
            if width is None:
                width = stream.get("width")
                height = stream.get("height")
            if video_bitrate is None:
                if "bit_rate" in stream:
                    video_bitrate = int(stream["bit_rate"])
                else:
                    duration = float(stream.get("duration", 0))
                    size = stream.get("size")
                    if duration > 0 and size:
                        video_bitrate = int(float(size) * 8 / duration)
            if fps is None:
                r_frame_rate = stream.get("r_frame_rate")
                if r_frame_rate:
                    try:
                        if '/' in r_frame_rate:
                            num, den = r_frame_rate.split('/')
                            fps = float(num) / float(den)
                        else:
                            fps = float(r_frame_rate)
                    except:
                        fps = None
            if video_total_frames is None:
                nb_frames = stream.get("nb_frames")
                if nb_frames is not None:
                    video_total_frames = int(nb_frames)
                else:
                    duration = float(stream.get("duration", 0))
                    if duration > 0 and fps is not None:
                        video_total_frames = int(duration * fps)
        elif codec_type == "audio" and not audio_info:
            audio_info = {
                "codec": stream.get("codec_name", "N/A"),
                "sample_rate": stream.get("sample_rate", "N/A"),
                "channels": stream.get("channels", "N/A"),
                "bitrate": int(stream.get("bit_rate", 0)) if stream.get("bit_rate") else None,
            }
            audio_bitrate = audio_info["bitrate"]
    pixel_count = width * height if width and height else None
    return {
        "total_bitrate_bps": total_bitrate,
        "video_bitrate_bps": video_bitrate,
        "audio_bitrate_bps": audio_bitrate,
        "width": width,
        "height": height,
        "pixel_count": pixel_count,
        "fps": fps,
        "video_total_frames": video_total_frames,
        "audio_info": audio_info if audio_info else None,
    }

def get_frame_sizes(file_path, stream_index=0):
    cmd = [
        FFPROBE_PATH, "-v", "quiet",
        "-select_streams", f"v:{stream_index}",
        "-show_packets", "-show_entries", "packet=size",
        "-of", "json", file_path
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        packets = json.loads(result.stdout).get("packets", [])
        return [int(p["size"]) for p in packets]
    except:
        return []

def format_bps(bps):
    if bps is None: return "N/A"
    if bps >= 1_000_000: return f"{bps/1_000_000:.2f} Mbps"
    elif bps >= 1_000: return f"{bps/1_000:.2f} kbps"
    return f"{bps} bps"

# ---------- 检测函数 ----------
def get_sample_times(duration, interval):
    if duration is None or duration <= 0:
        return []
    times = []
    t = 0.0
    while t < duration:
        times.append(t)
        t += interval
    return times

def grab_frame_at_time(cap, time_sec):
    cap.set(cv2.CAP_PROP_POS_MSEC, time_sec * 1000)
    ret, frame = cap.read()
    return frame if ret else None

def detect_sharpness(video_path, sample_interval=2.0):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        return None
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration = total_frames / fps if fps > 0 else 0
    times = get_sample_times(duration, sample_interval)
    if not times:
        cap.release()
        return None
    scores = []
    for t in times:
        frame = grab_frame_at_time(cap, t)
        if frame is None:
            continue
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        laplacian_var = cv2.Laplacian(gray, cv2.CV_64F).var()
        scores.append(laplacian_var)
    cap.release()
    if not scores:
        return None
    avg = np.mean(scores)
    judge = "清晰" if avg > 200 else "一般" if avg > 100 else "模糊"
    return {"average": round(avg, 1), "judgement": judge}

def detect_exposure(video_path, sample_interval=2.0):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        return None
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration = total_frames / fps if fps > 0 else 0
    times = get_sample_times(duration, sample_interval)
    if not times:
        cap.release()
        return None
    avg_brightness = []
    over_count = 0
    under_count = 0
    total = 0
    for t in times:
        frame = grab_frame_at_time(cap, t)
        if frame is None:
            continue
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        mean_val = np.mean(gray)
        avg_brightness.append(mean_val)
        if mean_val > OVEREXPOSURE_THRESH:
            over_count += 1
        if mean_val < UNDEREXPOSURE_THRESH:
            under_count += 1
        total += 1
    cap.release()
    if total == 0:
        return None
    overall_mean = np.mean(avg_brightness)
    over_ratio = over_count / total
    under_ratio = under_count / total
    if over_ratio > 0.3:
        judge = "过曝"
    elif under_ratio > 0.3:
        judge = "欠曝"
    elif overall_mean < 80:
        judge = "偏暗"
    elif overall_mean > 170:
        judge = "偏亮"
    else:
        judge = "正常"
    return {
        "average_brightness": round(overall_mean, 1),
        "overexposed_ratio": round(over_ratio * 100, 1),
        "underexposed_ratio": round(under_ratio * 100, 1),
        "judgement": judge
    }

def detect_stability(video_path, sample_interval=0.5):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        return None
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration = total_frames / fps if fps > 0 else 0
    times = get_sample_times(duration, sample_interval)
    if len(times) < 2:
        cap.release()
        return None
    prev_frame = grab_frame_at_time(cap, times[0])
    if prev_frame is None:
        cap.release()
        return None
    prev_gray = cv2.cvtColor(prev_frame, cv2.COLOR_BGR2GRAY)
    motion_magnitudes = []
    for t in times[1:]:
        curr_frame = grab_frame_at_time(cap, t)
        if curr_frame is None:
            continue
        curr_gray = cv2.cvtColor(curr_frame, cv2.COLOR_BGR2GRAY)
        flow = cv2.calcOpticalFlowFarneback(prev_gray, curr_gray, None, 0.5, 3, 15, 3, 5, 1.2, 0)
        mag, _ = cv2.cartToPolar(flow[..., 0], flow[..., 1])
        mean_mag = np.mean(mag)
        motion_magnitudes.append(mean_mag)
        prev_gray = curr_gray
    cap.release()
    if not motion_magnitudes:
        return None
    avg_motion = np.mean(motion_magnitudes)
    std_motion = np.std(motion_magnitudes)
    if std_motion > STABILITY_SHAKE_THRESH:
        judge = "明显抖动"
    elif std_motion > STABILITY_SHAKE_THRESH * 0.5:
        judge = "轻微抖动"
    else:
        judge = "稳定"
    max_std = max(STABILITY_SHAKE_THRESH * 2, std_motion * 1.5)
    score = max(0, 100 - (std_motion / max_std) * 100)
    return {"score": round(score, 1), "avg_motion": round(avg_motion, 2), "std_motion": round(std_motion, 2), "judgement": judge}

def detect_frame_drops(video_path):
    cmd = [
        FFPROBE_PATH, "-v", "quiet",
        "-select_streams", "v:0",
        "-show_packets",
        "-show_entries", "packet=pts_time",
        "-of", "json",
        video_path
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        packets = json.loads(result.stdout).get("packets", [])
        if len(packets) < 2:
            return None
        pts_times = [float(p["pts_time"]) for p in packets if "pts_time" in p]
        intervals = np.diff(pts_times)
        if len(intervals) == 0:
            return None
        nominal_fps = 1 / np.median(intervals) if np.median(intervals) != 0 else None
        actual_fps = len(pts_times) / (pts_times[-1] - pts_times[0]) if (pts_times[-1] - pts_times[0]) > 0 else None
        median_interval = np.median(intervals)
        abnormal_count = np.sum((intervals > median_interval * 2) | (intervals < median_interval * 0.5))
        drop_ratio = abnormal_count / len(intervals)
        judge = "疑似掉帧" if drop_ratio > 0.05 else "正常"
        return {
            "nominal_fps": round(nominal_fps, 2) if nominal_fps else "N/A",
            "actual_fps": round(actual_fps, 2) if actual_fps else "N/A",
            "abnormal_ratio": round(drop_ratio * 100, 1),
            "judgement": judge
        }
    except:
        return None

def detect_black_white(video_path, sample_interval=2.0):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        return None
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration = total_frames / fps if fps > 0 else 0
    times = get_sample_times(duration, sample_interval)
    if not times:
        cap.release()
        return None
    black_count = 0
    white_count = 0
    total = 0
    black_times = []
    white_times = []
    for t in times:
        frame = grab_frame_at_time(cap, t)
        if frame is None:
            continue
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        mean_val = np.mean(gray)
        if mean_val < BLACK_THRESH:
            black_count += 1
            black_times.append(t)
        elif mean_val > WHITE_THRESH:
            white_count += 1
            white_times.append(t)
        total += 1
    cap.release()
    if total == 0:
        return None
    black_ratio = black_count / total
    white_ratio = white_count / total
    issues = []
    if black_ratio > 0.1:
        issues.append("黑屏占比过高")
    if white_ratio > 0.1:
        issues.append("白屏占比过高")
    return {
        "black_ratio": round(black_ratio * 100, 1),
        "white_ratio": round(white_ratio * 100, 1),
        "black_samples": [format_time(t) for t in black_times[:3]],
        "white_samples": [format_time(t) for t in white_times[:3]],
        "judgement": " / ".join(issues) if issues else "正常"
    }

def detect_frozen_frames(video_path, sample_interval=1.0, freeze_thresh=FREEZE_MOTION_THRESH, consecutive=FREEZE_THRESHOLD_FRAMES):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        return None
    fps = cap.get(cv2.CAP_PROP_FPS)
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    duration = total_frames / fps if fps > 0 else 0
    times = get_sample_times(duration, sample_interval)
    if len(times) < consecutive:
        cap.release()
        return None
    prev_frame = None
    freeze_count = 0
    frozen_segments = []
    start_time = None
    for t in times:
        frame = grab_frame_at_time(cap, t)
        if frame is None:
            continue
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        if prev_frame is not None:
            diff = cv2.absdiff(prev_frame, gray)
            mean_diff = np.mean(diff)
            if mean_diff < freeze_thresh:
                if start_time is None:
                    start_time = t
                freeze_count += 1
            else:
                if freeze_count >= consecutive and start_time is not None:
                    frozen_segments.append((format_time(start_time), format_time(t)))
                start_time = None
                freeze_count = 0
        prev_frame = gray
    if freeze_count >= consecutive and start_time is not None:
        frozen_segments.append((format_time(start_time), format_time(times[-1])))
    cap.release()
    judge = "疑似画面冻结" if frozen_segments else "正常"
    return {"frozen_segments": frozen_segments, "judgement": judge}

def analyze_audio(video_path):
    cmd = [
        "ffmpeg", "-i", video_path,
        "-af", "silencedetect=noise=-30dB:d=0.5",
        "-f", "null", "-"
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, stderr=subprocess.STDOUT)
        output = result.stdout + result.stderr
        silence_starts = []
        silence_ends = []
        for line in output.splitlines():
            if "silence_start" in line:
                s = line.split("silence_start:")[1].strip().split()[0]
                silence_starts.append(float(s))
            elif "silence_end" in line:
                e = line.split("silence_end:")[1].strip().split()[0]
                silence_ends.append(float(e))
        silence_duration = sum(silence_ends) - sum(silence_starts) if silence_starts else 0
        has_silence = len(silence_starts) > 0
        return {
            "has_audio": True,
            "silence_detected": has_silence,
            "silence_segments_count": len(silence_starts),
            "total_silence_sec": round(silence_duration, 2)
        }
    except:
        return {"has_audio": False}

def run_quality_checks(video_path, info):
    quality = {}
    quality["sharpness"] = detect_sharpness(video_path, SHARPNESS_SAMPLE_INTERVAL)
    quality["exposure"] = detect_exposure(video_path, SHARPNESS_SAMPLE_INTERVAL)
    quality["stability"] = detect_stability(video_path, STABILITY_SAMPLE_INTERVAL)
    quality["frame_drops"] = detect_frame_drops(video_path)
    quality["black_white"] = detect_black_white(video_path, SHARPNESS_SAMPLE_INTERVAL)
    quality["frozen"] = detect_frozen_frames(video_path)
    audio_info = info.get("audio_info")
    if audio_info:
        quality["audio"] = analyze_audio(video_path)
    else:
        quality["audio"] = {"has_audio": False}
    return quality

# ---------- Excel 生成 ----------
def apply_wrap_and_width(ws, headers_count):
    wrap_align = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=headers_count):
        for cell in row:
            cell.alignment = wrap_align
    for col_idx in range(1, headers_count + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        length = sum(2 if ord(c) > 127 else 1 for c in line)
                        max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max_len + 2

def create_summary_sheet(wb, all_results, mode_config, include_frames):
    ws = wb.active
    ws.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A1:K1")
    ws["A1"] = f"视频码率分析报告（{mode_config['name']}）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:K2")
    ws["A2"] = f"扫描文件夹: {INPUT_FOLDER}"
    ws["A2"].font = Font(italic=True)
    start_row = 4
    headers = ["视频文件", "分辨率", "总码率", "视频码率", "音频码率", "帧率 (fps)", "视频总帧数", "码率判断"]
    if include_frames:
        headers.extend(["提取包数", "平均帧大小(字节)", "最小帧大小(字节)", "最大帧大小(字节)"])
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for row_idx, (video_name, info, frame_sizes) in enumerate(all_results, start_row + 1):
        ws.cell(row=row_idx, column=1, value=video_name)
        width = info.get("width")
        height = info.get("height")
        resolution_str = f"{width}x{height}" if width and height else "N/A"
        ws.cell(row=row_idx, column=2, value=resolution_str)
        ws.cell(row=row_idx, column=3, value=format_bps(info.get("total_bitrate_bps")))
        ws.cell(row=row_idx, column=4, value=format_bps(info.get("video_bitrate_bps")))
        ws.cell(row=row_idx, column=5, value=format_bps(info.get("audio_bitrate_bps")))
        fps_val = info.get("fps")
        ws.cell(row=row_idx, column=6, value=f"{fps_val:.2f}" if fps_val else "N/A")
        total_frames = info.get("video_total_frames")
        ws.cell(row=row_idx, column=7, value=total_frames if total_frames else "N/A")
        pixel_count = info.get("pixel_count")
        threshold = get_bitrate_threshold_mbps(mode_config["standards"], pixel_count, fps_val) if pixel_count and fps_val else None
        judgement = judge_bitrate(info.get("video_bitrate_bps"), threshold)
        judge_cell = ws.cell(row=row_idx, column=8, value=judgement)
        judge_cell.alignment = center_align
        if judgement == "OK":
            judge_cell.fill = green_fill
        elif judgement == "NG":
            judge_cell.fill = red_fill
        else:
            judge_cell.fill = gray_fill
        if include_frames:
            base_col = 9
            if frame_sizes:
                packet_count = len(frame_sizes)
                avg_size = sum(frame_sizes) / packet_count if packet_count > 0 else 0
                min_s = min(frame_sizes) if packet_count > 0 else 0
                max_s = max(frame_sizes) if packet_count > 0 else 0
                ws.cell(row=row_idx, column=base_col, value=packet_count)
                ws.cell(row=row_idx, column=base_col+1, value=round(avg_size, 2))
                ws.cell(row=row_idx, column=base_col+2, value=min_s)
                ws.cell(row=row_idx, column=base_col+3, value=max_s)
            else:
                for i in range(4):
                    ws.cell(row=row_idx, column=base_col+i, value="N/A")
    apply_wrap_and_width(ws, len(headers))
    ws.freeze_panes = f"A{start_row + 1}"

def add_standard_sheet(wb, mode_config):
    ws = wb.create_sheet("码率评判标准", 1)
    ws.merge_cells("A1:E1")
    ws["A1"] = f"评判标准 - {mode_config['name']}"
    ws["A1"].font = Font(bold=True, size=13)
    headers = ["拍摄模式/分辨率", "分辨率", "帧率", "建议码率", "检测判断"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for row_idx, row_data in enumerate(mode_config["table"], 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col_idx in range(1, len(headers)+1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    length = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                    max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max_len + 3
    ws.freeze_panes = "A4"

def create_quality_sheet(wb, all_results, quality_results):
    ws = wb.create_sheet("视频质量综合检测", 2)
    headers = [
        "视频文件", "清晰度(Laplacian)", "清晰度判断",
        "平均亮度", "过曝%", "欠曝%", "曝光判断",
        "稳定性评分", "稳定性判断",
        "标称fps", "实际fps", "异常帧间隔%", "掉帧判断",
        "黑屏%", "白屏%", "黑屏/白屏判断",
        "画面冻结判断", "冻结时段",
        "音频存在", "音频编码", "静音检测"
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    for row_idx, (video_name, info, _) in enumerate(all_results, 2):
        q = quality_results[video_name]
        ws.cell(row=row_idx, column=1, value=video_name)
        sharp = q.get("sharpness")
        ws.cell(row=row_idx, column=2, value=sharp["average"] if sharp else "N/A")
        ws.cell(row=row_idx, column=3, value=sharp["judgement"] if sharp else "N/A")
        exp = q.get("exposure")
        ws.cell(row=row_idx, column=4, value=exp["average_brightness"] if exp else "N/A")
        ws.cell(row=row_idx, column=5, value=exp["overexposed_ratio"] if exp else "N/A")
        ws.cell(row=row_idx, column=6, value=exp["underexposed_ratio"] if exp else "N/A")
        ws.cell(row=row_idx, column=7, value=exp["judgement"] if exp else "N/A")
        stab = q.get("stability")
        ws.cell(row=row_idx, column=8, value=stab["score"] if stab else "N/A")
        ws.cell(row=row_idx, column=9, value=stab["judgement"] if stab else "N/A")
        drops = q.get("frame_drops")
        ws.cell(row=row_idx, column=10, value=drops["nominal_fps"] if drops else "N/A")
        ws.cell(row=row_idx, column=11, value=drops["actual_fps"] if drops else "N/A")
        ws.cell(row=row_idx, column=12, value=drops["abnormal_ratio"] if drops else "N/A")
        ws.cell(row=row_idx, column=13, value=drops["judgement"] if drops else "N/A")
        bw = q.get("black_white")
        ws.cell(row=row_idx, column=14, value=bw["black_ratio"] if bw else "N/A")
        ws.cell(row=row_idx, column=15, value=bw["white_ratio"] if bw else "N/A")
        ws.cell(row=row_idx, column=16, value=bw["judgement"] if bw else "N/A")
        frozen = q.get("frozen")
        ws.cell(row=row_idx, column=17, value=frozen["judgement"] if frozen else "N/A")
        frozen_segs = ", ".join([f"{s}-{e}" for s, e in frozen["frozen_segments"]]) if frozen and frozen["frozen_segments"] else "无"
        ws.cell(row=row_idx, column=18, value=frozen_segs)
        audio = q.get("audio")
        if audio and audio.get("has_audio"):
            audio_info = info.get("audio_info")
            ws.cell(row=row_idx, column=19, value="有")
            ws.cell(row=row_idx, column=20, value=audio_info.get("codec", "") if audio_info else "")
            silence_str = f"检测到{audio['silence_segments_count']}段静音" if audio.get("silence_detected") else "无静音"
            ws.cell(row=row_idx, column=21, value=silence_str)
        else:
            ws.cell(row=row_idx, column=19, value="无")
            ws.cell(row=row_idx, column=20, value="N/A")
            ws.cell(row=row_idx, column=21, value="N/A")
    for col_idx in range(1, len(headers)+1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    length = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                    max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max_len + 3
    ws.freeze_panes = "A2"

def add_frame_sheet(wb, video_name, frame_sizes):
    sheet_name = re.sub(r'[\\/*?:\[\]]', '_', os.path.splitext(video_name)[0])[:31]
    original = sheet_name
    counter = 1
    while sheet_name in wb.sheetnames:
        sheet_name = f"{original}_{counter}"[:31]
        counter += 1
    ws = wb.create_sheet(title=sheet_name)
    bold_font = Font(bold=True)
    ws.append(["帧索引", "帧大小 (字节)"])
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["B1"].alignment = Alignment(horizontal="center")
    for idx, size in enumerate(frame_sizes, 1):
        ws.append([idx, size])
    apply_wrap_and_width(ws, 2)
    ws.freeze_panes = "A2"

# ---------- 主程序 ----------
def main():
    mode = select_mode()
    mode_config = STANDARDS_CONFIG[mode]
    print(f"当前模式：{mode_config['name']}\n")
    if not os.path.isdir(INPUT_FOLDER):
        print(f"错误: 文件夹不存在 - {INPUT_FOLDER}")
        sys.exit(1)
    video_files = []
    for fname in os.listdir(INPUT_FOLDER):
        fp = os.path.join(INPUT_FOLDER, fname)
        if os.path.isfile(fp) and os.path.splitext(fname)[1].lower() in VIDEO_EXTENSIONS:
            video_files.append(fp)
    if not video_files:
        print("未找到支持的视频文件。")
        return
    print(f"找到 {len(video_files)} 个视频，开始分析...\n")
    all_results = []
    quality_results = {}
    for vpath in video_files:
        vname = os.path.basename(vpath)
        print(f"分析: {vname}")
        info = extract_video_info(vpath)
        if not info:
            print("  无法提取媒体信息，跳过。")
            continue
        frame_sizes = get_frame_sizes(vpath) if INCLUDE_FRAMES else None
        all_results.append((vname, info, frame_sizes))
        print("  执行质量检测...")
        quality = run_quality_checks(vpath, info)
        quality_results[vname] = quality
        print(f"  清晰度: {quality['sharpness']['judgement'] if quality['sharpness'] else 'N/A'}, "
              f"曝光: {quality['exposure']['judgement'] if quality['exposure'] else 'N/A'}, "
              f"稳定性: {quality['stability']['judgement'] if quality['stability'] else 'N/A'}")
    wb = Workbook()
    create_summary_sheet(wb, all_results, mode_config, INCLUDE_FRAMES)
    add_standard_sheet(wb, mode_config)
    create_quality_sheet(wb, all_results, quality_results)
    if INCLUDE_FRAMES:
        for vname, _, frames in all_results:
            if frames:
                add_frame_sheet(wb, vname, frames)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"{OUTPUT_EXCEL_NAME}_{timestamp}.xlsx"
    out_path = os.path.join(INPUT_FOLDER, out_name)
    wb.save(out_path)
    print(f"\n报告已生成: {out_path}")

if __name__ == "__main__":
    main()