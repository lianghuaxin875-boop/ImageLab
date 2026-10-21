# -*- coding: utf-8 -*-
"""
视频码率提取与质量检测模块。

供 Web GUI 调用：扫描目录中的视频，使用 ffprobe 提取码率、分辨率、帧率和帧数，
按所选拍摄模式判断视频码率是否达标，并可生成 Excel 报告。
"""

import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_FFPROBE_PATH = (
    os.environ.get("FFPROBE_PATH")
    or shutil.which("ffprobe")
    or r"D:\码率工具FFepeg\ffmpeg-8.1.1-full_build\bin\ffprobe.exe"
)
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".mpg", ".mpeg"}

STANDARDS_CONFIG = {
    1: {
        "name": "普通手机拍摄模式",
        "standards": [
            (921600, 30, 4.0, None),
            (2073600, 30, 8.0, 12.0),
            (8294400, 30, 35.0, 50.0),
            (float("inf"), 30, 100.0, 150.0),
        ],
        "table": [
            ["720P 30fps", "1280x720", "30", "4 - 6 Mbps", "低于 4 Mbps 偏低"],
            ["1080P 30fps", "1920x1080", "30", "8 - 12 Mbps", "低于 8 Mbps 偏低"],
            ["1080P 60fps", "1920x1080", "60", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 30fps", "3840x2160", "30", "35 - 50 Mbps", "低于 35 Mbps 偏低"],
            ["4K 60fps", "3840x2160", "60", "50 - 80 Mbps", "低于 50 Mbps 偏低"],
        ],
    },
    2: {
        "name": "短视频平台拍摄/上传模式",
        "standards": [
            (921600, 30, 3.0, None),
            (2073600, 30, 6.0, 10.0),
            (3686400, 30, 12.0, 16.0),
            (8294400, 30, 30.0, 45.0),
            (float("inf"), 30, 30.0, 45.0),
        ],
        "table": [
            ["720P 竖屏", "720x1280", "30fps", "3 - 6 Mbps", "低于 3 Mbps 偏低"],
            ["1080P 竖屏", "1080x1920", "30fps", "6 - 10 Mbps", "低于 6 Mbps 偏低"],
            ["1080P 竖屏高帧率", "1080x1920", "60fps", "10 - 16 Mbps", "低于 10 Mbps 偏低"],
            ["2K 竖屏", "1440x2560", "30fps", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 竖屏", "2160x3840", "30fps", "30 - 45 Mbps", "低于 30 Mbps 偏低"],
        ],
    },
    3: {
        "name": "监控/低码率存储模式",
        "standards": [
            (921600, 30, 1.0, None),
            (2073600, 30, 2.0, None),
            (3686400, 30, 4.0, None),
            (8294400, 30, 8.0, None),
            (float("inf"), 30, 8.0, None),
        ],
        "table": [
            ["720P 监控", "1280x720", "15/25fps", "1 - 3 Mbps", "低于 1 Mbps 偏低"],
            ["1080P 监控", "1920x1080", "15/25fps", "2 - 6 Mbps", "低于 2 Mbps 偏低"],
            ["2K 监控", "2560x1440", "15/25fps", "4 - 10 Mbps", "低于 4 Mbps 偏低"],
            ["4K 监控", "3840x2160", "15/25fps", "8 - 20 Mbps", "低于 8 Mbps 偏低"],
        ],
    },
    4: {
        "name": "高质量拍摄/相机素材模式",
        "standards": [
            (2073600, 30, 20.0, 30.0),
            (8294400, 30, 60.0, 100.0),
            (18000000, 30, 200.0, 200.0),
            (float("inf"), 30, 200.0, 200.0),
        ],
        "table": [
            ["1080P 30fps", "1920x1080", "30", "20 - 50 Mbps", "低于 20 Mbps 偏低"],
            ["1080P 60fps", "1920x1080", "60", "30 - 80 Mbps", "低于 30 Mbps 偏低"],
            ["4K 24/25/30fps", "3840x2160", "24/25/30", "60 - 100 Mbps", "低于 60 Mbps 偏低"],
            ["4K 50/60fps", "3840x2160", "50/60", "100 - 200 Mbps", "低于 100 Mbps 偏低"],
            ["4K 10bit/LOG", "3840x2160", "24/25/30", "100 - 400 Mbps", "视编码而定"],
            ["6K/8K 高质量", "6K/8K", "24/30/60", "200 Mbps 以上", "低于 200 Mbps 偏低"],
        ],
    },
}


def get_bitrate_threshold_mbps(standards, pixel_count: Optional[int], fps: Optional[float]) -> Optional[float]:
    if pixel_count is None or fps is None:
        return None
    for max_pixels, fps_limit, normal_min, high_min in standards:
        if pixel_count <= max_pixels:
            return high_min if high_min is not None and fps > fps_limit else normal_min
    return None


def judge_bitrate(video_bitrate_bps: Optional[int], threshold_mbps: Optional[float]) -> str:
    if video_bitrate_bps is None or threshold_mbps is None:
        return "N/A"
    return "OK" if video_bitrate_bps >= threshold_mbps * 1_000_000 else "NG"


def format_bps(bps: Optional[int]) -> str:
    if bps is None:
        return "N/A"
    if bps >= 1_000_000:
        return f"{bps / 1_000_000:.2f} Mbps"
    if bps >= 1_000:
        return f"{bps / 1_000:.2f} kbps"
    return f"{bps} bps"


def parse_fps(rate: Optional[str]) -> Optional[float]:
    if not rate:
        return None
    try:
        if "/" in rate:
            num, den = rate.split("/")
            den_float = float(den)
            return float(num) / den_float if den_float else None
        return float(rate)
    except (ValueError, ZeroDivisionError):
        return None


def sanitize_sheet_name(name: str) -> str:
    return re.sub(r'[\\/*?:\[\]]', "_", name)[:31]


def score_white_balance_diff(max_diff: float) -> float:
    if max_diff <= 10:
        return round(100 - max_diff / 10 * 5, 1)
    if max_diff <= 15:
        return round(94 - (max_diff - 11) / 4 * 9, 1)
    if max_diff <= 25:
        return round(84 - (max_diff - 16) / 9 * 14, 1)
    if max_diff <= 35:
        return round(69 - (max_diff - 26) / 9 * 9, 1)
    return round(max(0, 59 - (max_diff - 35) / 65 * 59), 1)


def white_balance_level_from_score(score: float) -> str:
    if score >= 95:
        return "优秀"
    if score >= 85:
        return "OK"
    if score >= 70:
        return "轻微偏色"
    if score >= 60:
        return "偏色明显"
    return "NG"


def video_white_balance_level(abnormal_ratio: float) -> str:
    if abnormal_ratio <= 0.10:
        return "OK"
    if abnormal_ratio <= 0.30:
        return "警告：部分片段偏色"
    return "NG：视频整体白平衡异常"


class VideoBitrateDetector:
    def __init__(
        self,
        input_dir: str,
        report_dir: str,
        mode: int,
        ffprobe_path: str = DEFAULT_FFPROBE_PATH,
        include_frames: bool = False,
        log_callback: Optional[Callable[[str], None]] = None,
        progress_callback: Optional[Callable[[int, int], None]] = None,
    ):
        self.input_dir = Path(input_dir)
        self.report_dir = Path(report_dir)
        self.mode = int(mode)
        self.ffprobe_path = ffprobe_path
        self.include_frames = include_frames
        self.log = log_callback or (lambda message: None)
        self.progress = progress_callback or (lambda done, total: None)

    @property
    def mode_config(self):
        return STANDARDS_CONFIG[self.mode]

    def run_ffprobe(self, args: List[str], packet_mode: bool = False) -> dict:
        if not Path(self.ffprobe_path).exists():
            raise FileNotFoundError(f"未找到 ffprobe：{self.ffprobe_path}")
        base = [self.ffprobe_path, "-v", "quiet", "-print_format", "json"]
        if packet_mode:
            cmd = base + ["-show_packets", "-show_entries", "packet=size"] + args
        else:
            cmd = base + ["-show_format", "-show_streams"] + args
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return json.loads(result.stdout or "{}")

    def extract_video_info(self, file_path: Path) -> Dict:
        info = self.run_ffprobe([str(file_path)])
        fmt = info.get("format", {})
        total_bitrate = int(fmt["bit_rate"]) if fmt.get("bit_rate") else None
        video_bitrate = None
        audio_bitrate = None
        width = height = fps = video_total_frames = None

        for stream in info.get("streams", []):
            codec_type = stream.get("codec_type")
            if codec_type == "video":
                width = width or stream.get("width")
                height = height or stream.get("height")
                if video_bitrate is None and stream.get("bit_rate"):
                    video_bitrate = int(stream["bit_rate"])
                if fps is None:
                    fps = parse_fps(stream.get("r_frame_rate")) or parse_fps(stream.get("avg_frame_rate"))
                if video_total_frames is None:
                    if stream.get("nb_frames"):
                        video_total_frames = int(stream["nb_frames"])
                    elif stream.get("duration") and fps:
                        video_total_frames = int(float(stream["duration"]) * fps)
            elif codec_type == "audio" and audio_bitrate is None and stream.get("bit_rate"):
                audio_bitrate = int(stream["bit_rate"])

        pixel_count = width * height if width and height else None
        threshold = get_bitrate_threshold_mbps(self.mode_config["standards"], pixel_count, fps)
        judgement = judge_bitrate(video_bitrate, threshold)
        return {
            "文件名": file_path.name,
            "分辨率": f"{width}x{height}" if width and height else "N/A",
            "宽度": width,
            "高度": height,
            "像素数": pixel_count,
            "总码率": format_bps(total_bitrate),
            "视频码率": format_bps(video_bitrate),
            "音频码率": format_bps(audio_bitrate),
            "总码率_bps": total_bitrate,
            "视频码率_bps": video_bitrate,
            "音频码率_bps": audio_bitrate,
            "帧率": round(fps, 2) if fps else None,
            "视频总帧数": video_total_frames,
            "阈值Mbps": threshold,
            "码率判断": judgement,
            "检测模式": self.mode_config["name"],
        }

    def get_frame_sizes(self, file_path: Path) -> List[int]:
        data = self.run_ffprobe(["-select_streams", "v:0", str(file_path)], packet_mode=True)
        return [int(packet["size"]) for packet in data.get("packets", []) if packet.get("size")]

    def analyze_video_white_balance(self, file_path: Path, sample_count: int = 30) -> Dict:
        cap = cv2.VideoCapture(str(file_path))
        if not cap.isOpened():
            return {
                "白平衡评分": "",
                "视频白平衡判断": "N/A",
                "白平衡异常帧占比": "",
                "白平衡抽帧数量": 0,
                "视频白平衡说明": "视频无法打开，未进行白平衡抽帧检测",
            }

        total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        if total_frames > 0:
            frame_positions = np.linspace(0, max(total_frames - 1, 0), min(sample_count, total_frames), dtype=int)
        else:
            frame_positions = np.arange(sample_count, dtype=int)

        scores = []
        abnormal = 0
        sampled = 0
        for pos in frame_positions:
            if total_frames > 0:
                cap.set(cv2.CAP_PROP_POS_FRAMES, int(pos))
            ok, frame = cap.read()
            if not ok or frame is None:
                continue
            b_mean, g_mean, r_mean = [float(value) for value in cv2.mean(frame)[:3]]
            max_diff = max(r_mean, g_mean, b_mean) - min(r_mean, g_mean, b_mean)
            frame_score = score_white_balance_diff(max_diff)
            scores.append(frame_score)
            sampled += 1
            if frame_score < 70:
                abnormal += 1
        cap.release()

        if not sampled:
            return {
                "白平衡评分": "",
                "视频白平衡判断": "N/A",
                "白平衡异常帧占比": "",
                "白平衡抽帧数量": 0,
                "视频白平衡说明": "未成功抽取有效帧，无法判断视频白平衡",
            }

        avg_score = round(float(np.mean(scores)), 1)
        abnormal_ratio = abnormal / sampled
        judgement = video_white_balance_level(abnormal_ratio)
        return {
            "白平衡评分": avg_score,
            "视频白平衡判断": judgement,
            "白平衡异常帧占比": round(abnormal_ratio, 4),
            "白平衡抽帧数量": sampled,
            "视频白平衡说明": f"抽帧 {sampled} 帧，异常帧 {abnormal} 帧，平均帧白平衡状态：{white_balance_level_from_score(avg_score)}",
        }

    def run(self) -> Dict:
        if self.mode not in STANDARDS_CONFIG:
            raise ValueError("视频检测模式无效。")
        if not self.input_dir.is_dir():
            raise FileNotFoundError(f"视频输入目录不存在：{self.input_dir}")
        self.report_dir.mkdir(parents=True, exist_ok=True)

        files = [p for p in self.input_dir.iterdir() if p.is_file() and p.suffix.lower() in VIDEO_EXTENSIONS]
        if not files:
            raise RuntimeError(f"在 {self.input_dir} 中未找到支持的视频文件。")

        self.log(f"视频检测模式：{self.mode_config['name']}")
        self.log(f"找到 {len(files)} 个视频文件")
        results = []
        frame_data = {}
        for index, video_path in enumerate(files, 1):
            self.log(f"分析视频：{video_path.name}")
            try:
                info = self.extract_video_info(video_path)
                info.update(self.analyze_video_white_balance(video_path))
                if self.include_frames:
                    sizes = self.get_frame_sizes(video_path)
                    frame_data[video_path.name] = sizes
                    if sizes:
                        info["提取包数"] = len(sizes)
                        info["平均帧大小"] = round(sum(sizes) / len(sizes), 2)
                        info["最小帧大小"] = min(sizes)
                        info["最大帧大小"] = max(sizes)
                results.append(info)
            except Exception as exc:
                self.log(f"分析失败：{video_path.name}，原因：{exc}")
                results.append({"文件名": video_path.name, "码率判断": "ERROR", "错误信息": str(exc)})
            self.progress(index, len(files))

        return {"results": results, "frame_data": frame_data, "mode_config": self.mode_config}

    @staticmethod
    def _display_width(value) -> int:
        return sum(2 if ord(char) > 255 else 1 for char in str(value or ""))

    @classmethod
    def _style_sheet(cls, ws, table_name: Optional[str] = None, header_row: int = 1):
        border = Border(
            left=Side("thin", color="D9E2F3"),
            right=Side("thin", color="D9E2F3"),
            top=Side("thin", color="D9E2F3"),
            bottom=Side("thin", color="D9E2F3"),
        )
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_width = max(cls._display_width(ws.cell(row, col_idx).value) for row in range(1, ws.max_row + 1))
            ws.column_dimensions[letter].width = min(max(max_width + 2, 12), 38)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"A{header_row + 1}"

    @staticmethod
    def _summarize_video_report(results: List[Dict]) -> str:
        total = len(results)
        if not total:
            return "本次未检测到有效视频文件。"

        def has_issue(record):
            wb_judgement = str(record.get("视频白平衡判断", "OK"))
            return (
                record.get("码率判断") in ("NG", "ERROR", "N/A")
                or wb_judgement.startswith("警告")
                or wb_judgement.startswith("NG")
            )

        issue_count = sum(1 for record in results if has_issue(record))
        ok_count = total - issue_count
        issue_ratio = issue_count / total * 100 if total else 0
        issue_names = []
        if any(record.get("码率判断") == "NG" for record in results):
            issue_names.append("视频码率偏低")
        if any(record.get("码率判断") == "N/A" for record in results):
            issue_names.append("视频码率信息不足")
        if any(record.get("码率判断") == "ERROR" for record in results):
            issue_names.append("视频解析失败")
        if any(str(record.get("视频白平衡判断", "")).startswith("警告") for record in results):
            issue_names.append("部分视频片段偏色")
        if any(str(record.get("视频白平衡判断", "")).startswith("NG") for record in results):
            issue_names.append("视频整体白平衡异常")
        main_issues = "、".join(issue_names) if issue_names else "未发现集中性问题"

        return (
            f"本次共检测 {total} 个文件，其中图片 0 张、视频 {total} 个。\n"
            f"检测结果 OK {ok_count} 个，NG/需关注 {issue_count} 个，NG/需关注占比 {issue_ratio:.1f}%。\n"
            f"主要问题为：{main_issues}。"
        )

    def write_report(self, results: List[Dict], frame_data: Optional[Dict[str, List[int]]] = None) -> Path:
        self.report_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = self.report_dir / f"video_bitrate_report_{timestamp}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        headers = [
            "文件名", "分辨率", "总码率", "视频码率", "音频码率", "帧率", "视频总帧数",
            "阈值Mbps", "码率判断", "白平衡评分", "视频白平衡判断", "白平衡异常帧占比", "白平衡抽帧数量",
            "视频白平衡说明", "检测模式", "错误信息",
        ]
        ws.merge_cells("A1:P3")
        summary_cell = ws["A1"]
        summary_cell.value = self._summarize_video_report(results)
        summary_cell.font = Font(bold=True, size=12, color="1F4E78")
        summary_cell.fill = PatternFill(start_color="EAF4FF", end_color="EAF4FF", fill_type="solid")
        summary_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[2].height = 30
        ws.row_dimensions[3].height = 30
        header_row = 5
        for col_idx, header in enumerate(headers, 1):
            ws.cell(header_row, col_idx, header)
        fills = {
            "OK": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
            "NG": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
            "ERROR": PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid"),
            "N/A": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
        }
        for row_idx, record in enumerate(results, header_row + 1):
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row_idx, col_idx, record.get(header, ""))
            judge_cell = ws.cell(row_idx, headers.index("码率判断") + 1)
            judge_cell.fill = fills.get(judge_cell.value, fills["N/A"])
            judge_cell.alignment = Alignment(horizontal="center", vertical="center")
        self._style_sheet(ws, "VideoSummaryTable", header_row=header_row)

        std = wb.create_sheet("码率评判标准")
        std.append(["拍摄模式/分辨率", "分辨率", "帧率", "建议码率", "检测判断"])
        for row in self.mode_config["table"]:
            std.append(row)
        std.append([])
        std.append(["视频白平衡标准", "异常帧占比", "判断", "说明", ""])
        std.append(["视频抽帧白平衡", "≤10%", "OK", "视频整体白平衡稳定", ""])
        std.append(["视频抽帧白平衡", "10%-30%", "警告：部分片段偏色", "部分片段存在明显偏色", ""])
        std.append(["视频抽帧白平衡", ">30%", "NG：视频整体白平衡异常", "视频整体白平衡异常", ""])
        std.append([])
        std.append(["单帧白平衡评分", "RGB最大差值", "评分", "判断", ""])
        std.append(["单帧白平衡评分", "0-10", "95-100", "优秀", ""])
        std.append(["单帧白平衡评分", "11-15", "85-94", "OK", ""])
        std.append(["单帧白平衡评分", "16-25", "70-84", "轻微偏色", ""])
        std.append(["单帧白平衡评分", "26-35", "60-69", "偏色明显", ""])
        std.append(["单帧白平衡评分", ">35", "<60", "NG", ""])
        std.append([])
        std.append(["白平衡适用范围说明", "基于画面整体 RGB 通道差异进行估算", "适用于普通自然光、室内光、商品图、测试样张等场景", "夕阳、舞台灯光、蓝天、草地等强色彩场景可能存在误判", ""])
        self._style_sheet(std, "VideoStandardTable")

        if self.include_frames and frame_data:
            for video_name, sizes in frame_data.items():
                if not sizes:
                    continue
                sheet_name = sanitize_sheet_name(Path(video_name).stem)
                original = sheet_name
                counter = 1
                while sheet_name in wb.sheetnames:
                    sheet_name = f"{original}_{counter}"[:31]
                    counter += 1
                fs = wb.create_sheet(sheet_name)
                fs.append(["帧索引", "帧大小(字节)"])
                for idx, size in enumerate(sizes, 1):
                    fs.append([idx, size])
                self._style_sheet(fs)

        wb.save(report_path)
        return report_path
