# -*- coding: utf-8 -*-
"""
图片质量检测工具 Web 版

运行：
    python image_quality_web.py

浏览器访问：
    http://127.0.0.1:5000
"""

import os
import shutil
import threading
import time
import uuid
import importlib.util
import csv
import io
from types import SimpleNamespace
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from PIL import Image as PILImage, UnidentifiedImageError
from flask import Flask, Response, jsonify, render_template, request, send_file
from werkzeug.utils import secure_filename

from image_quality_gui import (
    DETECTOR_VERSION,
    IMAGE_EXTENSIONS,
    WEIGHT_MODES,
    ImageQualityDetector,
    artifact_desc_sug,
    calc_artifact,
    calc_clarity_metrics,
    calc_noise,
    clarity_desc_sug,
    contrast_desc_sug,
    exposure_desc_sug,
    level_from_score,
    noise_desc_sug,
    primary_issue,
    region_names_for,
    score_brightness,
    score_dynamic_range,
    score_laplacian,
    score_noise,
    score_overexposure,
    score_std,
    score_tenengrad,
)
from video_bitrate_tool import DEFAULT_FFPROBE_PATH, STANDARDS_CONFIG, VIDEO_EXTENSIONS, VideoBitrateDetector
from video_quality_scoring_tool import VideoQualityScoringDetector
from image_enhance.enhance_main import run_enhancement


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = BASE_DIR / "input_images"
DEFAULT_OK_DIR = BASE_DIR / "OK_images"
DEFAULT_NG_DIR = BASE_DIR / "NG_images"
DEFAULT_REPORT_DIR = BASE_DIR / "report"
DEFAULT_VIDEO_INPUT_DIR = BASE_DIR / "video_input"
DEFAULT_VIDEO_REPORT_DIR = BASE_DIR / "video_report"
DEFAULT_VIDEO_QUALITY_REPORT_DIR = BASE_DIR / "video_quality_report"
UPLOAD_DIR = BASE_DIR / "web_uploads"
VIDEO_UPLOAD_DIR = BASE_DIR / "web_video_uploads"
OCR_UPLOAD_DIR = BASE_DIR / "web_ocr_uploads"
OCR_OUTPUT_DIR = BASE_DIR / "web_ocr_results"
ENHANCE_UPLOAD_DIR = BASE_DIR / "web_enhance_uploads"
ENHANCE_OUTPUT_DIR = BASE_DIR / "web_enhance_results"
OCR_SCRIPT_PATH = BASE_DIR / "OCR Image Recognition2.0.py"
OCR_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
OCR_VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv", ".mpeg", ".mpg"}
OCR_EXTENSIONS = OCR_IMAGE_EXTENSIONS | OCR_VIDEO_EXTENSIONS

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024 * 1024

jobs = {}
jobs_lock = threading.Lock()
image_job_lock = threading.Lock()
ocr_module_cache = None
ocr_module_lock = threading.Lock()


def make_job():
    job_id = uuid.uuid4().hex
    with jobs_lock:
        jobs[job_id] = {
            "id": job_id,
            "status": "queued",
            "message": "任务已创建",
            "done": 0,
            "total": 0,
            "percent": 0,
            "eta_seconds": None,
            "elapsed_seconds": 0,
            "logs": [],
            "report_path": "",
            "results": [],
            "output_info": {},
            "config": {},
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    return job_id


def update_job(job_id, **changes):
    with jobs_lock:
        job = jobs[job_id]
        job.update(changes)


def append_log(job_id, message):
    with jobs_lock:
        job = jobs[job_id]
        now = datetime.now().strftime("%H:%M:%S")
        job["logs"].append(f"[{now}] {message}")
        job["logs"] = job["logs"][-300:]
        job["message"] = message
    # Keep phase diagnostics even if the hosting platform restarts the worker.
    print(f"[job {job_id}] {message}", flush=True)


def snapshot_job(job_id):
    with jobs_lock:
        job = jobs.get(job_id)
        return dict(job) if job else None


def json_safe(value):
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_safe(item) for item in value]
    if isinstance(value, tuple):
        return [json_safe(item) for item in value]
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    return value


def result_summary(results):
    rows = []
    for record in results:
        rows.append({
            "filename": record.get("文件名", ""),
            "score": float(record.get("综合评分", 0) or 0),
            "level": record.get("综合等级", ""),
            "clarity": float(record.get("清晰度得分", 0) or 0),
            "exposure": float(record.get("曝光得分", 0) or 0),
            "contrast": float(record.get("对比度得分", 0) or 0),
            "noise": float(record.get("噪声得分", 0) or 0),
            "artifact": float(record.get("伪影得分", 0) or 0),
            "summary": record.get("综合说明", ""),
            "location": record.get("存放位置", ""),
            "output_path": record.get("输出图片路径", ""),
            "output_status": record.get("标注图状态", ""),
            "color_cast": record.get("色偏结果", ""),
            "white_balance_score": record.get("白平衡评分", ""),
            "white_balance": record.get("白平衡状态", ""),
            "histogram": record.get("亮度直方图判断", ""),
        })
    return rows


def output_info(config, results):
    saved = [r for r in results if r.get("输出图片路径")]
    failed = [r for r in results if r.get("标注图状态", "").startswith("保存失败")]
    return {
        "ok_dir": config.get("ok_dir", ""),
        "ng_dir": config.get("ng_dir", ""),
        "report_dir": config.get("report_dir", ""),
        "saved_count": len(saved),
        "failed_count": len(failed),
    }


def parse_weights(form):
    mode = form.get("weight_mode", "默认综合模式")
    if mode in WEIGHT_MODES:
        return mode, WEIGHT_MODES[mode]

    labels = ["clarity", "exposure", "contrast", "noise", "artifact"]
    weights = tuple(float(form.get(f"weight_{name}", "0")) for name in labels)
    if any(weight < 0 for weight in weights) or abs(sum(weights) - 1.0) > 0.001:
        raise ValueError("自定义权重必须为非负数，且总和必须为 1.0。")
    return "自定义模式", weights


def save_uploads(files):
    if not files:
        return None

    batch_dir = UPLOAD_DIR / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    batch_dir.mkdir(parents=True, exist_ok=True)
    saved = 0

    for file_storage in files:
        if not file_storage or not file_storage.filename:
            continue
        suffix = Path(file_storage.filename).suffix.lower()
        if suffix not in IMAGE_EXTENSIONS:
            continue
        filename = secure_filename(file_storage.filename)
        if not filename:
            filename = f"image_{saved + 1}{suffix}"
        file_storage.save(batch_dir / filename)
        saved += 1

    return batch_dir if saved else None


def save_video_uploads(files):
    if not files:
        return None

    batch_dir = VIDEO_UPLOAD_DIR / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    batch_dir.mkdir(parents=True, exist_ok=True)
    saved = 0

    for file_storage in files:
        if not file_storage or not file_storage.filename:
            continue
        suffix = Path(file_storage.filename).suffix.lower()
        if suffix not in VIDEO_EXTENSIONS:
            continue
        filename = secure_filename(file_storage.filename)
        if not filename:
            filename = f"video_{saved + 1}{suffix}"
        file_storage.save(batch_dir / filename)
        saved += 1

    return batch_dir if saved else None


def save_ocr_uploads(files):
    if not files:
        return None

    batch_dir = OCR_UPLOAD_DIR / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    batch_dir.mkdir(parents=True, exist_ok=True)
    saved = 0

    for file_storage in files:
        if not file_storage or not file_storage.filename:
            continue
        suffix = Path(file_storage.filename).suffix.lower()
        if suffix not in OCR_EXTENSIONS:
            continue
        filename = secure_filename(file_storage.filename)
        if not filename:
            filename = f"ocr_file_{saved + 1}{suffix}"
        file_storage.save(batch_dir / filename)
        saved += 1

    return batch_dir if saved else None


def save_enhance_uploads(files):
    if not files:
        return None

    batch_dir = ENHANCE_UPLOAD_DIR / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    batch_dir.mkdir(parents=True, exist_ok=True)
    saved = 0

    for file_storage in files:
        if not file_storage or not file_storage.filename:
            continue
        suffix = Path(file_storage.filename).suffix.lower()
        if suffix not in IMAGE_EXTENSIONS:
            continue
        filename = secure_filename(file_storage.filename)
        if not filename:
            filename = f"enhance_image_{saved + 1}{suffix}"
        file_storage.save(batch_dir / filename)
        saved += 1

    return batch_dir if saved else None


def load_ocr_module():
    global ocr_module_cache
    with ocr_module_lock:
        if ocr_module_cache is not None:
            return ocr_module_cache
        if not OCR_SCRIPT_PATH.exists():
            raise FileNotFoundError(f"OCR 2.0 脚本不存在：{OCR_SCRIPT_PATH}")
        spec = importlib.util.spec_from_file_location("ocr_image_recognition_2", OCR_SCRIPT_PATH)
        if spec is None or spec.loader is None:
            raise RuntimeError("无法加载 OCR 2.0 脚本")
        module = importlib.util.module_from_spec(spec)
        try:
            spec.loader.exec_module(module)
        except SystemExit as exc:
            raise RuntimeError("OCR 2.0 依赖未安装，请先安装 paddleocr / paddlepaddle / paddlex") from exc
        ocr_module_cache = module
        return module


def ocr_result_summary(records):
    rows = []
    for record in records:
        rows.append({
            "filename": record.get("image_name", ""),
            "preprocess": record.get("preprocess", ""),
            "quality_score": record.get("quality_score", ""),
            "quality_level": record.get("quality_level", ""),
            "review_required": record.get("review_required", False),
            "review_advice": record.get("review_advice", ""),
            "text": record.get("full_text", ""),
            "ocr_count": record.get("ocr_count", 0),
            "avg_confidence": record.get("avg_confidence", 0),
            "min_confidence": record.get("min_confidence", 0),
            "low_conf_ratio": record.get("low_conf_ratio", 0),
            "hit_rate": record.get("hit_rate", 0),
            "status": record.get("status", ""),
            "sub_codes": ", ".join(record.get("sub_codes", [])),
            "reasons": "；".join(record.get("reasons", [])),
        })
    return rows


def assess_ocr_image_quality(image, threshold=70.0):
    if image is None:
        return {
            "score": 0.0,
            "level": "ERROR",
            "issue": "图片无法读取",
            "review_required": True,
            "advice": "识别结果建议人工复核",
        }

    height, width = image.shape[:2]
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY) if len(image.shape) == 3 else image
    weights = WEIGHT_MODES.get("默认综合模式", (0.4, 0.2, 0.2, 0.1, 0.1))

    lap_val, ten_val = calc_clarity_metrics(gray)
    lap_score = round(score_laplacian(lap_val), 1)
    ten_score = round(score_tenengrad(ten_val), 1)
    clarity_score = round(lap_score * 0.5 + ten_score * 0.5, 1)
    clarity_desc, _ = clarity_desc_sug(lap_score, ten_score, clarity_score)

    gray_min, gray_max = int(gray.min()), int(gray.max())
    gray_range = gray_max - gray_min
    std_val = np.std(gray)
    std_score = round(score_std(std_val), 1)
    range_score = round(score_dynamic_range(gray_range), 1)
    contrast_score = round(std_score * 0.7 + range_score * 0.3, 1)
    contrast_desc, _ = contrast_desc_sug(std_score, range_score, contrast_score)

    mean_bright = gray.mean()
    over_ratio = np.sum(gray > 240) / gray.size
    under_ratio = np.sum(gray < 30) / gray.size
    bright_score = round(score_brightness(mean_bright), 1)
    over_score = round(score_overexposure(over_ratio), 1)
    under_score = round(score_overexposure(under_ratio), 1)
    exposure_score = round(bright_score * 0.6 + over_score * 0.2 + under_score * 0.2, 1)
    exposure_desc, _ = exposure_desc_sug(mean_bright, bright_score, over_score, under_score, exposure_score)

    noise_val = calc_noise(gray)
    noise_score = round(score_noise(noise_val), 1)
    noise_desc, _ = noise_desc_sug(noise_score)

    artifact_val = calc_artifact(gray, width, height)
    artifact_score = round(score_noise(artifact_val), 1)
    artifact_desc, _ = artifact_desc_sug(artifact_score)

    total_score = round(
        clarity_score * weights[0]
        + exposure_score * weights[1]
        + contrast_score * weights[2]
        + noise_score * weights[3]
        + artifact_score * weights[4],
        1,
    )
    level = level_from_score(total_score)
    issue = primary_issue([
        ("清晰度", clarity_score, clarity_desc),
        ("曝光", exposure_score, exposure_desc),
        ("对比度", contrast_score, contrast_desc),
        ("噪声", noise_score, noise_desc),
        ("压缩", artifact_score, artifact_desc),
    ])
    review_required = total_score < threshold or level in {"WARNING", "NG", "ERROR"}
    return {
        "score": total_score,
        "level": level,
        "issue": issue,
        "review_required": review_required,
        "advice": "识别结果建议人工复核" if review_required else "",
    }


def add_ocr_quality_to_report(report_path, records):
    try:
        from openpyxl import load_workbook
        from copy import copy
    except Exception:
        return

    workbook = load_workbook(report_path)
    sheet = workbook["OCR详细结果"] if "OCR详细结果" in workbook.sheetnames else workbook[workbook.sheetnames[1]]
    start_col = sheet.max_column + 1
    headers = ["图像质量评分", "图像质量等级", "人工复核建议"]

    def copy_cell_style(source, target):
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)

    for offset, header in enumerate(headers):
        cell = sheet.cell(row=1, column=start_col + offset, value=header)
        copy_cell_style(sheet.cell(row=1, column=start_col - 1), cell)
    for row_index, record in enumerate(records, start=2):
        values = [
            record.get("quality_score", ""),
            record.get("quality_level", ""),
            record.get("review_advice", ""),
        ]
        for offset, value in enumerate(values):
            cell = sheet.cell(row=row_index, column=start_col + offset, value=value)
            copy_cell_style(sheet.cell(row=row_index, column=start_col - 1), cell)

    widths = [14, 14, 24]
    for offset, width in enumerate(widths):
        col_letter = sheet.cell(row=1, column=start_col + offset).column_letter
        sheet.column_dimensions[col_letter].width = width
    workbook.save(report_path)


def cleanup_uploaded_batch(config):
    upload_dir = config.get("upload_dir")
    if not upload_dir:
        return False
    path = Path(upload_dir)
    if not path.exists():
        return False
    try:
        path.relative_to(BASE_DIR)
    except ValueError:
        return False
    shutil.rmtree(path, ignore_errors=True)
    return not path.exists()


def run_detection_job(job_id, config):
    # Large images must not multiply peak memory usage on a small web instance.
    append_log(job_id, "任务已排队，等待图片检测资源")
    with image_job_lock:
        _run_image_detection_job(job_id, config)


def validate_web_image_inputs(input_dir, log):
    directory = Path(input_dir)
    if not directory.is_dir():
        return
    for image_path in directory.iterdir():
        if not image_path.is_file() or image_path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        if image_path.stat().st_size > 30 * 1024 * 1024:
            raise ValueError(f"{image_path.name} 超过单张 30MB 的在线处理限制，请减小文件后重试。")
        try:
            with PILImage.open(image_path) as header:
                width, height = header.size
                baseline_jpeg = header.format == "JPEG" and header.mode in ("RGB", "L") and not (header.info.get("progressive") or header.info.get("progression"))
                max_pixels = 48_000_000 if baseline_jpeg else 24_000_000
        except PILImage.DecompressionBombError as exc:
            raise ValueError(f"{image_path.name} 声明的分辨率过大，请降低分辨率后重试。") from exc
        except (UnidentifiedImageError, OSError):
            # Preserve the detector's per-image ERROR result for unreadable inputs.
            continue
        log(f"图片输入：{image_path.name}，{width}×{height}，{image_path.stat().st_size / 1024 / 1024:.2f} MB")
        if width * height > max_pixels or max(width, height) > 16384:
            raise ValueError(
                f"{image_path.name} 分辨率 {width}×{height} 超过当前在线实例的处理限制"
                f"（此格式最多 {max_pixels // 1_000_000}MP，最长边 16384 像素）。"
                "请降低分辨率后重试；本任务未自动缩放图片。"
            )


def _run_image_detection_job(job_id, config):
    started_at = time.time()
    try:
        update_job(job_id, status="running", config=config, started_at_ts=started_at, elapsed_seconds=0, eta_seconds=None)
        append_log(job_id, "开始检测任务")
        append_log(job_id, f"报告目录：{config['report_dir']}")

        def progress(done, total):
            percent = round(done / total * 100) if total else 0
            elapsed = max(0, time.time() - started_at)
            eta = round(elapsed / done * (total - done)) if done and total else None
            update_job(job_id, done=done, total=total, percent=percent, elapsed_seconds=round(elapsed), eta_seconds=eta)

        web_config = dict(config)
        web_config["generate_report"] = False
        web_config["return_results"] = True
        validate_web_image_inputs(config["input_dir"], lambda msg: append_log(job_id, msg))
        detector = ImageQualityDetector(web_config, lambda msg: append_log(job_id, msg), progress)
        payload = detector.run()
        results = payload["results"]
        info = output_info(config, results)
        elapsed = round(time.time() - started_at, 1)
        update_job(job_id, status="done", percent=100, elapsed_seconds=elapsed, eta_seconds=0, results=results, result_rows=result_summary(results), output_info=info)
        append_log(job_id, f"检测完成，耗时 {elapsed} 秒，标注图片已保存 {info['saved_count']} 张，失败 {info['failed_count']} 张")
        append_log(job_id, "Excel 不会自动生成；需要时点击“下载 Excel 报告”")
    except Exception as exc:
        update_job(job_id, status="error", elapsed_seconds=round(time.time() - started_at, 1))
        append_log(job_id, f"检测失败：{exc}")
        app.logger.exception("Image detection job %s failed", job_id)


@app.route("/")
def index():
    return render_template(
        "index.html",
        defaults={
            "input_dir": str(DEFAULT_INPUT_DIR),
            "ok_dir": str(DEFAULT_OK_DIR),
            "ng_dir": str(DEFAULT_NG_DIR),
            "report_dir": str(DEFAULT_REPORT_DIR),
        },
        weight_modes=WEIGHT_MODES,
        detector_version=DETECTOR_VERSION,
        video_defaults={
            "input_dir": str(DEFAULT_VIDEO_INPUT_DIR),
            "report_dir": str(DEFAULT_VIDEO_REPORT_DIR),
            "quality_report_dir": str(DEFAULT_VIDEO_QUALITY_REPORT_DIR),
            "ffprobe_path": DEFAULT_FFPROBE_PATH,
        },
        video_modes={key: value["name"] for key, value in STANDARDS_CONFIG.items()},
    )


@app.post("/api/detect")
def detect():
    try:
        weight_mode, weights = parse_weights(request.form)
        upload_dir = save_uploads(request.files.getlist("images"))
        input_dir = upload_dir or Path(request.form.get("input_dir", str(DEFAULT_INPUT_DIR)).strip())

        config = {
            "input_dir": str(input_dir),
            "ok_dir": request.form.get("ok_dir", str(DEFAULT_OK_DIR)).strip(),
            "ng_dir": request.form.get("ng_dir", str(DEFAULT_NG_DIR)).strip(),
            "report_dir": request.form.get("report_dir", str(DEFAULT_REPORT_DIR)).strip(),
            "analysis_mode": request.form.get("analysis_mode", "整图检测"),
            "region_mode": request.form.get("region_mode", "九宫格全部"),
            "weight_mode": weight_mode,
            "weights": weights,
        }

        if upload_dir:
            batch_output = BASE_DIR / "web_results" / upload_dir.name
            config["ok_dir"] = str(batch_output / "OK_images")
            config["ng_dir"] = str(batch_output / "NG_images")
            config["report_dir"] = str(batch_output / "report")

        for key in ("input_dir", "ok_dir", "ng_dir", "report_dir"):
            if not config[key]:
                raise ValueError("目录路径不能为空。")

        job_id = make_job()
        thread = threading.Thread(target=run_detection_job, args=(job_id, config), daemon=True)
        thread.start()
        return jsonify({"ok": True, "job_id": job_id})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


def run_ocr_job(job_id, config):
    try:
        update_job(job_id, status="running", config=config, job_type="ocr", started_at_ts=time.time(), eta_seconds=None, elapsed_seconds=0)
        append_log(job_id, "开始 OCR 2.0 识别任务")
        append_log(job_id, f"输入路径：{config['source']}")
        append_log(job_id, f"识别模式：{config['mode']}")
        append_log(job_id, f"输出目录：{config['output_dir']}")

        module = load_ocr_module()
        preprocess_list = [s.strip() for s in config.get("preprocess", "").split(",") if s.strip()]
        preproc_kwargs = {
            "clahe_clip": config["clahe_clip"],
            "denoise_method": config["denoise_method"],
            "denoise_ksize": config["denoise_ksize"],
            "sharpen_strength": config["sharpen_strength"],
            "scale_factor": config["scale_factor"],
            "gamma": config["gamma"],
        }
        if preprocess_list:
            append_log(job_id, f"预处理：{', '.join(preprocess_list)}")
        if config.get("quality_gate_enabled", True):
            append_log(job_id, f"OCR 前置质量预检已开启，复核阈值：{config['quality_review_threshold']}")

        input_module = module.InputModule(config["source"], config["mode"], config["frame_interval"])
        ocr_engine = module.OCREngine(config["lang"], config["use_gpu"])
        analyzer = module.ResultAnalyzer(config["keywords"], module.normalize_confidence(config["min_char_confidence"]))
        judge = module.QualityJudge(
            config["min_confidence"],
            config["max_garbled"],
            module.normalize_confidence(config["min_char_confidence"]),
            config["max_low_conf_ratio"],
            config["min_ocr_count"],
            config["min_text_area"],
        )
        writer = module.ReportWriter(config["output_dir"])

        records = []
        total_hint = config.get("total_hint") or 0
        done = ok = ng = 0
        started_at = time.time()

        for name, img, orig_path in input_module.get_data():
            done += 1
            elapsed = max(0, time.time() - started_at)
            percent = round(done / total_hint * 100) if total_hint else 0
            eta = round(elapsed / done * (total_hint - done)) if total_hint and done < total_hint else 0 if total_hint else None
            update_job(job_id, done=done, total=total_hint, percent=percent, eta_seconds=eta, elapsed_seconds=round(elapsed))
            append_log(job_id, f"[{done}] 正在处理：{name}")

            if img is None:
                quality = assess_ocr_image_quality(None, config["quality_review_threshold"])
                record = {
                    "image_name": name,
                    "preprocess": "None",
                    "quality_score": quality["score"],
                    "quality_level": quality["level"],
                    "quality_issue": quality["issue"],
                    "review_required": quality["review_required"],
                    "review_advice": quality["advice"],
                    "full_text": "",
                    "ocr_count": 0,
                    "avg_confidence": 0,
                    "min_confidence": 0,
                    "max_confidence": 0,
                    "low_conf_count": 0,
                    "low_conf_ratio": 0,
                    "text_area": 0,
                    "keywords": config["keywords"],
                    "hit_keywords": [],
                    "miss_keywords": config["keywords"][:],
                    "hit_rate": 0,
                    "status": "NG",
                    "sub_codes": ["NG-07"],
                    "reasons": ["图片读取失败"],
                }
                records.append(record)
                ng += 1
                append_log(job_id, f"{name} 图片读取失败 (NG-07)")
                continue

            quality = assess_ocr_image_quality(img, config["quality_review_threshold"]) if config.get("quality_gate_enabled", True) else {
                "score": "",
                "level": "",
                "issue": "",
                "review_required": False,
                "advice": "",
            }
            if quality["review_required"]:
                append_log(job_id, f"{name} 图像质量 {quality['level']} / {quality['score']}，识别结果建议人工复核")

            processed_img, preproc_desc = module.ImagePreprocessor.process(img, preprocess_list, **preproc_kwargs)
            ocr_results = ocr_engine.recognize(processed_img)
            analysis = analyzer.analyze(ocr_results)
            judgement = judge.judge(analysis)
            if judgement["status"] == "OK":
                ok += 1
            else:
                ng += 1

            record = {
                "image_name": name,
                "preprocess": preproc_desc,
                "quality_score": quality["score"],
                "quality_level": quality["level"],
                "quality_issue": quality["issue"],
                "review_required": quality["review_required"],
                "review_advice": quality["advice"],
                "full_text": analysis["full_text"],
                "ocr_count": analysis["ocr_count"],
                "avg_confidence": analysis["avg_confidence"],
                "min_confidence": analysis["min_confidence"],
                "max_confidence": analysis["max_confidence"],
                "low_conf_count": analysis["low_conf_count"],
                "low_conf_ratio": analysis["low_conf_ratio"],
                "text_area": analysis["text_area"],
                "keywords": config["keywords"],
                "hit_keywords": analysis["hit_keywords"],
                "miss_keywords": analysis["miss_keywords"],
                "hit_rate": analysis["hit_rate"],
                "status": judgement["status"],
                "sub_codes": judgement["sub_codes"],
                "reasons": judgement["reasons"],
            }
            records.append(record)
            writer.save_text(name, analysis["full_text"])
            writer.save_annotated_image(name, processed_img, ocr_results, judgement["status"], orig_path)
            advice = f"，{quality['advice']}" if quality["review_required"] else ""
            quality_text = f"，质量 {quality['level']} {quality['score']}" if quality["level"] else ""
            append_log(job_id, f"{name} => {judgement['status']}，文字块 {analysis['ocr_count']}，均置信度 {analysis['avg_confidence']:.3f}{quality_text}{advice}")

        if not records:
            raise RuntimeError("没有找到可处理的 OCR 文件")

        report_path = writer.generate_report(records)
        add_ocr_quality_to_report(report_path, records)
        output_info = {
            "report_dir": str(writer.report_dir),
            "text_dir": str(writer.text_dir),
            "annotated_dir": str(writer.annotated_dir),
            "ng_dir": str(writer.ng_dir),
            "ok_count": ok,
            "ng_count": ng,
            "saved_count": len(records),
            "failed_count": ng,
        }
        update_job(
            job_id,
            status="done",
            percent=100,
            eta_seconds=0,
            elapsed_seconds=round(time.time() - started_at),
            report_path=str(report_path),
            results=records,
            result_rows=ocr_result_summary(records),
            output_info=output_info,
        )
        append_log(job_id, f"OCR 2.0 完成：总数 {len(records)}，OK {ok}，NG {ng}")
        append_log(job_id, f"Excel 报告已生成：{report_path}")
    except Exception as exc:
        update_job(job_id, status="error", job_type="ocr")
        append_log(job_id, f"OCR 2.0 任务失败：{exc}")


@app.post("/api/ocr/detect")
def ocr_detect():
    try:
        upload_dir = save_ocr_uploads(request.files.getlist("ocr_files"))
        source = upload_dir or Path(request.form.get("ocr_input_dir", str(BASE_DIR / "OCR测试")).strip())
        source = Path(source)
        if not source.exists():
            raise ValueError(f"OCR 输入路径不存在：{source}")

        requested_mode = request.form.get("ocr_mode", "auto")
        if requested_mode == "auto":
            suffix = source.suffix.lower()
            mode = "batch" if source.is_dir() else "video" if suffix in OCR_VIDEO_EXTENSIONS else "single"
        else:
            mode = requested_mode

        output_dir = request.form.get("ocr_output_dir", "").strip()
        if upload_dir:
            output_dir = str(OCR_OUTPUT_DIR / upload_dir.name)
        elif not output_dir:
            output_dir = str((source if source.is_dir() else source.parent) / "ocr_output")

        keywords = [kw.strip() for kw in request.form.get("ocr_keywords", "").replace("，", ",").split(",") if kw.strip()]
        total_hint = 0
        if mode == "batch" and source.is_dir():
            total_hint = sum(1 for path in source.rglob("*") if path.suffix.lower() in OCR_IMAGE_EXTENSIONS)
        elif mode == "single":
            total_hint = 1

        config = {
            "source": str(source),
            "mode": mode,
            "output_dir": output_dir,
            "frame_interval": max(1, int(request.form.get("ocr_frame_interval", "30") or 30)),
            "lang": request.form.get("ocr_lang", "ch").strip() or "ch",
            "use_gpu": request.form.get("ocr_use_gpu") == "on",
            "keywords": keywords,
            "min_confidence": float(request.form.get("ocr_min_confidence", "0.8") or 0.8),
            "max_garbled": float(request.form.get("ocr_max_garbled", "0.2") or 0.2),
            "min_char_confidence": float(request.form.get("ocr_min_char_confidence", "0.5") or 0.5),
            "max_low_conf_ratio": float(request.form.get("ocr_max_low_conf_ratio", "0.3") or 0.3),
            "min_ocr_count": int(request.form.get("ocr_min_ocr_count", "1") or 1),
            "min_text_area": float(request.form.get("ocr_min_text_area", "0") or 0),
            "quality_gate_enabled": "on" in request.form.getlist("ocr_quality_gate"),
            "quality_review_threshold": float(request.form.get("ocr_quality_review_threshold", "70") or 70),
            "preprocess": request.form.get("ocr_preprocess", "").strip(),
            "clahe_clip": float(request.form.get("ocr_clahe_clip", "2.0") or 2.0),
            "denoise_method": request.form.get("ocr_denoise_method", "gaussian").strip() or "gaussian",
            "denoise_ksize": int(request.form.get("ocr_denoise_ksize", "5") or 5),
            "sharpen_strength": float(request.form.get("ocr_sharpen_strength", "1.0") or 1.0),
            "scale_factor": float(request.form.get("ocr_scale_factor", "2.0") or 2.0),
            "gamma": float(request.form.get("ocr_gamma", "1.2") or 1.2),
            "upload_dir": str(upload_dir) if upload_dir else "",
            "total_hint": total_hint,
        }

        job_id = make_job()
        thread = threading.Thread(target=run_ocr_job, args=(job_id, config), daemon=True)
        thread.start()
        return jsonify({"ok": True, "job_id": job_id})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


def video_result_summary(results):
    rows = []
    for record in results:
        rows.append({
            "filename": record.get("文件名", ""),
            "resolution": record.get("分辨率", ""),
            "total_bitrate": record.get("总码率", ""),
            "video_bitrate": record.get("视频码率", ""),
            "audio_bitrate": record.get("音频码率", ""),
            "fps": record.get("帧率") or "",
            "frames": record.get("视频总帧数") or "",
            "threshold": record.get("阈值Mbps") or "",
            "judgement": record.get("码率判断", ""),
            "white_balance_score": record.get("白平衡评分", ""),
            "white_balance": record.get("视频白平衡判断", ""),
            "white_balance_abnormal": record.get("白平衡异常帧占比", ""),
            "error": record.get("错误信息", ""),
        })
    return rows


def video_quality_result_summary(results):
    rows = []
    for record in results:
        rows.append({
            "filename": record.get("文件名", ""),
            "resolution": record.get("分辨率", ""),
            "score": record.get("综合评分", 0),
            "level": record.get("综合等级", ""),
            "bitrate": record.get("码率判断", ""),
            "sharpness": record.get("清晰度判断", ""),
            "exposure": record.get("曝光判断", ""),
            "stability": record.get("稳定性判断", ""),
            "drops": record.get("掉帧判断", ""),
            "black_white": record.get("黑白屏判断", ""),
            "frozen": record.get("冻结判断", ""),
            "audio": record.get("音频判断", ""),
            "error": record.get("错误信息", ""),
        })
    return rows


def enhance_result_summary(rows):
    summary_rows = []
    for row in rows:
        summary_rows.append({
            "filename": row.get("文件名", ""),
            "before_score": row.get("原图综合分", 0),
            "after_score": row.get("增强后综合分", 0),
            "delta": row.get("分数变化", 0),
            "effect": row.get("增强效果判断", ""),
            "before_level": row.get("原图质量等级", ""),
            "after_level": row.get("增强后质量等级", ""),
            "before_issue": row.get("原图问题原因", ""),
            "steps": row.get("执行增强方法", ""),
            "after_issue": row.get("增强后问题原因", ""),
            "advice": row.get("处理建议", ""),
            "enhanced_path": row.get("增强图路径", ""),
            "compare_path": row.get("对比图路径", ""),
            "before_clarity": row.get("原图清晰度分", ""),
            "after_clarity": row.get("增强后清晰度分", ""),
            "before_exposure": row.get("原图曝光分", ""),
            "after_exposure": row.get("增强后曝光分", ""),
            "before_contrast": row.get("原图对比度分", ""),
            "after_contrast": row.get("增强后对比度分", ""),
            "before_noise": row.get("原图噪声分", ""),
            "after_noise": row.get("增强后噪声分", ""),
            "before_artifact": row.get("原图压缩伪影分", ""),
            "after_artifact": row.get("增强后压缩伪影分", ""),
        })
    return summary_rows


def build_enhance_args(form):
    def get_float(name, default):
        try:
            return float(form.get(name, default))
        except (TypeError, ValueError):
            return default

    def get_int(name, default):
        try:
            return int(float(form.get(name, default)))
        except (TypeError, ValueError):
            return default

    return SimpleNamespace(
        steps=form.get("enhance_steps", ""),
        enhance_strength=form.get("enhance_strength", "standard"),
        gamma=get_float("enhance_gamma", 0.85),
        darken_alpha=get_float("enhance_darken_alpha", 0.96),
        darken_beta=get_int("enhance_darken_beta", -10),
        highlight_threshold=get_int("enhance_highlight_threshold", 205),
        highlight_strength=get_float("enhance_highlight_strength", 0.35),
        clahe_clip=get_float("enhance_clahe_clip", 1.8),
        clahe_grid=get_int("enhance_clahe_grid", 8),
        denoise_h=get_float("enhance_denoise_h", 6.0),
        denoise_h_color=get_float("enhance_denoise_h_color", 6.0),
        median_ksize=get_int("enhance_median_ksize", 3),
        bilateral_diameter=get_int("enhance_bilateral_diameter", 5),
        bilateral_sigma_color=get_float("enhance_bilateral_sigma_color", 35.0),
        bilateral_sigma_space=get_float("enhance_bilateral_sigma_space", 35.0),
        sharpen_amount=get_float("enhance_sharpen_amount", 0.25),
        sharpen_sigma=get_float("enhance_sharpen_sigma", 1.0),
        ocr_clip=get_float("enhance_ocr_clip", 2.0),
        ocr_grid=get_int("enhance_ocr_grid", 8),
        ocr_denoise_h=get_float("enhance_ocr_denoise_h", 6.0),
    )


def run_enhance_job(job_id, config):
    try:
        update_job(job_id, status="running", config=config, job_type="enhance", started_at_ts=time.time(), eta_seconds=None, elapsed_seconds=0)
        append_log(job_id, "开始规则版图像增强任务")
        append_log(job_id, f"输入：{config['input_path']}")
        append_log(job_id, f"输出：{config['output_dir']}")
        append_log(job_id, f"模式：{config['mode']}")

        started = time.time()
        result = run_enhancement(
            input_path=config["input_path"],
            output_dir=config["output_dir"],
            mode=config["mode"],
            args=config["args"],
            ocr_mode=config["ocr_mode"],
        )
        rows = result.get("rows") or []
        elapsed = round(time.time() - started)
        update_job(
            job_id,
            status="done",
            done=result.get("success", len(rows)),
            total=result.get("total", len(rows)),
            percent=100,
            eta_seconds=0,
            elapsed_seconds=elapsed,
            results=rows,
            result_rows=enhance_result_summary(rows),
            report_path=result.get("excel_report", ""),
            csv_report_path=result.get("csv_report", ""),
            output_info={
                "enhanced_dir": result.get("enhanced_dir", ""),
                "compare_dir": result.get("compare_dir", ""),
                "report_dir": str(Path(result.get("excel_report", "")).parent) if result.get("excel_report") else "",
                "log_path": result.get("log_path", ""),
                "saved_count": result.get("success", len(rows)),
                "failed_count": result.get("failed", 0),
            },
        )
        append_log(job_id, f"图像增强完成：成功 {result.get('success', len(rows))} 张，失败 {result.get('failed', 0)} 张")
        append_log(job_id, f"Excel 报告：{result.get('excel_report', '')}")
    except Exception as exc:
        update_job(job_id, status="error", job_type="enhance")
        append_log(job_id, f"图像增强失败：{exc}")


def run_video_detection_job(job_id, config):
    try:
        update_job(job_id, status="running", config=config, job_type="video", started_at_ts=time.time(), eta_seconds=None, elapsed_seconds=0)
        append_log(job_id, "开始视频码率检测任务")
        append_log(job_id, f"ffprobe：{config['ffprobe_path']}")
        append_log(job_id, f"视频目录：{config['input_dir']}")
        append_log(job_id, f"报告目录：{config['report_dir']}")

        def progress(done, total):
            percent = round(done / total * 100) if total else 0
            with jobs_lock:
                started_at = jobs[job_id].get("started_at_ts")
            elapsed = max(0, time.time() - started_at) if started_at else 0
            eta = round(elapsed / done * (total - done)) if done and total and done < total else 0 if done == total and total else None
            update_job(job_id, done=done, total=total, percent=percent, eta_seconds=eta, elapsed_seconds=round(elapsed))

        detector = VideoBitrateDetector(
            input_dir=config["input_dir"],
            report_dir=config["report_dir"],
            mode=int(config["mode"]),
            ffprobe_path=config["ffprobe_path"],
            include_frames=config["include_frames"],
            log_callback=lambda msg: append_log(job_id, msg),
            progress_callback=progress,
        )
        payload = detector.run()
        update_job(
            job_id,
            status="done",
            percent=100,
            eta_seconds=0,
            results=payload["results"],
            frame_data=payload["frame_data"],
            result_rows=video_result_summary(payload["results"]),
            output_info={
                "report_dir": config["report_dir"],
                "mode_name": payload["mode_config"]["name"],
                "saved_count": len(payload["results"]),
                "failed_count": sum(1 for row in payload["results"] if row.get("码率判断") == "ERROR"),
            },
        )
        if cleanup_uploaded_batch(config):
            append_log(job_id, "已自动清理上传视频副本，保留检测报告和结果文件")
        append_log(job_id, "视频检测完成；需要 Excel 时点击下载报告")
    except Exception as exc:
        update_job(job_id, status="error", job_type="video")
        append_log(job_id, f"视频检测失败：{exc}")


def run_video_quality_job(job_id, config):
    try:
        update_job(job_id, status="running", config=config, job_type="video_quality", started_at_ts=time.time(), eta_seconds=None, elapsed_seconds=0)
        append_log(job_id, "开始视频综合质量评分任务")
        append_log(job_id, f"ffprobe：{config['ffprobe_path']}")
        append_log(job_id, f"视频目录：{config['input_dir']}")
        append_log(job_id, f"报告目录：{config['report_dir']}")

        def progress(done, total):
            percent = round(done / total * 100) if total else 0
            with jobs_lock:
                started_at = jobs[job_id].get("started_at_ts")
            elapsed = max(0, time.time() - started_at) if started_at else 0
            eta = round(elapsed / done * (total - done)) if done and total and done < total else 0 if done == total and total else None
            update_job(job_id, done=done, total=total, percent=percent, eta_seconds=eta, elapsed_seconds=round(elapsed))

        detector = VideoQualityScoringDetector(
            input_dir=config["input_dir"],
            report_dir=config["report_dir"],
            mode=int(config["mode"]),
            ffprobe_path=config["ffprobe_path"],
            log_callback=lambda msg: append_log(job_id, msg),
            progress_callback=progress,
        )
        payload = detector.run()
        update_job(
            job_id,
            status="done",
            percent=100,
            eta_seconds=0,
            results=payload["results"],
            quality_payload=payload,
            result_rows=video_quality_result_summary(payload["results"]),
            output_info={
                "report_dir": config["report_dir"],
                "mode_name": payload["mode_config"]["name"],
                "saved_count": len(payload["results"]),
                "failed_count": sum(1 for row in payload["results"] if row.get("综合等级") == "ERROR"),
            },
        )
        if cleanup_uploaded_batch(config):
            append_log(job_id, "已自动清理上传视频副本，保留综合报告、预览图和异常截图")
        append_log(job_id, "视频综合质量评分完成；需要 Excel 时点击下载报告")
    except Exception as exc:
        update_job(job_id, status="error", job_type="video_quality")
        append_log(job_id, f"视频综合质量评分失败：{exc}")


@app.post("/api/video/detect")
def video_detect():
    try:
        upload_dir = save_video_uploads(request.files.getlist("videos"))
        input_dir = upload_dir or Path(request.form.get("video_input_dir", str(DEFAULT_VIDEO_INPUT_DIR)).strip())
        config = {
            "input_dir": str(input_dir),
            "report_dir": request.form.get("video_report_dir", str(DEFAULT_VIDEO_REPORT_DIR)).strip(),
            "ffprobe_path": request.form.get("ffprobe_path", DEFAULT_FFPROBE_PATH).strip(),
            "mode": int(request.form.get("video_mode", "1")),
            "include_frames": request.form.get("include_frames") == "on",
        }
        if upload_dir:
            config["report_dir"] = str(BASE_DIR / "web_video_results" / upload_dir.name / "report")
            config["upload_dir"] = str(upload_dir)
        for key in ("input_dir", "report_dir", "ffprobe_path"):
            if not config[key]:
                raise ValueError("视频目录、报告目录和 ffprobe 路径不能为空。")
        job_id = make_job()
        thread = threading.Thread(target=run_video_detection_job, args=(job_id, config), daemon=True)
        thread.start()
        return jsonify({"ok": True, "job_id": job_id})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/video-quality/detect")
def video_quality_detect():
    try:
        upload_dir = save_video_uploads(request.files.getlist("quality_videos"))
        input_dir = upload_dir or Path(request.form.get("quality_video_input_dir", str(DEFAULT_VIDEO_INPUT_DIR)).strip())
        config = {
            "input_dir": str(input_dir),
            "report_dir": request.form.get("quality_video_report_dir", str(DEFAULT_VIDEO_QUALITY_REPORT_DIR)).strip(),
            "ffprobe_path": request.form.get("quality_ffprobe_path", DEFAULT_FFPROBE_PATH).strip(),
            "mode": int(request.form.get("quality_video_mode", "1")),
        }
        if upload_dir:
            config["report_dir"] = str(BASE_DIR / "web_video_results" / upload_dir.name / "quality_report")
            config["upload_dir"] = str(upload_dir)
        for key in ("input_dir", "report_dir", "ffprobe_path"):
            if not config[key]:
                raise ValueError("视频目录、报告目录和 ffprobe 路径不能为空。")
        job_id = make_job()
        thread = threading.Thread(target=run_video_quality_job, args=(job_id, config), daemon=True)
        thread.start()
        return jsonify({"ok": True, "job_id": job_id})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/enhance/detect")
def enhance_detect():
    try:
        upload_dir = save_enhance_uploads(request.files.getlist("enhance_images"))
        input_path = upload_dir or Path(request.form.get("enhance_input_dir", str(DEFAULT_INPUT_DIR)).strip())
        output_dir = Path(request.form.get("enhance_output_dir", str(ENHANCE_OUTPUT_DIR)).strip())
        mode = request.form.get("enhance_mode", "rule")
        if mode not in {"manual", "fixed", "rule"}:
            raise ValueError("图像增强模式只能是 manual / fixed / rule")
        if upload_dir:
            output_dir = ENHANCE_OUTPUT_DIR / upload_dir.name
        if not str(input_path).strip() or not str(output_dir).strip():
            raise ValueError("图像增强输入路径和输出目录不能为空。")
        config = {
            "input_path": str(input_path),
            "output_dir": str(output_dir),
            "mode": mode,
            "ocr_mode": request.form.get("enhance_ocr_mode") == "on",
            "args": build_enhance_args(request.form),
            "upload_dir": str(upload_dir) if upload_dir else "",
        }
        job_id = make_job()
        thread = threading.Thread(target=run_enhance_job, args=(job_id, config), daemon=True)
        thread.start()
        return jsonify({"ok": True, "job_id": job_id})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.get("/api/jobs/<job_id>")
def job_status(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    public_job = dict(job)
    if job.get("status") == "running" and job.get("started_at_ts"):
        public_job["elapsed_seconds"] = round(max(0, time.time() - job["started_at_ts"]))
    public_job.pop("results", None)
    public_job.pop("quality_payload", None)
    public_job.pop("config", None)
    return jsonify({"ok": True, "job": json_safe(public_job)})


@app.get("/api/jobs/<job_id>/csv")
def download_job_csv(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    rows = json_safe(job.get("result_rows") or [])
    if not rows:
        return jsonify({"ok": False, "error": "没有可导出的结果"}), 400

    fieldnames = []
    for row in rows:
        for key in row.keys():
            if key not in fieldnames:
                fieldnames.append(key)
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    filename = f"{job.get('job_type') or 'image'}_results_{job_id[:8]}.csv"
    return Response(
        "\ufeff" + buffer.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/files")
def serve_local_file():
    raw_path = request.args.get("path", "")
    if not raw_path:
        return jsonify({"ok": False, "error": "文件路径不能为空"}), 400
    try:
        file_path = Path(raw_path).resolve()
        file_path.relative_to(BASE_DIR)
    except Exception:
        return jsonify({"ok": False, "error": "只能预览项目目录内的文件"}), 403
    if not file_path.exists() or not file_path.is_file():
        return jsonify({"ok": False, "error": "文件不存在"}), 404
    return send_file(file_path)


@app.get("/api/video/jobs/<job_id>/report")
def download_video_report(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    config = dict(job.get("config") or {})
    if not job.get("results") or not config:
        return jsonify({"ok": False, "error": "没有可生成报告的视频结果"}), 400

    detector = VideoBitrateDetector(
        input_dir=config["input_dir"],
        report_dir=config["report_dir"],
        mode=int(config["mode"]),
        ffprobe_path=config["ffprobe_path"],
        include_frames=config["include_frames"],
    )
    try:
        report_path = detector.write_report(job["results"], job.get("frame_data") or {})
    except OSError as exc:
        append_log(job_id, f"指定视频报告目录不可写，已改存默认目录：{DEFAULT_VIDEO_REPORT_DIR}，原因：{exc}")
        detector.report_dir = DEFAULT_VIDEO_REPORT_DIR
        report_path = detector.write_report(job["results"], job.get("frame_data") or {})
    update_job(job_id, report_path=str(report_path))
    append_log(job_id, f"视频 Excel 报告已生成：{report_path}")
    return send_file(report_path, as_attachment=True, download_name=Path(report_path).name)


@app.get("/api/video-quality/jobs/<job_id>/report")
def download_video_quality_report(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    config = dict(job.get("config") or {})
    payload = job.get("quality_payload")
    if not payload or not config:
        return jsonify({"ok": False, "error": "没有可生成报告的视频综合质量结果"}), 400

    detector = VideoQualityScoringDetector(
        input_dir=config["input_dir"],
        report_dir=config["report_dir"],
        mode=int(config["mode"]),
        ffprobe_path=config["ffprobe_path"],
    )
    try:
        report_path = detector.write_report(payload)
    except OSError as exc:
        append_log(job_id, f"指定视频综合报告目录不可写，已改存默认目录：{DEFAULT_VIDEO_QUALITY_REPORT_DIR}，原因：{exc}")
        detector.report_dir = DEFAULT_VIDEO_QUALITY_REPORT_DIR
        report_path = detector.write_report(payload)
    update_job(job_id, report_path=str(report_path))
    append_log(job_id, f"视频综合 Excel 报告已生成：{report_path}")
    return send_file(report_path, as_attachment=True, download_name=Path(report_path).name)


@app.get("/api/jobs/<job_id>/report")
def download_report(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    if not job.get("results"):
        return jsonify({"ok": False, "error": "没有可生成报告的检测结果"}), 400

    config = dict(job.get("config") or {})
    if not config:
        return jsonify({"ok": False, "error": "任务配置缺失，无法生成报告"}), 400

    _, sub_label = region_names_for(config.get("region_mode", "九宫格全部"))
    detector = ImageQualityDetector(config)
    try:
        report_path = detector._write_report(job["results"], Path(config["report_dir"]), sub_label)
    except OSError as exc:
        fallback_dir = DEFAULT_REPORT_DIR
        append_log(job_id, f"指定报告目录不可写，已改存默认目录：{fallback_dir}，原因：{exc}")
        report_path = detector._write_report(job["results"], fallback_dir, sub_label)
    update_job(job_id, report_path=str(report_path))
    append_log(job_id, f"Excel 报告已按需生成：{report_path}")
    return send_file(report_path, as_attachment=True, download_name=Path(report_path).name)


@app.get("/api/ocr/jobs/<job_id>/report")
def download_ocr_report(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    report_path = Path(job.get("report_path") or "")
    if not report_path.exists():
        return jsonify({"ok": False, "error": "OCR 报告文件不存在"}), 404
    return send_file(report_path, as_attachment=True, download_name=report_path.name)


@app.get("/api/enhance/jobs/<job_id>/report")
def download_enhance_report(job_id):
    job = snapshot_job(job_id)
    if not job:
        return jsonify({"ok": False, "error": "任务不存在"}), 404
    if job.get("status") != "done":
        return jsonify({"ok": False, "error": "任务尚未完成"}), 400
    report_path = Path(job.get("report_path") or "")
    if not report_path.exists():
        return jsonify({"ok": False, "error": "图像增强报告文件不存在"}), 404
    return send_file(report_path, as_attachment=True, download_name=report_path.name)


@app.post("/api/clear-uploads")
def clear_uploads():
    cleared = []
    for directory in (UPLOAD_DIR, BASE_DIR / "web_results", VIDEO_UPLOAD_DIR, BASE_DIR / "web_video_results", OCR_UPLOAD_DIR, OCR_OUTPUT_DIR, ENHANCE_UPLOAD_DIR, ENHANCE_OUTPUT_DIR):
        if directory.exists():
            shutil.rmtree(directory, ignore_errors=True)
        directory.mkdir(parents=True, exist_ok=True)
        cleared.append(str(directory))
    return jsonify({"ok": True, "cleared": cleared})


if __name__ == "__main__":
    for directory in (DEFAULT_INPUT_DIR, DEFAULT_OK_DIR, DEFAULT_NG_DIR, DEFAULT_REPORT_DIR, DEFAULT_VIDEO_INPUT_DIR, DEFAULT_VIDEO_REPORT_DIR, DEFAULT_VIDEO_QUALITY_REPORT_DIR, UPLOAD_DIR, VIDEO_UPLOAD_DIR, OCR_UPLOAD_DIR, OCR_OUTPUT_DIR, ENHANCE_UPLOAD_DIR, ENHANCE_OUTPUT_DIR):
        directory.mkdir(parents=True, exist_ok=True)
    # 托管平台通常通过 PORT 注入监听端口；本地运行仍默认使用 5000。
    port = int(os.environ.get("PORT") or os.environ.get("WEB_PORT", "5000"))
    host = os.environ.get("WEB_HOST", "0.0.0.0")
    debug = os.environ.get("WEB_DEBUG", "0").lower() in {"1", "true", "yes"}
    app.run(host=host, port=port, debug=debug, use_reloader=False)
