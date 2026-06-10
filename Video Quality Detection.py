#!/usr/bin/env python3
"""
视频综合质量评分工具（100分制）+ 图表 + 异常截图 + 预览图
功能：码率检测、清晰度、曝光、稳定性、掉帧、黑屏/白屏、画面冻结、音频基础
输出 Excel 包含独立模块报告、评分标准、综合评分、统计图表、预览图
依赖：ffprobe, ffmpeg, opencv-python, numpy, openpyxl
"""

import os, sys, json, subprocess, re, cv2, numpy as np
from datetime import datetime, timedelta
from typing import Dict, List, Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.drawing.image import Image as XLImage
except ImportError:
    print("请安装 openpyxl: pip install openpyxl")
    sys.exit(1)

# ==================== 配置区域 ====================
FFPROBE_PATH = r"D:\码率工具FFepeg\ffmpeg-8.1.1-full_build\bin\ffprobe.exe"   # 你的 ffprobe 路径
FFMPEG_PATH = os.path.join(os.path.dirname(FFPROBE_PATH), "ffmpeg.exe")
VIDEO_FOLDER = r"D:\Codex\码率测试\input"                                      # 视频文件夹
OUTPUT_NAME = "video_scoring_report"
VIDEO_EXTS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".mpg", ".mpeg"}

# 检测阈值
SHARPNESS_THRESH = 80
EXP_UNDER_SEVERE = 30; EXP_UNDER = 60; EXP_OVER = 191; EXP_OVER_SEVERE = 225
BLACK_PIX_RATIO = 0.9; SUS_BLACK_RATIO = 0.8; WHITE_PIX_RATIO = 0.9; SUS_WHITE_RATIO = 0.8
FREEZE_LOW_DIFF = 2; STABILITY_OK_THRESH = 5
SILENCE_DBFS = -60; LOW_VOL_DBFS = -35; HIGH_VOL_DBFS = -3

# 截图和预览配置
ABNORMAL_SAVE_DIR = "异常截图"
PREVIEW_ROWS, PREVIEW_COLS = 3, 3
PREVIEW_SIZE = (320, 180)
# =================================================

# ---------- 完整模式定义 ----------
STANDARDS_CONFIG = {
    1: {
        "name": "普通手机拍摄模式",
        "standards": [
            (921600, 30, 4.0, None),
            (2073600, 30, 8.0, 12.0),
            (8294400, 30, 35.0, 50.0),
            (float("inf"), 30, 100.0, 150.0),
        ],
        "table": [
            ["720P 30fps", "1280×720", "30", "4 - 6 Mbps", "低于 4 Mbps 偏低"],
            ["1080P 30fps", "1920×1080", "30", "8 - 12 Mbps", "低于 8 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 30fps", "3840×2160", "30", "35 - 50 Mbps", "低于 35 Mbps 偏低"],
            ["4K 60fps", "3840×2160", "60", "50 - 80 Mbps", "低于 50 Mbps 偏低"],
        ]
    },
    2: {
        "name": "短视频平台拍摄/上传模式",
        "standards": [
            (921600, 30, 3.0, None),
            (2073600, 30, 6.0, 10.0),
            (3686400, 30, 12.0, 16.0),
            (8294400, 30, 30.0, 45.0),
            (float("inf"), 30, 30.0, 45.0),
        ],
        "table": [
            ["720P 竖屏", "720×1280", "30fps", "3 - 6 Mbps", "低于 3 Mbps 偏低"],
            ["1080P 竖屏", "1080×1920", "30fps", "6 - 10 Mbps", "低于 6 Mbps 偏低"],
            ["1080P 竖屏高帧率", "1080×1920", "60fps", "10 - 16 Mbps", "低于 10 Mbps 偏低"],
            ["2K 竖屏", "1440×2560", "30fps", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 竖屏", "2160×3840", "30fps", "30 - 45 Mbps", "低于 30 Mbps 偏低"],
        ]
    },
    3: {
        "name": "监控/低码率存储模式",
        "standards": [
            (921600, 30, 1.0, None),
            (2073600, 30, 2.0, None),
            (3686400, 30, 4.0, None),
            (8294400, 30, 8.0, None),
            (float("inf"), 30, 8.0, None),
        ],
        "table": [
            ["720P 监控", "1280×720", "15/25fps", "1 - 3 Mbps", "低于 1 Mbps 偏低"],
            ["1080P 监控", "1920×1080", "15/25fps", "2 - 6 Mbps", "低于 2 Mbps 偏低"],
            ["2K 监控", "2560×1440", "15/25fps", "4 - 10 Mbps", "低于 4 Mbps 偏低"],
            ["4K 监控", "3840×2160", "15/25fps", "8 - 20 Mbps", "低于 8 Mbps 偏低"],
        ]
    },
    4: {
        "name": "高质量拍摄/相机素材模式",
        "standards": [
            (2073600, 30, 20.0, 30.0),
            (8294400, 30, 60.0, 100.0),
            (18000000, 30, 200.0, 200.0),
            (float("inf"), 30, 200.0, 200.0),
        ],
        "table": [
            ["1080P 30fps", "1920×1080", "30", "20 - 50 Mbps", "低于 20 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "30 - 80 Mbps", "低于 30 Mbps 偏低"],
            ["4K 24/25/30fps", "3840×2160", "24/25/30", "60 - 100 Mbps", "低于 60 Mbps 偏低"],
            ["4K 50/60fps", "3840×2160", "50/60", "100 - 200 Mbps", "低于 100 Mbps 偏低"],
            ["4K 10bit/LOG", "3840×2160", "24/25/30", "100 - 400 Mbps", "视编码而定"],
            ["6K/8K 高质量", "6K/8K", "24/30/60", "200 Mbps 以上", "低于 200 Mbps 偏低"],
        ]
    }
}

# ---------- 辅助函数 ----------
def select_mode():
    print("请选择拍摄模式：\n1 - 普通手机拍摄模式\n2 - 短视频平台模式\n3 - 监控/低码率存储模式\n4 - 高质量拍摄/相机素材模式")
    while True:
        try:
            m = int(input("请输入数字 (1-4): ").strip())
            if m in STANDARDS_CONFIG: return m
        except: pass
        print("输入无效")

def get_video_duration(path):
    cmd = [FFPROBE_PATH, "-v", "quiet", "-print_format", "json", "-show_format", path]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        dur = json.loads(res.stdout).get("format", {}).get("duration")
        return float(dur) if dur else None
    except: return None

def smart_interval(dur):
    if dur <= 10: return 1.0
    elif dur <= 60: return 2.0
    elif dur <= 300: return 5.0
    else: return 10.0

def format_time(s): return str(timedelta(seconds=s))[:10]
def format_bps(b):
    if b is None: return "N/A"
    if b >= 1e6: return f"{b/1e6:.2f} Mbps"
    if b >= 1e3: return f"{b/1e3:.2f} kbps"
    return f"{b} bps"

def run_ffprobe(args):
    cmd = [FFPROBE_PATH, "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams"] + args
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return json.loads(res.stdout)
    except subprocess.CalledProcessError as e:
        print(f"ffprobe 执行失败: {e.stderr}", file=sys.stderr)
        return {}
    except FileNotFoundError:
        print(f"错误: 未找到 ffprobe，请检查路径: {FFPROBE_PATH}", file=sys.stderr)
        sys.exit(1)

def extract_video_info(path):
    info = run_ffprobe([path])
    if not info: return {}
    fmt = info.get("format", {})
    total_br = int(fmt["bit_rate"]) if fmt.get("bit_rate") else None
    v_br = a_br = None; w = h = None; fps = None; total_f = None; audio = {}
    for s in info.get("streams", []):
        if s["codec_type"] == "video":
            if w is None:
                w = s.get("width"); h = s.get("height")
            if v_br is None:
                if "bit_rate" in s: v_br = int(s["bit_rate"])
                else:
                    dur = float(s.get("duration", 0)); sz = s.get("size")
                    if dur>0 and sz: v_br = int(float(sz)*8/dur)
            if fps is None:
                r = s.get("r_frame_rate")
                if r:
                    try:
                        if '/' in r: num, den = r.split('/'); fps = float(num)/float(den)
                        else: fps = float(r)
                    except: pass
            if total_f is None:
                nb = s.get("nb_frames")
                if nb is not None: total_f = int(nb)
                else:
                    dur = float(s.get("duration",0))
                    if dur>0 and fps: total_f = int(dur*fps)
        elif s["codec_type"] == "audio" and not audio:
            audio = {
                "codec": s.get("codec_name","N/A"),
                "sample_rate": s.get("sample_rate","N/A"),
                "channels": s.get("channels","N/A"),
                "bitrate": int(s.get("bit_rate",0)) if s.get("bit_rate") else None
            }
            a_br = audio["bitrate"]
    pix = w*h if w and h else None
    return {
        "total_bitrate": total_br, "video_bitrate": v_br, "audio_bitrate": a_br,
        "width": w, "height": h, "pixel_count": pix, "fps": fps, "total_frames": total_f,
        "audio_info": audio if audio else None,
    }

# ---------- 图片写入支持中文路径 ----------
def imwrite_chinese(path, img):
    """解决 OpenCV 不支持中文路径的问题"""
    try:
        ext = os.path.splitext(path)[1] or '.jpg'
        success, buf = cv2.imencode(ext, img)
        if not success:
            return False
        with open(path, 'wb') as f:
            f.write(buf.tobytes())
        return True
    except Exception as e:
        print(f"    [写入异常] {e}")
        return False

def save_abnormal_frame(frame, video_name, timestamp, anomaly_type, save_root):
    try:
        video_safe = re.sub(r'[\\/*?:\[\]]', '_', video_name)
        anomaly_dir = os.path.join(save_root, video_safe, anomaly_type)
        os.makedirs(anomaly_dir, exist_ok=True)
        time_str = format_time(timestamp).replace(":", "-")
        filename = f"{time_str}.jpg"
        filepath = os.path.join(anomaly_dir, filename)
        if imwrite_chinese(filepath, frame):
            print(f"    [截图] {filepath}")
            return filepath
        else:
            print(f"    [截图写入失败] {filepath}")
            return None
    except Exception as e:
        print(f"    [截图异常] {e}")
        return None

def generate_preview_image(video_path, output_path, rows=3, cols=3, thumb_size=(320,180)):
    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        print(f"    [预览] 无法打开视频: {video_path}")
        return None
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    fps = cap.get(cv2.CAP_PROP_FPS)
    dur = total_frames / fps if fps > 0 else 0
    if dur <= 0: cap.release(); return None
    num_slots = rows * cols
    step = dur / (num_slots + 1)
    times = [step * (i+1) for i in range(num_slots)]
    thumbs = []
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
        ret, frame = cap.read()
        if not ret: thumbs.append(np.zeros((thumb_size[1], thumb_size[0], 3), dtype=np.uint8))
        else: thumbs.append(cv2.resize(frame, thumb_size))
    cap.release()
    rows_list = [np.hstack([thumbs[r*cols + c] for c in range(cols)]) for r in range(rows)]
    preview = np.vstack(rows_list)
    if imwrite_chinese(output_path, preview):
        print(f"    [预览图] {output_path}")
        return output_path
    else:
        print(f"    [预览图写入失败] {output_path}")
        return None

# ---------- 七大质量检测函数 ----------
def detect_sharpness(path, video_name, abnormal_root):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened(): print(f"    [警告] 无法打开视频 (清晰度): {path}"); return None
    dur = get_video_duration(path)
    if not dur or dur<=0:
        fps = cap.get(cv2.CAP_PROP_FPS); tot = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        dur = tot/fps if fps>0 else 0
    inv = smart_interval(dur)
    times = [t for t in np.arange(0, dur, inv)]
    if not times: cap.release(); return None
    vals = []; abnormal_count = 0
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t*1000); ret, frm = cap.read()
        if not ret: continue
        gray = cv2.cvtColor(frm, cv2.COLOR_BGR2GRAY)
        lap = cv2.Laplacian(gray, cv2.CV_64F).var()
        vals.append(lap)
        if lap < SHARPNESS_THRESH:
            abnormal_count += 1
            save_abnormal_frame(frm, video_name, t, "模糊", abnormal_root)
    cap.release()
    if not vals: return None
    vals = np.array(vals); avg = np.mean(vals); mn = np.min(vals); mx = np.max(vals)
    bl = int(np.sum(vals < SHARPNESS_THRESH)); tot = len(vals); br = bl/tot*100 if tot else 0
    if avg >= 150 and br <= 10: judge = "OK"
    elif avg >= 80 or br <= 30: judge = "警告"
    else: judge = "NG"
    return {"avg": round(avg,1), "min": round(mn,1), "max": round(mx,1),
            "blurry_count": bl, "total": tot, "blurry_ratio": round(br,1), "judge": judge, "abnormal_screenshots": abnormal_count}

def detect_exposure(path, video_name, abnormal_root):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened(): print(f"    [警告] 无法打开视频 (曝光): {path}"); return None
    dur = get_video_duration(path)
    if not dur or dur<=0:
        fps = cap.get(cv2.CAP_PROP_FPS); tot = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        dur = tot/fps if fps>0 else 0
    inv = smart_interval(dur)
    times = [t for t in np.arange(0, dur, inv)]
    if not times: cap.release(); return None
    brights = []; u_s=0; u=0; n=0; o=0; o_s=0; total=0; abnormal_screenshots = 0
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t*1000); ret, frm = cap.read()
        if not ret: continue
        gray = cv2.cvtColor(frm, cv2.COLOR_BGR2GRAY); mean_v = np.mean(gray)
        brights.append(mean_v); total+=1
        is_abnormal = False; anomaly_type = None
        if mean_v <= EXP_UNDER_SEVERE: u_s+=1; is_abnormal = True; anomaly_type = "严重欠曝"
        elif mean_v <= EXP_UNDER: u+=1; is_abnormal = True; anomaly_type = "欠曝"
        elif mean_v <= EXP_OVER: n+=1
        elif mean_v <= EXP_OVER_SEVERE: o+=1; is_abnormal = True; anomaly_type = "过曝"
        else: o_s+=1; is_abnormal = True; anomaly_type = "严重过曝"
        if is_abnormal:
            abnormal_screenshots += 1
            save_abnormal_frame(frm, video_name, t, anomaly_type, abnormal_root)
    cap.release()
    if total==0: return None
    avg_b = np.mean(brights); mn_b = np.min(brights); mx_b = np.max(brights)
    under_r = (u_s+u)/total*100; over_r = (o+o_s)/total*100; normal_r = n/total*100
    if 61<=avg_b<=190 and (under_r+over_r)<=10: judge="OK"
    elif (31<=avg_b<=60 or 191<=avg_b<=225) and (under_r+over_r)<=30: judge="警告"
    else: judge="NG"
    return {"avg_b": round(avg_b,1), "min_b": round(mn_b,1), "max_b": round(mx_b,1),
            "under_ratio": round(under_r,1), "over_ratio": round(over_r,1), "normal_ratio": round(normal_r,1),
            "judge": judge, "abnormal_screenshots": abnormal_screenshots}

def detect_stability(path, video_name, abnormal_root):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened(): return None
    dur = get_video_duration(path)
    if not dur or dur<=0:
        fps = cap.get(cv2.CAP_PROP_FPS); tot = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        dur = tot/fps if fps>0 else 0
    inv = max(0.5, smart_interval(dur)/2)
    times = [t for t in np.arange(0, dur, inv)]
    if len(times)<2: cap.release(); return None
    prev = None; disps=[]
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t*1000); ret, frm = cap.read()
        if not ret: continue
        gray = cv2.cvtColor(frm, cv2.COLOR_BGR2GRAY)
        if prev is not None:
            flow = cv2.calcOpticalFlowFarneback(prev, gray, None, 0.5, 3, 15, 3, 5, 1.2, 0)
            mag, _ = cv2.cartToPolar(flow[...,0], flow[...,1])
            disps.append(np.mean(mag))
        prev = gray
    cap.release()
    if not disps: return None
    disps = np.array(disps); avg_d = np.mean(disps)
    shake_r = np.sum(disps > STABILITY_OK_THRESH)/len(disps)*100; std_d = np.std(disps)
    score = max(0, 100 - (std_d*10 + shake_r)); score = min(100, int(score))
    if score>=90: judge="优秀"
    elif score>=75: judge="OK"
    elif score>=60: judge="轻微抖动"
    else: judge="NG：抖动明显"
    return {"avg_disp": round(avg_d,2), "shake_ratio": round(shake_r,1), "score": score, "judge": judge}

def detect_frame_drops(path):
    cmd = [FFPROBE_PATH, "-v", "quiet", "-select_streams", "v:0", "-show_packets",
           "-show_entries", "packet=pts_time", "-of", "json", path]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        pts = [float(p["pts_time"]) for p in json.loads(res.stdout)["packets"] if "pts_time" in p]
        if len(pts)<2: return None
        intervals = np.diff(pts)
        nom_fps = 1.0/np.median(intervals) if np.median(intervals)!=0 else None
        act_fps = len(pts)/(pts[-1]-pts[0]) if pts[-1]>pts[0] else None
        th_int = 1.0/round(nom_fps) if nom_fps else None
        ab_mask = np.zeros(len(intervals), dtype=bool)
        if th_int: ab_mask = (intervals > th_int*1.5) | (intervals < th_int*0.5)
        ab_cnt = int(np.sum(ab_mask)); ab_ratio = ab_cnt/len(intervals)*100; max_int = np.max(intervals)*1000
        diff_pct = None
        if nom_fps and act_fps: diff_pct = abs(nom_fps - act_fps)/nom_fps*100
        if diff_pct is not None:
            if diff_pct<=3: judge="OK"
            elif diff_pct<=8: judge="警告：帧率略不稳定"
            else: judge="NG：疑似掉帧"
        else: judge="无法判断"
        return {"nom_fps": round(nom_fps,2) if nom_fps else "N/A", "act_fps": round(act_fps,2) if act_fps else "N/A",
                "diff_pct": round(diff_pct,1) if diff_pct else "N/A", "max_interval_ms": round(max_int,1),
                "abnormal_count": ab_cnt, "abnormal_ratio": round(ab_ratio,1), "judge": judge}
    except: return None

def detect_black_white(path, video_name, abnormal_root):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened(): return None
    dur = get_video_duration(path)
    if not dur or dur<=0:
        fps = cap.get(cv2.CAP_PROP_FPS); tot = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        dur = tot/fps if fps>0 else 0
    inv = smart_interval(dur)
    times = [t for t in np.arange(0, dur, inv)]
    if not times: cap.release(); return None
    b_segs=[]; w_segs=[]; b_start=None; w_start=None; b_cnt=0; w_cnt=0; total=0; abnormal_screenshots=0
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t*1000); ret, frm = cap.read()
        if not ret: continue; total+=1
        gray = cv2.cvtColor(frm, cv2.COLOR_BGR2GRAY); mean_v = np.mean(gray)
        b_pix = np.sum(gray<15)/gray.size; w_pix = np.sum(gray>240)/gray.size
        is_b = (mean_v<15 and b_pix>BLACK_PIX_RATIO) or (mean_v<30 and b_pix>SUS_BLACK_RATIO)
        is_w = (mean_v>240 and w_pix>WHITE_PIX_RATIO) or (mean_v>225 and w_pix>SUS_WHITE_RATIO)
        if is_b:
            b_cnt+=1; abnormal_screenshots+=1; save_abnormal_frame(frm, video_name, t, "黑屏", abnormal_root)
            if b_start is None: b_start=t
        else:
            if b_start is not None: b_segs.append((b_start,t)); b_start=None
        if is_w:
            w_cnt+=1; abnormal_screenshots+=1; save_abnormal_frame(frm, video_name, t, "白屏", abnormal_root)
            if w_start is None: w_start=t
        else:
            if w_start is not None: w_segs.append((w_start,t)); w_start=None
    if b_start: b_segs.append((b_start, times[-1]))
    if w_start: w_segs.append((w_start, times[-1]))
    cap.release()
    if total==0: return None
    b_ratio = b_cnt/total*100; w_ratio = w_cnt/total*100
    long_b = max([e-s for s,e in b_segs]) if b_segs else 0
    long_w = max([e-s for s,e in w_segs]) if w_segs else 0
    issues=[]
    if long_b>=3: issues.append("NG：黑屏≥3秒")
    elif long_b>=1: issues.append("警告：存在短暂黑屏")
    if long_w>=3: issues.append("NG：白屏≥3秒")
    elif long_w>=1: issues.append("警告：存在短暂白屏")
    judge = "；".join(issues) if issues else "正常"
    return {"b_segs": [f"{format_time(s)}-{format_time(e)}" for s,e in b_segs[:3]],
            "w_segs": [f"{format_time(s)}-{format_time(e)}" for s,e in w_segs[:3]],
            "long_b": round(long_b,2), "long_w": round(long_w,2), "b_ratio": round(b_ratio,1), "w_ratio": round(w_ratio,1),
            "judge": judge, "abnormal_screenshots": abnormal_screenshots}

def detect_frozen(path, video_name, abnormal_root):
    cap = cv2.VideoCapture(path)
    if not cap.isOpened(): return None
    dur = get_video_duration(path)
    if not dur or dur<=0:
        fps = cap.get(cv2.CAP_PROP_FPS); tot = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        dur = tot/fps if fps>0 else 0
    inv = smart_interval(dur)
    times = [t for t in np.arange(0, dur, inv)]
    if len(times)<2: cap.release(); return None
    prev=None; segs=[]; f_start=None; f_cnt=0; abnormal_screenshots=0
    for t in times:
        cap.set(cv2.CAP_PROP_POS_MSEC, t*1000); ret, frm = cap.read()
        if not ret: continue
        gray = cv2.cvtColor(frm, cv2.COLOR_BGR2GRAY)
        if prev is not None:
            diff = cv2.absdiff(prev, gray); mean_d = np.mean(diff)
            if mean_d < FREEZE_LOW_DIFF:
                if f_start is None: f_start = t - inv
                f_cnt+=1
                if f_cnt == 1:
                    abnormal_screenshots+=1; save_abnormal_frame(frm, video_name, t, "冻结", abnormal_root)
            else:
                if f_cnt>0 and f_start is not None: segs.append((f_start, t, t-f_start))
                f_start=None; f_cnt=0
        prev=gray
    if f_cnt>0 and f_start is not None: segs.append((f_start, times[-1], times[-1]-f_start))
    cap.release()
    if not segs: return {"segs":[], "longest":0, "count":0, "total_dur":0, "ratio":0, "judge":"正常", "abnormal_screenshots":0}
    longest = max(s[2] for s in segs); total_dur = sum(s[2] for s in segs)
    ratio = total_dur/dur*100 if dur>0 else 0; count = len(segs)
    if longest>3: judge="严重 NG"
    elif longest>1: judge="NG：存在明显画面冻结"
    elif longest>0.5: judge="警告"
    else: judge="正常"
    return {"segs": [f"{format_time(s)}-{format_time(e)}" for s,e,_ in segs],
            "longest": round(longest,2), "count": count, "total_dur": round(total_dur,2), "ratio": round(ratio,1),
            "judge": judge, "abnormal_screenshots": abnormal_screenshots}

def analyze_audio(path):
    info = extract_video_info(path).get("audio_info")
    if not info: return {"has_audio": False}
    cmd_s = [FFMPEG_PATH, "-i", path, "-af", "silencedetect=noise=-30dB:d=0.5", "-f", "null", "-"]
    sil_segs=[]
    try:
        res=subprocess.run(cmd_s, capture_output=True, text=True, stderr=subprocess.STDOUT)
        out=res.stdout+res.stderr; starts=[]; ends=[]
        for line in out.splitlines():
            if "silence_start" in line: starts.append(float(line.split("silence_start:")[1].strip().split()[0]))
            elif "silence_end" in line: ends.append(float(line.split("silence_end:")[1].strip().split()[0]))
        sil_segs=[(s,e) for s,e in zip(starts, ends)]
    except: pass
    cmd_v = [FFMPEG_PATH, "-i", path, "-af", "volumedetect", "-f", "null", "-"]
    mean_db=None
    try:
        res=subprocess.run(cmd_v, capture_output=True, text=True, stderr=subprocess.STDOUT)
        out=res.stdout+res.stderr
        for line in out.splitlines():
            if "mean_volume:" in line: mean_db=float(line.split("mean_volume:")[1].strip().split()[0])
    except: pass
    issues=[]
    br=info.get("bitrate")
    if br:
        if br<64000: issues.append("NG：音频码率过低")
        elif br<96000: issues.append("警告：音质可能偏低")
    sr=info.get("sample_rate")
    if sr.isdigit():
        srv=int(sr)
        if srv<22050: issues.append("NG：采样率过低")
        elif srv<32000: issues.append("警告：采样率略低")
    if mean_db is not None:
        if mean_db<SILENCE_DBFS: issues.append("静音/接近无声")
        elif mean_db<LOW_VOL_DBFS: issues.append("音量偏低")
        elif mean_db>HIGH_VOL_DBFS: issues.append("可能爆音/削波")
    sil_cnt=len(sil_segs)
    if sil_cnt>0: issues.append(f"检测到{sil_cnt}段静音")
    judge="；".join(issues) if issues else "OK"
    return {"has_audio":True, "codec":info.get("codec","N/A"), "sample_rate":sr,
            "channels":info.get("channels","N/A"), "bitrate_kbps":round(br/1000,1) if br else "N/A",
            "mean_volume_db": round(mean_db,1) if mean_db else "N/A",
            "silence_segs": [f"{format_time(s)}-{format_time(e)}" for s,e in sil_segs], "judge": judge}

# ---------- 评分计算（已增加空值保护）----------
def score_from_judge(judge_str):
    if not judge_str: return 50
    if judge_str == "OK" or judge_str == "优秀" or judge_str == "正常": return 100
    if "警告" in judge_str: return 60
    if "NG" in judge_str: return 0
    if "轻微抖动" in judge_str: return 60
    return 50

def compute_scores(vname, bitrate_judge, quality):
    scores = {}
    scores["码率"] = score_from_judge(bitrate_judge)

    # 使用 or {} 防止 None
    sharp = quality.get("sharpness") or {}
    scores["清晰度"] = score_from_judge(sharp.get("judge", ""))

    exp = quality.get("exposure") or {}
    scores["曝光"] = score_from_judge(exp.get("judge", ""))

    stab = quality.get("stability") or {}
    scores["稳定性"] = stab.get("score", 50)

    drops = quality.get("frame_drops") or {}
    scores["掉帧"] = score_from_judge(drops.get("judge", ""))

    bw = (quality.get("black_white") or {}).get("judge", "正常")
    fz = (quality.get("frozen") or {}).get("judge", "正常")
    bw_fz_score = 0 if "NG" in bw or "NG" in fz else (60 if "警告" in bw or "警告" in fz else 100)
    scores["黑屏/白屏/冻结"] = bw_fz_score

    audio = quality.get("audio") or {}
    if not audio.get("has_audio"):
        scores["音频"] = 100
    else:
        scores["音频"] = score_from_judge(audio.get("judge", ""))

    weights = {"码率":0.20, "清晰度":0.25, "曝光":0.15, "稳定性":0.15, "掉帧":0.10, "黑屏/白屏/冻结":0.10, "音频":0.05}
    total = sum(scores[k]*weights[k] for k in weights)
    scores["总分"] = round(total, 1)
    if total >= 90: level = "优秀"
    elif total >= 75: level = "良好"
    elif total >= 60: level = "一般"
    else: level = "较差"
    scores["等级"] = level
    return scores

# ---------- Excel 生成（不变）----------
def set_header(ws, headers, row=1, start_col=1):
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    al = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, h in enumerate(headers, start=start_col):
        cell = ws.cell(row=row, column=i, value=h); cell.font = hf; cell.fill = hfill; cell.alignment = al

def apply_style(ws, ncols):
    for col in range(1, ncols+1):
        letter = get_column_letter(col); max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        length = sum(2 if ord(c)>127 else 1 for c in line)
                        max_len = max(max_len, length)
        ws.column_dimensions[letter].width = max_len + 3
    ws.freeze_panes = "A2"

def add_summary_sheet(wb, all_bitrate, mode_config):
    ws = wb.active; ws.title = "Summary"
    ws.merge_cells("A1:K1"); ws["A1"] = f"视频码率分析报告（{mode_config['name']}）"; ws["A1"].font = Font(bold=True, size=14)
    headers = ["视频文件", "分辨率", "总码率", "视频码率", "音频码率", "帧率", "总帧数", "码率判断"]
    set_header(ws, headers, row=4)
    for i, (vname, info, judge) in enumerate(all_bitrate, 5):
        ws.cell(row=i, column=1, value=vname)
        ws.cell(row=i, column=2, value=f"{info.get('width')}x{info.get('height')}" if info.get('width') else "N/A")
        ws.cell(row=i, column=3, value=format_bps(info.get("total_bitrate")))
        ws.cell(row=i, column=4, value=format_bps(info.get("video_bitrate")))
        ws.cell(row=i, column=5, value=format_bps(info.get("audio_bitrate")))
        ws.cell(row=i, column=6, value=f"{info.get('fps'):.2f}" if info.get('fps') else "N/A")
        ws.cell(row=i, column=7, value=info.get("total_frames") if info.get("total_frames") else "N/A")
        cell = ws.cell(row=i, column=8, value=judge); cell.alignment = Alignment(horizontal="center")
        if judge == "OK": cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        elif judge == "NG": cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    apply_style(ws, len(headers))

def add_standard_sheet(wb, mode_config):
    ws = wb.create_sheet("码率评判标准", 1)
    ws.merge_cells("A1:E1"); ws["A1"] = f"评判标准 - {mode_config['name']}"; ws["A1"].font = Font(bold=True, size=13)
    hds = ["拍摄模式/分辨率", "分辨率", "帧率", "建议码率", "检测判断"]
    set_header(ws, hds, row=3)
    for i, row_data in enumerate(mode_config["table"], 4):
        for j, val in enumerate(row_data, 1):
            ws.cell(row=i, column=j, value=val).alignment = Alignment(vertical="center", wrap_text=True)
    apply_style(ws, len(hds))

def add_quality_sheets(wb, video_names, quality_data):
    # 清晰度
    ws = wb.create_sheet("清晰度检测")
    hds = ["视频文件", "平均清晰度", "最低清晰度", "最高清晰度", "模糊帧数量", "抽帧总数", "模糊帧占比(%)", "判断", "异常截图数"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("sharpness") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=d.get("avg", "N/A"))
        ws.cell(row=i, column=3, value=d.get("min", "N/A"))
        ws.cell(row=i, column=4, value=d.get("max", "N/A"))
        ws.cell(row=i, column=5, value=d.get("blurry_count", "N/A"))
        ws.cell(row=i, column=6, value=d.get("total", "N/A"))
        ws.cell(row=i, column=7, value=d.get("blurry_ratio", "N/A"))
        ws.cell(row=i, column=8, value=d.get("judge", "N/A"))
        ws.cell(row=i, column=9, value=d.get("abnormal_screenshots", 0))
    apply_style(ws, len(hds))

    # 曝光
    ws = wb.create_sheet("曝光检测")
    hds = ["视频文件", "平均亮度", "最低亮度", "最高亮度", "欠曝帧占比(%)", "过曝帧占比(%)", "正常帧占比(%)", "判断", "异常截图数"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("exposure") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=d.get("avg_b", "N/A"))
        ws.cell(row=i, column=3, value=d.get("min_b", "N/A"))
        ws.cell(row=i, column=4, value=d.get("max_b", "N/A"))
        ws.cell(row=i, column=5, value=d.get("under_ratio", "N/A"))
        ws.cell(row=i, column=6, value=d.get("over_ratio", "N/A"))
        ws.cell(row=i, column=7, value=d.get("normal_ratio", "N/A"))
        ws.cell(row=i, column=8, value=d.get("judge", "N/A"))
        ws.cell(row=i, column=9, value=d.get("abnormal_screenshots", 0))
    apply_style(ws, len(hds))

    # 稳定性
    ws = wb.create_sheet("稳定性检测")
    hds = ["视频文件", "平均位移量(px)", "抖动帧占比(%)", "稳定性评分", "判断"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("stability") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=d.get("avg_disp", "N/A"))
        ws.cell(row=i, column=3, value=d.get("shake_ratio", "N/A"))
        ws.cell(row=i, column=4, value=d.get("score", "N/A"))
        ws.cell(row=i, column=5, value=d.get("judge", "N/A"))
    apply_style(ws, len(hds))

    # 掉帧
    ws = wb.create_sheet("掉帧检测")
    hds = ["视频文件", "标称fps", "实际fps", "帧率差值(%)", "最大帧间隔(ms)", "异常帧数量", "异常帧占比(%)", "判断"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("frame_drops") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=d.get("nom_fps", "N/A"))
        ws.cell(row=i, column=3, value=d.get("act_fps", "N/A"))
        ws.cell(row=i, column=4, value=d.get("diff_pct", "N/A"))
        ws.cell(row=i, column=5, value=d.get("max_interval_ms", "N/A"))
        ws.cell(row=i, column=6, value=d.get("abnormal_count", "N/A"))
        ws.cell(row=i, column=7, value=d.get("abnormal_ratio", "N/A"))
        ws.cell(row=i, column=8, value=d.get("judge", "N/A"))
    apply_style(ws, len(hds))

    # 黑屏白屏
    ws = wb.create_sheet("黑屏白屏检测")
    hds = ["视频文件", "黑屏片段", "白屏片段", "最长黑屏(s)", "最长白屏(s)", "黑屏帧占比(%)", "白屏帧占比(%)", "判断", "异常截图数"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("black_white") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=", ".join(d.get("b_segs", [])) or "无")
        ws.cell(row=i, column=3, value=", ".join(d.get("w_segs", [])) or "无")
        ws.cell(row=i, column=4, value=d.get("long_b", "N/A"))
        ws.cell(row=i, column=5, value=d.get("long_w", "N/A"))
        ws.cell(row=i, column=6, value=d.get("b_ratio", "N/A"))
        ws.cell(row=i, column=7, value=d.get("w_ratio", "N/A"))
        ws.cell(row=i, column=8, value=d.get("judge", "N/A"))
        ws.cell(row=i, column=9, value=d.get("abnormal_screenshots", 0))
    apply_style(ws, len(hds))

    # 画面冻结
    ws = wb.create_sheet("画面冻结检测")
    hds = ["视频文件", "冻结片段", "最长冻结(s)", "冻结次数", "冻结总时长(s)", "冻结占比(%)", "判断", "异常截图数"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("frozen") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value=", ".join(d.get("segs", [])) or "无")
        ws.cell(row=i, column=3, value=d.get("longest", "N/A"))
        ws.cell(row=i, column=4, value=d.get("count", "N/A"))
        ws.cell(row=i, column=5, value=d.get("total_dur", "N/A"))
        ws.cell(row=i, column=6, value=d.get("ratio", "N/A"))
        ws.cell(row=i, column=7, value=d.get("judge", "N/A"))
        ws.cell(row=i, column=8, value=d.get("abnormal_screenshots", 0))
    apply_style(ws, len(hds))

    # 音频
    ws = wb.create_sheet("音频基础检测")
    hds = ["视频文件", "是否有音频", "编码", "采样率(Hz)", "声道", "码率(kbps)", "平均音量(dBFS)", "静音片段", "判断"]
    set_header(ws, hds)
    for i, vn in enumerate(video_names, 2):
        d = quality_data[vn].get("audio") or {}
        ws.cell(row=i, column=1, value=vn)
        ws.cell(row=i, column=2, value="有" if d.get("has_audio") else "无")
        ws.cell(row=i, column=3, value=d.get("codec", "N/A"))
        ws.cell(row=i, column=4, value=d.get("sample_rate", "N/A"))
        ws.cell(row=i, column=5, value=d.get("channels", "N/A"))
        ws.cell(row=i, column=6, value=d.get("bitrate_kbps", "N/A"))
        ws.cell(row=i, column=7, value=d.get("mean_volume_db", "N/A"))
        ws.cell(row=i, column=8, value=", ".join(d.get("silence_segs", [])) or "无")
        ws.cell(row=i, column=9, value=d.get("judge", "N/A"))
    apply_style(ws, len(hds))

def add_scoring_standard_sheet(wb):
    ws = wb.create_sheet("评分标准")
    ws.merge_cells("A1:D1"); ws["A1"] = "视频综合评分权重与等级"; ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A3:B3"); ws["A3"] = "评分模块与权重"
    set_header(ws, ["模块", "权重"], row=4)
    weights = [("码率/分辨率/帧率基础参数", "20%"), ("视频清晰度", "25%"), ("视频曝光", "15%"),
               ("视频稳定性", "15%"), ("卡顿/掉帧", "10%"), ("黑屏/白屏/冻结", "10%"), ("音频基础检测", "5%")]
    for i, (mod, w) in enumerate(weights, 5):
        ws.cell(row=i, column=1, value=mod); ws.cell(row=i, column=2, value=w)
    ws.merge_cells("D3:E3"); ws["D3"] = "等级划分"
    set_header(ws, ["综合评分", "等级", "判断"], row=4, start_col=4)
    levels = [(90, 100, "优秀", "OK"), (75, 89, "良好", "OK"), (60, 74, "一般", "警告"), (0, 59, "较差", "NG")]
    for i, (low, high, lvl, judge) in enumerate(levels, 5):
        ws.cell(row=i, column=4, value=f"{low}-{high}"); ws.cell(row=i, column=5, value=lvl); ws.cell(row=i, column=6, value=judge)
    apply_style(ws, 6)

def add_scores_sheet(wb, video_names, scores):
    ws = wb.create_sheet("综合评分")
    headers = ["视频文件", "码率(20%)", "清晰度(25%)", "曝光(15%)", "稳定性(15%)",
               "掉帧(10%)", "黑屏/白屏/冻结(10%)", "音频(5%)", "加权总分", "等级"]
    set_header(ws, headers)
    for i, vn in enumerate(video_names, 2):
        ws.cell(row=i, column=1, value=vn); sc = scores[vn]
        ws.cell(row=i, column=2, value=sc["码率"]); ws.cell(row=i, column=3, value=sc["清晰度"]); ws.cell(row=i, column=4, value=sc["曝光"]); ws.cell(row=i, column=5, value=sc["稳定性"]); ws.cell(row=i, column=6, value=sc["掉帧"]); ws.cell(row=i, column=7, value=sc["黑屏/白屏/冻结"]); ws.cell(row=i, column=8, value=sc["音频"]); ws.cell(row=i, column=9, value=sc["总分"]); ws.cell(row=i, column=10, value=sc["等级"])
    apply_style(ws, len(headers))

def add_preview_sheet(wb, video_names, preview_paths):
    ws = wb.create_sheet("视频预览图")
    ws.merge_cells("A1:E1"); ws["A1"] = "视频抽帧预览图"; ws["A1"].font = Font(bold=True, size=14)
    for col in range(1, 3): ws.column_dimensions[get_column_letter(col)].width = 40
    for row_idx, vname in enumerate(video_names, 2):
        ws.cell(row=row_idx, column=1, value=vname)
        path = preview_paths.get(vname)
        if path and os.path.exists(path):
            try:
                img = XLImage(path); img.width = 300; img.height = int(img.height * (300 / img.width))
                ws.add_image(img, f"B{row_idx}"); ws.row_dimensions[row_idx].height = img.height + 10
            except Exception as e:
                ws.cell(row=row_idx, column=2, value=f"预览图加载失败: {e}")

def add_charts_sheet(wb, all_bitrate, quality_data):
    if not all_bitrate: return
    ws = wb.create_sheet("统计图表")
    video_names = []; sharpness_avg = []; exposure_avg = []; bitrate_mbps = []; drop_ratio = []
    ok_count = ng_count = na_count = 0
    anomaly_keys = ["清晰度", "曝光", "稳定性", "掉帧", "黑屏/白屏/冻结", "音频"]
    anomaly_counts = {k: 0 for k in anomaly_keys}
    for vname, info, judge in all_bitrate:
        video_names.append(vname)
        if judge == "OK": ok_count += 1
        elif judge == "NG": ng_count += 1
        else: na_count += 1
        v_br = info.get("video_bitrate"); bitrate_mbps.append(round(v_br/1e6,2) if v_br else 0)
        q = quality_data.get(vname, {})
        sharp = q.get("sharpness") or {}; sharpness_avg.append(sharp.get("avg",0))
        exp = q.get("exposure") or {}; exposure_avg.append(exp.get("avg_b",0))
        drops = q.get("frame_drops") or {}; drop_ratio.append(drops.get("abnormal_ratio",0))
        if sharp.get("judge","") in ("警告","NG"): anomaly_counts["清晰度"] += 1
        if exp.get("judge","") in ("警告","NG"): anomaly_counts["曝光"] += 1
        stab = q.get("stability") or {}
        if stab.get("judge","") not in ("优秀","OK","正常"): anomaly_counts["稳定性"] += 1
        if drops.get("judge","") in ("警告：帧率略不稳定","NG：疑似掉帧"): anomaly_counts["掉帧"] += 1
        bw = q.get("black_white") or {}; fz = q.get("frozen") or {}
        if "NG" in bw.get("judge","") or "NG" in fz.get("judge","") or "警告" in bw.get("judge","") or "警告" in fz.get("judge",""):
            anomaly_counts["黑屏/白屏/冻结"] += 1
        audio = q.get("audio") or {}
        if audio.get("has_audio") and audio.get("judge","") not in ("OK","正常"): anomaly_counts["音频"] += 1

    for i, name in enumerate(video_names, 1):
        ws.cell(row=i, column=1, value=name); ws.cell(row=i, column=2, value=sharpness_avg[i-1]); ws.cell(row=i, column=3, value=exposure_avg[i-1]); ws.cell(row=i, column=4, value=bitrate_mbps[i-1]); ws.cell(row=i, column=5, value=drop_ratio[i-1])
    ws.cell(row=1, column=7, value="OK"); ws.cell(row=2, column=7, value="NG"); ws.cell(row=3, column=7, value="N/A")
    ws.cell(row=1, column=8, value=ok_count); ws.cell(row=2, column=8, value=ng_count); ws.cell(row=3, column=8, value=na_count)
    for i, key in enumerate(anomaly_keys, 1): ws.cell(row=i, column=10, value=key); ws.cell(row=i, column=11, value=anomaly_counts[key])

    if len(video_names) > 0:
        pie = PieChart(); pie.title="码率判断 OK/NG 分布"; pie.style=10
        pie_data = Reference(ws, min_col=8, min_row=1, max_row=3); pie_cats = Reference(ws, min_col=7, min_row=1, max_row=3)
        pie.add_data(pie_data); pie.set_categories(pie_cats); pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent=True; pie.dataLabels.showVal=True
        ws.add_chart(pie, "G5")
        bar1 = BarChart(); bar1.type="col"; bar1.title="各模块异常视频数量"; bar1.y_axis.title="视频数"; bar1.style=10
        bar1_data = Reference(ws, min_col=11, min_row=1, max_row=len(anomaly_keys)); bar1_cats = Reference(ws, min_col=10, min_row=1, max_row=len(anomaly_keys))
        bar1.add_data(bar1_data); bar1.set_categories(bar1_cats); ws.add_chart(bar1, "G20")
        bar2 = BarChart(); bar2.type="col"; bar2.title="清晰度分布"; bar2.y_axis.title="平均清晰度"; bar2.style=10
        bar2_data = Reference(ws, min_col=2, min_row=1, max_row=len(video_names)); bar2_cats = Reference(ws, min_col=1, min_row=1, max_row=len(video_names))
        bar2.add_data(bar2_data); bar2.set_categories(bar2_cats); ws.add_chart(bar2, "G35")
        bar3 = BarChart(); bar3.type="col"; bar3.title="曝光亮度分布"; bar3.y_axis.title="平均亮度"; bar3.style=10
        bar3_data = Reference(ws, min_col=3, min_row=1, max_row=len(video_names)); bar3_cats = Reference(ws, min_col=1, min_row=1, max_row=len(video_names))
        bar3.add_data(bar3_data); bar3.set_categories(bar3_cats); ws.add_chart(bar3, "G50")
        bar4 = BarChart(); bar4.type="col"; bar4.title="码率分布 (Mbps)"; bar4.y_axis.title="Mbps"; bar4.style=10
        bar4_data = Reference(ws, min_col=4, min_row=1, max_row=len(video_names)); bar4_cats = Reference(ws, min_col=1, min_row=1, max_row=len(video_names))
        bar4.add_data(bar4_data); bar4.set_categories(bar4_cats); ws.add_chart(bar4, "G65")
        bar5 = BarChart(); bar5.type="col"; bar5.title="掉帧异常占比 (%)"; bar5.y_axis.title="%"; bar5.style=10
        bar5_data = Reference(ws, min_col=5, min_row=1, max_row=len(video_names)); bar5_cats = Reference(ws, min_col=1, min_row=1, max_row=len(video_names))
        bar5.add_data(bar5_data); bar5.set_categories(bar5_cats); ws.add_chart(bar5, "G80")
    ws.column_dimensions['A'].width=18; ws.column_dimensions['G'].width=20; ws.column_dimensions['J'].width=20

# ---------- 主程序 ----------
def main():
    mode = select_mode()
    mode_config = STANDARDS_CONFIG[mode]
    print(f"当前模式：{mode_config['name']}")
    if not os.path.isdir(VIDEO_FOLDER):
        print(f"文件夹不存在：{VIDEO_FOLDER}"); return
    videos = []
    for f in os.listdir(VIDEO_FOLDER):
        fp = os.path.join(VIDEO_FOLDER, f)
        if os.path.isfile(fp) and os.path.splitext(f)[1].lower() in VIDEO_EXTS:
            videos.append(fp)
    if not videos:
        print("未找到视频文件"); return
    print(f"找到 {len(videos)} 个视频，开始分析...\n")

    abnormal_root = os.path.join(VIDEO_FOLDER, ABNORMAL_SAVE_DIR)
    preview_dir = os.path.join(VIDEO_FOLDER, "预览图")
    os.makedirs(preview_dir, exist_ok=True)
    print(f"异常截图目录: {abnormal_root}")
    print(f"预览图目录: {preview_dir}\n")

    all_bitrate = []; quality_data = {}; preview_paths = {}
    for vpath in videos:
        vname = os.path.basename(vpath)
        print(f"分析: {vname}")
        info = extract_video_info(vpath)
        if not info: print("  无法提取媒体信息，跳过"); continue
        pixel = info.get("pixel_count"); fps = info.get("fps"); threshold = None
        if pixel and fps:
            for max_pix, f_lim, n_min, h_min in mode_config["standards"]:
                if pixel <= max_pix:
                    threshold = h_min if (h_min and fps > f_lim) else n_min; break
        v_br = info.get("video_bitrate")
        judge = "OK" if (v_br and threshold and v_br >= threshold*1e6) else ("NG" if v_br and threshold else "N/A")
        all_bitrate.append((vname, info, judge))

        qual = {
            "sharpness": detect_sharpness(vpath, vname, abnormal_root),
            "exposure": detect_exposure(vpath, vname, abnormal_root),
            "stability": detect_stability(vpath, vname, abnormal_root),
            "frame_drops": detect_frame_drops(vpath),
            "black_white": detect_black_white(vpath, vname, abnormal_root),
            "frozen": detect_frozen(vpath, vname, abnormal_root),
            "audio": analyze_audio(vpath)
        }
        quality_data[vname] = qual

        preview_path = os.path.join(preview_dir, f"{os.path.splitext(vname)[0]}_preview.jpg")
        if generate_preview_image(vpath, preview_path, PREVIEW_ROWS, PREVIEW_COLS, PREVIEW_SIZE):
            preview_paths[vname] = preview_path

    if not all_bitrate:
        print("没有成功分析任何视频，退出。")
        return

    scores = {}
    for vname, info, judge in all_bitrate:
        scores[vname] = compute_scores(vname, judge, quality_data[vname])

    wb = Workbook()
    add_summary_sheet(wb, all_bitrate, mode_config)
    add_standard_sheet(wb, mode_config)
    add_quality_sheets(wb, [v[0] for v in all_bitrate], quality_data)
    add_scoring_standard_sheet(wb)
    add_scores_sheet(wb, [v[0] for v in all_bitrate], scores)
    add_preview_sheet(wb, [v[0] for v in all_bitrate], preview_paths)
    add_charts_sheet(wb, all_bitrate, quality_data)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(VIDEO_FOLDER, f"{OUTPUT_NAME}_{timestamp}.xlsx")
    wb.save(out_path)
    print(f"\n报告已生成: {out_path}")

if __name__ == "__main__":
    main()