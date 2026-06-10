#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
OCR 图像识别一体化工具（测试工程师专业版）—— 完整交互 + 报告美化
"""

import os
os.environ["FLAGS_use_mkldnn"] = "0"
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
os.environ["PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK"] = "True"

import re
import cv2
import shutil
import argparse
import datetime
import numpy as np
from typing import List, Dict, Tuple, Generator, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

try:
    from paddleocr import PaddleOCR
except ImportError:
    print("❌ 缺少 paddleocr，请执行：pip install paddleocr")
    raise SystemExit(1)

IMAGE_EXTS = [".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"]
VIDEO_EXTS = [".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv", ".mpeg", ".mpg"]

# ===================== 样式定义 =====================
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="微软雅黑", size=10)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

def apply_header_style(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

def apply_data_style(ws, start_row=2, center_cols=None):
    center_cols = center_cols or []
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if cell.column in center_cols:
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN

# ===================== 工具函数 =====================
def read_image_chinese_path(path: str) -> Optional[np.ndarray]:
    if not os.path.exists(path):
        return None
    try:
        data = np.fromfile(path, dtype=np.uint8)
        img = cv2.imdecode(data, cv2.IMREAD_COLOR)
        return img
    except Exception:
        return None

def write_image_chinese_path(path: str, image: np.ndarray) -> bool:
    try:
        ext = os.path.splitext(path)[1] or ".jpg"
        ok, buf = cv2.imencode(ext, image)
        if ok:
            buf.tofile(path)
            return True
        return False
    except Exception:
        return False

def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", name)

def normalize_confidence(value: float) -> float:
    return value / 100 if value > 1 else value

def infer_mode(source: str) -> str:
    if os.path.isdir(source):
        return "batch"
    if os.path.splitext(source)[1].lower() in VIDEO_EXTS:
        return "video"
    return "single"

def ensure_dir(path: str):
    os.makedirs(path, exist_ok=True)

def polygon_area(points):
    if len(points) < 3:
        return 0.0
    area = 0.0
    n = len(points)
    for i in range(n):
        x1, y1 = points[i]
        x2, y2 = points[(i+1) % n]
        area += x1 * y2 - x2 * y1
    return abs(area) / 2.0

# ===================== 图像预处理模块 =====================
class ImagePreprocessor:
    @staticmethod
    def clahe(img, clip=2.0, tile=8):
        if len(img.shape) == 3:
            lab = cv2.cvtColor(img, cv2.COLOR_BGR2LAB)
            l, a, b = cv2.split(lab)
            clahe = cv2.createCLAHE(clipLimit=clip, tileGridSize=(tile, tile))
            l = clahe.apply(l)
            return cv2.cvtColor(cv2.merge((l, a, b)), cv2.COLOR_LAB2BGR)
        else:
            return cv2.createCLAHE(clipLimit=clip, tileGridSize=(tile, tile)).apply(img)

    @staticmethod
    def denoise(img, method="gaussian", ksize=5):
        if ksize % 2 == 0: ksize += 1
        if method == "gaussian":
            return cv2.GaussianBlur(img, (ksize, ksize), 0)
        elif method == "median":
            return cv2.medianBlur(img, ksize)
        elif method == "bilateral":
            return cv2.bilateralFilter(img, d=ksize, sigmaColor=75, sigmaSpace=75)
        return img

    @staticmethod
    def sharpen(img, strength=1.0):
        blurred = cv2.GaussianBlur(img, (0,0), 3)
        return cv2.addWeighted(img, 1.0+strength, blurred, -strength, 0)

    @staticmethod
    def scale(img, factor=2.0):
        h, w = img.shape[:2]
        return cv2.resize(img, (int(w*factor), int(h*factor)), interpolation=cv2.INTER_CUBIC)

    @staticmethod
    def rotate_correct(img):
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY) if len(img.shape)==3 else img
        edges = cv2.Canny(gray, 50, 150)
        lines = cv2.HoughLinesP(edges, 1, np.pi/180, threshold=100, minLineLength=100, maxLineGap=10)
        if lines is None: return img
        angles = [np.arctan2(y2-y1, x2-x1)*180/np.pi for x1,y1,x2,y2 in lines[:,0]]
        med = np.median(angles)
        if abs(med) < 0.5: return img
        h, w = img.shape[:2]
        M = cv2.getRotationMatrix2D((w/2, h/2), med, 1.0)
        return cv2.warpAffine(img, M, (w, h), borderMode=cv2.BORDER_CONSTANT, borderValue=(255,255,255))

    @staticmethod
    def to_gray(img):
        if len(img.shape) == 2: return cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
        return cv2.cvtColor(cv2.cvtColor(img, cv2.COLOR_BGR2GRAY), cv2.COLOR_GRAY2BGR)

    @staticmethod
    def gamma_correction(img, gamma=1.2):
        inv = 1.0/gamma
        table = np.array([(i/255.0)**inv*255 for i in range(256)]).astype("uint8")
        return cv2.LUT(img, table)

    @classmethod
    def process(cls, img, steps, **kwargs):
        out = img.copy()
        applied = []
        for s in steps:
            s = s.strip().lower()
            if s == "clahe":
                out = cls.clahe(out, kwargs.get("clahe_clip",2.0))
                applied.append(f"CLAHE(clip={kwargs.get('clahe_clip',2.0)})")
            elif s == "denoise":
                m = kwargs.get("denoise_method","gaussian")
                k = kwargs.get("denoise_ksize",5)
                out = cls.denoise(out, m, k)
                applied.append(f"Denoise({m}, ksize={k})")
            elif s == "sharpen":
                out = cls.sharpen(out, kwargs.get("sharpen_strength",1.0))
                applied.append(f"Sharpen(strength={kwargs.get('sharpen_strength',1.0)})")
            elif s == "scale":
                out = cls.scale(out, kwargs.get("scale_factor",2.0))
                applied.append(f"Scale(factor={kwargs.get('scale_factor',2.0)})")
            elif s == "rotate":
                out = cls.rotate_correct(out)
                applied.append("RotateCorrect")
            elif s == "gray":
                out = cls.to_gray(out)
                applied.append("Grayscale")
            elif s == "gamma":
                out = cls.gamma_correction(out, kwargs.get("gamma",1.2))
                applied.append(f"Gamma(gamma={kwargs.get('gamma',1.2)})")
            else:
                print(f"⚠️ 未知预处理步骤：{s}")
        return out, " -> ".join(applied) if applied else "None"

# ===================== 输入模块 =====================
class InputModule:
    def __init__(self, source, mode="single", frame_interval=30):
        self.source = source
        self.mode = mode
        self.frame_interval = max(1, frame_interval)

    def get_data(self) -> Generator[Tuple[str, Optional[np.ndarray], Optional[str]], None, None]:
        if self.mode == "single":
            yield from self._read_single()
        elif self.mode == "batch":
            yield from self._read_batch()
        elif self.mode == "video":
            yield from self._read_video()
        else:
            raise ValueError(f"不支持的模式：{self.mode}")

    def _read_single(self):
        img = read_image_chinese_path(self.source)
        name = os.path.basename(self.source)
        yield name, img, self.source

    def _read_batch(self):
        if not os.path.isdir(self.source):
            raise ValueError("批量模式需要文件夹")
        paths = []
        skip = {"ocr_output","reports","texts","annotated_images","ng_images"}
        for root, dirs, files in os.walk(self.source):
            dirs[:] = [d for d in dirs if d not in skip]
            for f in files:
                if os.path.splitext(f)[1].lower() in IMAGE_EXTS:
                    paths.append(os.path.join(root, f))
        if not paths:
            print("⚠️ 文件夹中没有图片")
            return
        for p in sorted(paths):
            img = read_image_chinese_path(p)
            name = os.path.basename(p)
            yield name, img, p

    def _read_video(self):
        if not os.path.isfile(self.source):
            raise ValueError(f"视频模式需要文件，当前路径：{self.source}")
        cap = cv2.VideoCapture(self.source)
        if not cap.isOpened():
            raise IOError(f"无法打开视频：{self.source}")
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = cap.get(cv2.CAP_PROP_FPS)
        print(f"视频总帧数：{total}，帧率：{fps:.2f}，抽帧间隔：{self.frame_interval}")
        idx = 0
        saved = 0
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            if idx % self.frame_interval == 0:
                yield f"frame_{saved:06d}.jpg", frame, None
                saved += 1
            idx += 1
        cap.release()
        print(f"视频抽帧结束，共输出 {saved} 帧")

# ===================== OCR 引擎 =====================
class OCREngine:
    def __init__(self, lang="ch", use_gpu=False):
        print("初始化 PaddleOCR...")
        self.ocr = PaddleOCR(lang=lang, device="gpu" if use_gpu else "cpu",
                             use_doc_orientation_classify=False,
                             use_doc_unwarping=False,
                             use_textline_orientation=True)
        print("PaddleOCR 初始化完成。\n")

    def recognize(self, img):
        if img is None: return []
        if len(img.shape) == 2:
            img = cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
        try:
            raw = self.ocr.predict(img)
        except Exception as e:
            print(f"⚠️ OCR错误：{e}")
            return []
        return self._parse(raw)

    def _parse(self, raw):
        results = []
        for item in raw:
            data = self._to_dict(item)
            if not data: continue
            if "res" in data and isinstance(data["res"], dict):
                data = data["res"]
            texts = data.get("rec_texts", [])
            scores = data.get("rec_scores", [])
            boxes = data.get("rec_polys") or data.get("dt_polys", [])
            for i, txt in enumerate(texts):
                if not txt or str(txt).strip() == "": continue
                conf = float(scores[i]) if i < len(scores) else 0.0
                bbox = boxes[i] if i < len(boxes) else [[0,0],[0,0],[0,0],[0,0]]
                if hasattr(bbox, "tolist"): bbox = bbox.tolist()
                results.append({"bbox": bbox, "text": str(txt), "confidence": conf})
        return results

    def _to_dict(self, item):
        if isinstance(item, dict): return item
        for attr in ["to_dict","json","res"]:
            if hasattr(item, attr):
                try:
                    v = getattr(item, attr)
                    return v() if callable(v) else v
                except: pass
        return item.__dict__ if hasattr(item, "__dict__") else None

# ===================== 结果分析 =====================
class ResultAnalyzer:
    def __init__(self, keywords=None, min_char_conf=0.5):
        self.keywords = keywords or []
        self.min_char_conf = min_char_conf

    def analyze(self, ocr_results):
        texts = [r["text"] for r in ocr_results]
        full_text = " ".join(texts)
        count = len(ocr_results)
        confs = [r["confidence"] for r in ocr_results]
        avg_conf = sum(confs)/count if count else 0.0
        min_conf = min(confs) if confs else 0.0
        max_conf = max(confs) if confs else 0.0
        low_count = sum(1 for c in confs if c < self.min_char_conf)
        low_ratio = low_count/count if count else 0.0

        hit_kws, miss_kws = [], []
        for kw in self.keywords:
            (hit_kws if kw in full_text else miss_kws).append(kw)
        all_found = len(miss_kws) == 0
        hit_rate = len(hit_kws)/len(self.keywords)*100 if self.keywords else 100.0

        normal = re.compile(r'[\u4e00-\u9fffA-Za-z0-9\s\.\,\-\_\:\/\(\)\[\]\{\}，。！？、；：“”￥%℃#]+')
        normal_chunks = normal.findall(full_text)
        normal_len = sum(len(c) for c in normal_chunks)
        garbled = 1 - normal_len/max(len(full_text),1)

        total_area = sum(polygon_area(r["bbox"]) for r in ocr_results)

        return {
            "full_text": full_text,
            "ocr_count": count,
            "avg_confidence": round(avg_conf,4),
            "min_confidence": round(min_conf,4),
            "max_confidence": round(max_conf,4),
            "low_conf_count": low_count,
            "low_conf_ratio": round(low_ratio,4),
            "hit_keywords": hit_kws,
            "miss_keywords": miss_kws,
            "all_keywords_found": all_found,
            "hit_rate": round(hit_rate,2),
            "garbled_ratio": round(garbled,4),
            "text_area": round(total_area,2)
        }

# ===================== 质量判定（NG子代码） =====================
class QualityJudge:
    def __init__(self, min_avg_conf=0.8, max_garbled=0.2,
                 min_char_conf=0.5, max_low_ratio=0.3,
                 min_ocr_count=1, min_text_area=0):
        self.min_avg_conf = normalize_confidence(min_avg_conf)
        self.max_garbled = max_garbled
        self.min_char_conf = min_char_conf
        self.max_low_ratio = max_low_ratio
        self.min_ocr_count = min_ocr_count
        self.min_text_area = min_text_area

    def judge(self, analysis):
        reasons = []
        sub_codes = []

        if analysis["ocr_count"] == 0:
            reasons.append("未识别到文字")
            sub_codes.append("NG-01")
        else:
            if analysis["avg_confidence"] < self.min_avg_conf:
                reasons.append(f"平均置信度过低：{analysis['avg_confidence']:.2f} < {self.min_avg_conf:.2f}")
                sub_codes.append("NG-02")
            if analysis["miss_keywords"]:
                reasons.append(f"缺失关键字：{analysis['miss_keywords']}")
                sub_codes.append("NG-03")
            if analysis["garbled_ratio"] > self.max_garbled:
                reasons.append(f"乱码比例过高：{analysis['garbled_ratio']:.2f} > {self.max_garbled:.2f}")
                sub_codes.append("NG-04")
            if analysis["ocr_count"] < self.min_ocr_count:
                reasons.append(f"文字块数量过少：{analysis['ocr_count']} < {self.min_ocr_count}")
                sub_codes.append("NG-05")
            if self.min_text_area > 0 and analysis["text_area"] < self.min_text_area:
                reasons.append(f"文字区域过小：{analysis['text_area']:.0f} < {self.min_text_area}")
                sub_codes.append("NG-06")
            if analysis["low_conf_ratio"] > self.max_low_ratio:
                reasons.append(f"低置信度文字块比例过高：{analysis['low_conf_ratio']:.2f} > {self.max_low_ratio:.2f}")
                sub_codes.append("NG-08")

        status = "OK" if not reasons else "NG"
        return {"status": status, "reasons": reasons, "sub_codes": sub_codes}

# ===================== 报告输出（多Sheet + 美化） =====================
class ReportWriter:
    def __init__(self, output_dir):
        self.output_dir = output_dir
        self.report_dir = os.path.join(output_dir, "reports")
        self.text_dir = os.path.join(output_dir, "texts")
        self.annotated_dir = os.path.join(output_dir, "annotated_images")
        self.ng_dir = os.path.join(output_dir, "ng_images")
        for d in [self.output_dir, self.report_dir, self.text_dir, self.annotated_dir, self.ng_dir]:
            ensure_dir(d)

    def save_text(self, img_name, text):
        base = os.path.splitext(safe_filename(img_name))[0]
        with open(os.path.join(self.text_dir, f"{base}.txt"), "w", encoding="utf-8") as f:
            f.write(text)

    def save_annotated_image(self, img_name, img, ocr_results, status, orig_path=None):
        annotated = img.copy()
        for r in ocr_results:
            pts = np.array(r["bbox"], dtype=np.int32).reshape((-1,1,2))
            cv2.polylines(annotated, [pts], True, (0,255,0), 2)
            x, y = pts[0][0]
            cv2.putText(annotated, f"{r['confidence']:.2f}", (x, max(y-5,18)),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.55, (0,0,255), 1)
        out_path = os.path.join(self.annotated_dir, safe_filename(img_name))
        write_image_chinese_path(out_path, annotated)
        if status == "NG" and orig_path and os.path.exists(orig_path):
            try:
                shutil.copy2(orig_path, os.path.join(self.ng_dir, safe_filename(img_name)))
            except Exception as e:
                print(f"⚠️ NG图片保存失败：{e}")

    def generate_report(self, records: List[Dict]) -> str:
        wb = Workbook()
        # ---------- Sheet1: 检测总览 ----------
        ws1 = wb.active
        ws1.title = "检测总览"
        total = len(records)
        ok_count = sum(1 for r in records if r["status"] == "OK")
        ng_count = total - ok_count
        ok_rate = ok_count/total*100 if total else 0
        conf_list = [(r["avg_confidence"], r["image_name"]) for r in records if r["ocr_count"]>0]
        avg_conf = np.mean([c for c, _ in conf_list]) if conf_list else 0
        min_conf_img = min(conf_list, key=lambda x: x[0])[1] if conf_list else "无"
        unrecognized = sum(1 for r in records if r["ocr_count"]==0)
        kw_missing = sum(1 for r in records if r.get("miss_keywords") and len(r["miss_keywords"])>0)
        ng01 = sum(1 for r in records if "NG-01" in r.get("sub_codes",[]))
        ng02 = sum(1 for r in records if "NG-02" in r.get("sub_codes",[]))
        ng03 = sum(1 for r in records if "NG-03" in r.get("sub_codes",[]))
        ng04 = sum(1 for r in records if "NG-04" in r.get("sub_codes",[]))
        ng05 = sum(1 for r in records if "NG-05" in r.get("sub_codes",[]))
        ng06 = sum(1 for r in records if "NG-06" in r.get("sub_codes",[]))
        ng07 = sum(1 for r in records if "NG-07" in r.get("sub_codes",[]))
        ng08 = sum(1 for r in records if "NG-08" in r.get("sub_codes",[]))

        overview_data = [
            ["项目", "数值"],
            ["检测总图片数", total],
            ["OK数量", ok_count],
            ["NG数量", ng_count],
            ["OK率", f"{ok_rate:.1f}%"],
            ["平均OCR置信度", f"{avg_conf:.4f}"],
            ["最低置信度图片", min_conf_img],
            ["未识别图片数量", unrecognized],
            ["关键字缺失图片数量", kw_missing],
            ["NG-01 (无文字)", ng01],
            ["NG-02 (均置信度低)", ng02],
            ["NG-03 (关键字缺失)", ng03],
            ["NG-04 (疑似乱码)", ng04],
            ["NG-05 (文字块过少)", ng05],
            ["NG-06 (文字区域小)", ng06],
            ["NG-07 (读取失败)", ng07],
            ["NG-08 (疑似模糊)", ng08],
        ]
        for row in overview_data:
            ws1.append(row)
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        apply_header_style(ws1, 1)
        apply_data_style(ws1, 2, center_cols=[2])

        # ---------- Sheet2: OCR详细结果 ----------
        ws2 = wb.create_sheet("OCR详细结果")
        headers2 = ["序号","图片名称","预处理","识别文字","文字块数","平均置信度","最低置信度","最高置信度",
                    "低置信度块数","低置信度比例","文字区域面积","必检关键字","命中关键字","缺失关键字","命中率%",
                    "OK/NG","NG子代码","异常原因"]
        ws2.append(headers2)
        for i, r in enumerate(records, 1):
            ws2.append([i, r["image_name"], r.get("preprocess",""),
                        r["full_text"], r["ocr_count"], r["avg_confidence"], r["min_confidence"],
                        r["max_confidence"], r["low_conf_count"], r["low_conf_ratio"],
                        r.get("text_area",""), ", ".join(r.get("keywords",[])),
                        ", ".join(r.get("hit_keywords",[])), ", ".join(r.get("miss_keywords",[])),
                        r["hit_rate"], r["status"], ", ".join(r.get("sub_codes",[])),
                        "；".join(r["reasons"])])
        col_widths2 = [6,30,30,60,10,12,12,12,12,12,12,20,20,20,10,10,15,50]
        for idx, w in enumerate(col_widths2, 1):
            ws2.column_dimensions[chr(64+idx) if idx<=26 else None].width = w
        apply_header_style(ws2, 1)
        center_cols2 = [1,5,6,7,8,9,10,11,15,16,17]
        apply_data_style(ws2, 2, center_cols=center_cols2)

        # ---------- Sheet3: NG异常明细 ----------
        ws3 = wb.create_sheet("NG异常明细")
        ng_records = [r for r in records if r["status"]=="NG"]
        headers3 = ["序号","图片名称","NG子代码","异常原因","文字块数","平均置信度","最低置信度","命中关键字","缺失关键字"]
        ws3.append(headers3)
        for i, r in enumerate(ng_records, 1):
            ws3.append([i, r["image_name"], ", ".join(r.get("sub_codes",[])),
                        "；".join(r["reasons"]), r["ocr_count"], r["avg_confidence"],
                        r["min_confidence"], ", ".join(r.get("hit_keywords",[])),
                        ", ".join(r.get("miss_keywords",[]))])
        for idx, w in enumerate([6,30,20,60,10,12,12,25,25], 1):
            ws3.column_dimensions[chr(64+idx)].width = w
        apply_header_style(ws3, 1)
        apply_data_style(ws3, 2, center_cols=[1,5,6,7])

        # ---------- Sheet4: 关键字检测结果 ----------
        ws4 = wb.create_sheet("关键字检测结果")
        headers4 = ["序号","图片名称","必检关键字","命中关键字","缺失关键字","命中率%","OK/NG"]
        ws4.append(headers4)
        for i, r in enumerate(records, 1):
            ws4.append([i, r["image_name"], ", ".join(r.get("keywords",[])),
                        ", ".join(r.get("hit_keywords",[])), ", ".join(r.get("miss_keywords",[])),
                        r["hit_rate"], r["status"]])
        for idx, w in enumerate([6,30,30,30,30,10,10], 1):
            ws4.column_dimensions[chr(64+idx)].width = w
        apply_header_style(ws4, 1)
        apply_data_style(ws4, 2, center_cols=[1,6,7])

        # ---------- Sheet5: 统计汇总 ----------
        ws5 = wb.create_sheet("统计汇总")
        summary = [
            ["统计项", "数值"],
            ["总检测图片", total],
            ["OK数量", ok_count],
            ["NG数量", ng_count],
            ["OK率", f"{ok_rate:.1f}%"],
            ["平均OCR置信度", f"{avg_conf:.4f}"],
            ["最低置信度图片", min_conf_img],
            ["未识别图片数量", unrecognized],
            ["关键字缺失图片数量", kw_missing],
            ["NG-01 (无文字)", ng01],
            ["NG-02 (均置信度低)", ng02],
            ["NG-03 (关键字缺失)", ng03],
            ["NG-04 (疑似乱码)", ng04],
            ["NG-05 (文字块过少)", ng05],
            ["NG-06 (文字区域小)", ng06],
            ["NG-07 (读取失败)", ng07],
            ["NG-08 (疑似模糊)", ng08],
        ]
        for row in summary:
            ws5.append(row)
        ws5.column_dimensions['A'].width = 25
        ws5.column_dimensions['B'].width = 20
        apply_header_style(ws5, 1)
        apply_data_style(ws5, 2, center_cols=[2])

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = os.path.join(self.report_dir, f"OCR检测报告_{timestamp}.xlsx")
        wb.save(report_path)
        return report_path

# ===================== 交互输入（完整版） =====================
def interactive_input():
    print("\n🔍 OCR 图像识别工具 —— 交互模式")
    print("="*50)
    print("请选择识别模式：")
    print("  1. 单张图片")
    print("  2. 批量文件夹")
    print("  3. 视频抽帧")
    choice = input("请输入数字（1/2/3）：").strip()
    while choice not in ["1","2","3"]:
        choice = input("输入无效，请重新选择（1/2/3）：").strip()

    source = input("请输入图片 / 文件夹 / 视频路径：").strip().strip('"').strip("'")
    while not source:
        source = input("路径不能为空，请重新输入：").strip().strip('"').strip("'")
    source = os.path.abspath(source)

    if choice == "1":
        mode = "single"
        frame_interval = 30
    elif choice == "2":
        mode = "batch"
        frame_interval = 30
    else:
        mode = "video"
        val = input("请输入抽帧间隔（单位：帧，默认30）：").strip()
        frame_interval = int(val) if val.isdigit() else 30

    return source, mode, frame_interval

# ===================== 参数解析 =====================
def parse_args():
    p = argparse.ArgumentParser(description="OCR图像识别一体化工具（专业版）")
    p.add_argument("--source", type=str, default=None)
    p.add_argument("--mode", type=str, choices=["single","batch","video"])
    p.add_argument("--frame_interval", type=int, default=30)
    p.add_argument("--lang", type=str, default="ch")
    p.add_argument("--use_gpu", action="store_true")
    p.add_argument("--keywords", nargs="+", default=[])
    p.add_argument("--min_confidence", type=float, default=0.8)
    p.add_argument("--max_garbled", type=float, default=0.2)
    p.add_argument("--min_char_confidence", type=float, default=0.5)
    p.add_argument("--max_low_conf_ratio", type=float, default=0.3)
    p.add_argument("--min_ocr_count", type=int, default=1)
    p.add_argument("--min_text_area", type=float, default=0)
    p.add_argument("--output_dir", type=str, default=None)
    p.add_argument("--preprocess", type=str, default="")
    p.add_argument("--clahe_clip", type=float, default=2.0)
    p.add_argument("--denoise_method", type=str, default="gaussian")
    p.add_argument("--denoise_ksize", type=int, default=5)
    p.add_argument("--sharpen_strength", type=float, default=1.0)
    p.add_argument("--scale_factor", type=float, default=2.0)
    p.add_argument("--gamma", type=float, default=1.2)
    return p.parse_args()

# ===================== 主流程 =====================
def main():
    args = parse_args()
    if args.source:
        source = os.path.abspath(args.source.strip().strip('"').strip("'"))
        mode = args.mode if args.mode else infer_mode(source)
        frame_interval = args.frame_interval
    else:
        source, mode, frame_interval = interactive_input()

    if not os.path.exists(source):
        print(f"❌ 路径不存在：{source}")
        return

    output_dir = args.output_dir or (
        os.path.join(source, "ocr_output") if mode == "batch" else os.path.join(os.path.dirname(source) or ".", "ocr_output"))
    preprocess_list = [s.strip() for s in args.preprocess.split(",") if s.strip()] if args.preprocess else []
    preproc_kwargs = {
        "clahe_clip": args.clahe_clip,
        "denoise_method": args.denoise_method,
        "denoise_ksize": args.denoise_ksize,
        "sharpen_strength": args.sharpen_strength,
        "scale_factor": args.scale_factor,
        "gamma": args.gamma,
    }

    print("="*50)
    print(f"源：{source}  模式：{mode}  输出：{output_dir}")
    if preprocess_list: print(f"预处理：{preprocess_list}")
    print("="*50)

    input_module = InputModule(source, mode, frame_interval)
    ocr_engine = OCREngine(args.lang, args.use_gpu)
    analyzer = ResultAnalyzer(args.keywords, normalize_confidence(args.min_char_confidence))
    judge = QualityJudge(args.min_confidence, args.max_garbled,
                         normalize_confidence(args.min_char_confidence),
                         args.max_low_conf_ratio, args.min_ocr_count, args.min_text_area)
    writer = ReportWriter(output_dir)

    records = []
    total = ok = ng = 0

    for name, img, orig_path in input_module.get_data():
        total += 1
        if img is None:
            record = {
                "image_name": name,
                "preprocess": "None",
                "full_text": "",
                "ocr_count": 0,
                "avg_confidence": 0,
                "min_confidence": 0,
                "max_confidence": 0,
                "low_conf_count": 0,
                "low_conf_ratio": 0,
                "text_area": 0,
                "keywords": args.keywords,
                "hit_keywords": [],
                "miss_keywords": args.keywords[:],
                "hit_rate": 0,
                "status": "NG",
                "sub_codes": ["NG-07"],
                "reasons": ["图片读取失败"]
            }
            records.append(record)
            ng += 1
            print(f"[{total}] {name} ❌ 图片读取失败 (NG-07)")
            continue

        processed_img, preproc_desc = ImagePreprocessor.process(img, preprocess_list, **preproc_kwargs)
        ocr_results = ocr_engine.recognize(processed_img)
        analysis = analyzer.analyze(ocr_results)
        judgement = judge.judge(analysis)

        if judgement["status"] == "OK":
            ok += 1
        else:
            ng += 1

        record = {
            "image_name": name,
            "preprocess": preproc_desc,
            "full_text": analysis["full_text"],
            "ocr_count": analysis["ocr_count"],
            "avg_confidence": analysis["avg_confidence"],
            "min_confidence": analysis["min_confidence"],
            "max_confidence": analysis["max_confidence"],
            "low_conf_count": analysis["low_conf_count"],
            "low_conf_ratio": analysis["low_conf_ratio"],
            "text_area": analysis["text_area"],
            "keywords": args.keywords,
            "hit_keywords": analysis["hit_keywords"],
            "miss_keywords": analysis["miss_keywords"],
            "hit_rate": analysis["hit_rate"],
            "status": judgement["status"],
            "sub_codes": judgement["sub_codes"],
            "reasons": judgement["reasons"],
        }
        records.append(record)

        writer.save_text(name, analysis["full_text"])
        writer.save_annotated_image(name, processed_img, ocr_results, judgement["status"], orig_path)

        print(f"[{total}] {name} 预处理: {preproc_desc}")
        print(f"    状态: {judgement['status']}  子代码: {judgement['sub_codes']}")
        print(f"    文字块: {analysis['ocr_count']}  均置信度: {analysis['avg_confidence']:.3f}  最低置信度: {analysis['min_confidence']:.3f}")
        if judgement["reasons"]:
            print(f"    原因: {'；'.join(judgement['reasons'])}")
        print()

    if total == 0:
        print("⚠️ 没有图片被处理")
        return

    report_path = writer.generate_report(records)
    print("========== 检测完成 ==========")
    print(f"总图片: {total}  OK: {ok}  NG: {ng}  报告: {report_path}")
    print("==============================")

if __name__ == "__main__":
    main()