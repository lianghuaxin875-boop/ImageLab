# -*- coding: utf-8 -*-
"""
视频综合质量评分模块。

供 Web GUI 调用：码率、清晰度、曝光、稳定性、掉帧、黑/白屏、
画面冻结、音频基础检测，并生成多 Sheet Excel 报告、统计图表、
异常截图和视频预览图。
"""

import json
import os
import re
import shutil
import subprocess
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable, Dict, List, Optional

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from video_bitrate_tool import DEFAULT_FFPROBE_PATH, STANDARDS_CONFIG, VIDEO_EXTENSIONS, format_bps


SHARPNESS_THRESH = 80
EXP_UNDER_SEVERE = 30
EXP_UNDER = 60
EXP_OVER = 191
EXP_OVER_SEVERE = 225
BLACK_PIX_RATIO = 0.9
SUS_BLACK_RATIO = 0.8
WHITE_PIX_RATIO = 0.9
SUS_WHITE_RATIO = 0.8
FREEZE_LOW_DIFF = 2
STABILITY_OK_THRESH = 5
SILENCE_DBFS = -60
LOW_VOL_DBFS = -35
HIGH_VOL_DBFS = -3
PREVIEW_ROWS = 3
PREVIEW_COLS = 3
PREVIEW_SIZE = (320, 180)


def format_time(seconds):
    return str(timedelta(seconds=float(seconds)))[:10]


def smart_interval(duration):
    if duration <= 10:
        return 1.0
    if duration <= 60:
        return 2.0
    if duration <= 300:
        return 5.0
    return 10.0


def safe_sheet_name(name):
    return re.sub(r'[\\/*?:\[\]]', "_", name)[:31]


def imwrite_chinese(path, image):
    ext = Path(path).suffix or ".jpg"
    ok, buf = cv2.imencode(ext, image)
    if not ok:
        return False
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_bytes(buf.tobytes())
    return True


def score_from_judge(judge):
    text = str(judge or "")
    if text in ("OK", "优秀", "正常"):
        return 100
    if "警告" in text or "轻微" in text:
        return 60
    if "NG" in text or "严重" in text:
        return 0
    return 50


class VideoQualityScoringDetector:
    def __init__(
        self,
        input_dir: str,
        report_dir: str,
        mode: int,
        ffprobe_path: str = DEFAULT_FFPROBE_PATH,
        log_callback: Optional[Callable[[str], None]] = None,
        progress_callback: Optional[Callable[[int, int], None]] = None,
    ):
        self.input_dir = Path(input_dir)
        self.report_dir = Path(report_dir)
        self.mode = int(mode)
        self.ffprobe_path = ffprobe_path
        self.ffmpeg_path = (
            os.environ.get("FFMPEG_PATH")
            or shutil.which("ffmpeg")
            or str(Path(ffprobe_path).with_name("ffmpeg.exe"))
        )
        self.log = log_callback or (lambda message: None)
        self.progress = progress_callback or (lambda done, total: None)

    @property
    def mode_config(self):
        return STANDARDS_CONFIG[self.mode]

    def run_ffprobe(self, args, show_packets=False):
        if not Path(self.ffprobe_path).exists():
            raise FileNotFoundError(f"未找到 ffprobe：{self.ffprobe_path}")
        cmd = [self.ffprobe_path, "-v", "quiet", "-print_format", "json"]
        if show_packets:
            cmd += ["-select_streams", "v:0", "-show_packets", "-show_entries", "packet=pts_time"]
        else:
            cmd += ["-show_format", "-show_streams"]
        result = subprocess.run(cmd + list(args), capture_output=True, text=True, check=True)
        return json.loads(result.stdout or "{}")

    def get_duration(self, path):
        try:
            data = self.run_ffprobe([str(path)])
            duration = data.get("format", {}).get("duration")
            return float(duration) if duration else None
        except Exception:
            return None

    def extract_video_info(self, path):
        data = self.run_ffprobe([str(path)])
        fmt = data.get("format", {})
        total_bitrate = int(fmt["bit_rate"]) if fmt.get("bit_rate") else None
        video_bitrate = audio_bitrate = None
        width = height = fps = total_frames = None
        audio_info = None

        for stream in data.get("streams", []):
            codec_type = stream.get("codec_type")
            if codec_type == "video":
                width = width or stream.get("width")
                height = height or stream.get("height")
                if video_bitrate is None and stream.get("bit_rate"):
                    video_bitrate = int(stream["bit_rate"])
                if fps is None:
                    rate = stream.get("avg_frame_rate") or stream.get("r_frame_rate")
                    fps = self._parse_fps(rate)
                if total_frames is None:
                    if stream.get("nb_frames"):
                        total_frames = int(stream["nb_frames"])
                    elif stream.get("duration") and fps:
                        total_frames = int(float(stream["duration"]) * fps)
            elif codec_type == "audio" and audio_info is None:
                audio_bitrate = int(stream["bit_rate"]) if stream.get("bit_rate") else None
                audio_info = {
                    "codec": stream.get("codec_name", "N/A"),
                    "sample_rate": stream.get("sample_rate", "N/A"),
                    "channels": stream.get("channels", "N/A"),
                    "bitrate": audio_bitrate,
                }

        pixel_count = width * height if width and height else None
        threshold = self._bitrate_threshold(pixel_count, fps)
        judgement = "OK" if video_bitrate and threshold and video_bitrate >= threshold * 1_000_000 else "NG" if video_bitrate and threshold else "N/A"
        return {
            "filename": path.name,
            "total_bitrate": total_bitrate,
            "video_bitrate": video_bitrate,
            "audio_bitrate": audio_bitrate,
            "width": width,
            "height": height,
            "pixel_count": pixel_count,
            "fps": fps,
            "total_frames": total_frames,
            "audio_info": audio_info,
            "threshold_mbps": threshold,
            "bitrate_judge": judgement,
        }

    @staticmethod
    def _parse_fps(rate):
        if not rate:
            return None
        try:
            if "/" in rate:
                n, d = rate.split("/")
                d = float(d)
                return float(n) / d if d else None
            return float(rate)
        except Exception:
            return None

    def _bitrate_threshold(self, pixel_count, fps):
        if not pixel_count or not fps:
            return None
        for max_pixels, fps_limit, normal_min, high_min in self.mode_config["standards"]:
            if pixel_count <= max_pixels:
                return high_min if high_min and fps > fps_limit else normal_min
        return None

    def _sample_frames(self, path, interval_scale=1.0):
        cap = cv2.VideoCapture(str(path))
        if not cap.isOpened():
            return []
        duration = self.get_duration(path)
        if not duration or duration <= 0:
            fps = cap.get(cv2.CAP_PROP_FPS)
            total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            duration = total / fps if fps else 0
        interval = max(0.5, smart_interval(duration) * interval_scale)
        frames = []
        for second in np.arange(0, duration, interval):
            cap.set(cv2.CAP_PROP_POS_MSEC, float(second) * 1000)
            ok, frame = cap.read()
            if ok and frame is not None:
                frames.append((float(second), frame))
        cap.release()
        return frames

    def save_abnormal_frame(self, frame, video_name, timestamp, anomaly_type, root):
        safe_video = re.sub(r'[\\/*?:\[\]]', "_", video_name)
        path = Path(root) / safe_video / anomaly_type / f"{format_time(timestamp).replace(':', '-')}.jpg"
        return str(path) if imwrite_chinese(path, frame) else None

    def detect_sharpness(self, path, abnormal_root):
        vals = []
        shots = 0
        for second, frame in self._sample_frames(path):
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            lap = cv2.Laplacian(gray, cv2.CV_64F).var()
            vals.append(lap)
            if lap < SHARPNESS_THRESH:
                shots += 1
                self.save_abnormal_frame(frame, path.name, second, "模糊", abnormal_root)
        if not vals:
            return None
        arr = np.array(vals)
        blurry_count = int(np.sum(arr < SHARPNESS_THRESH))
        blurry_ratio = blurry_count / len(arr) * 100
        avg = float(np.mean(arr))
        judge = "OK" if avg >= 150 and blurry_ratio <= 10 else "警告" if avg >= 80 or blurry_ratio <= 30 else "NG"
        return {
            "avg": round(avg, 1),
            "min": round(float(np.min(arr)), 1),
            "max": round(float(np.max(arr)), 1),
            "blurry_count": blurry_count,
            "total": len(arr),
            "blurry_ratio": round(blurry_ratio, 1),
            "judge": judge,
            "abnormal_screenshots": shots,
        }

    def detect_exposure(self, path, abnormal_root):
        values = []
        counts = {"under_severe": 0, "under": 0, "normal": 0, "over": 0, "over_severe": 0}
        shots = 0
        for second, frame in self._sample_frames(path):
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            mean = float(np.mean(gray))
            values.append(mean)
            anomaly = None
            if mean <= EXP_UNDER_SEVERE:
                counts["under_severe"] += 1
                anomaly = "严重欠曝"
            elif mean <= EXP_UNDER:
                counts["under"] += 1
                anomaly = "欠曝"
            elif mean <= EXP_OVER:
                counts["normal"] += 1
            elif mean <= EXP_OVER_SEVERE:
                counts["over"] += 1
                anomaly = "过曝"
            else:
                counts["over_severe"] += 1
                anomaly = "严重过曝"
            if anomaly:
                shots += 1
                self.save_abnormal_frame(frame, path.name, second, anomaly, abnormal_root)
        if not values:
            return None
        total = len(values)
        avg = float(np.mean(values))
        under_ratio = (counts["under_severe"] + counts["under"]) / total * 100
        over_ratio = (counts["over"] + counts["over_severe"]) / total * 100
        normal_ratio = counts["normal"] / total * 100
        judge = "OK" if 61 <= avg <= 190 and under_ratio + over_ratio <= 10 else "警告" if 31 <= avg <= 225 and under_ratio + over_ratio <= 30 else "NG"
        return {
            "avg_b": round(avg, 1),
            "min_b": round(min(values), 1),
            "max_b": round(max(values), 1),
            "under_ratio": round(under_ratio, 1),
            "over_ratio": round(over_ratio, 1),
            "normal_ratio": round(normal_ratio, 1),
            "judge": judge,
            "abnormal_screenshots": shots,
        }

    def detect_stability(self, path):
        frames = self._sample_frames(path, interval_scale=0.5)
        if len(frames) < 2:
            return None
        prev = None
        disps = []
        for _second, frame in frames:
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            if prev is not None:
                flow = cv2.calcOpticalFlowFarneback(prev, gray, None, 0.5, 3, 15, 3, 5, 1.2, 0)
                mag, _ = cv2.cartToPolar(flow[..., 0], flow[..., 1])
                disps.append(float(np.mean(mag)))
            prev = gray
        if not disps:
            return None
        arr = np.array(disps)
        shake_ratio = float(np.sum(arr > STABILITY_OK_THRESH) / len(arr) * 100)
        score = int(min(100, max(0, 100 - (np.std(arr) * 10 + shake_ratio))))
        judge = "优秀" if score >= 90 else "OK" if score >= 75 else "轻微抖动" if score >= 60 else "NG：抖动明显"
        return {"avg_disp": round(float(np.mean(arr)), 2), "shake_ratio": round(shake_ratio, 1), "score": score, "judge": judge}

    def detect_frame_drops(self, path):
        try:
            data = self.run_ffprobe([str(path)], show_packets=True)
            pts = [float(p["pts_time"]) for p in data.get("packets", []) if p.get("pts_time")]
            if len(pts) < 2:
                return None
            intervals = np.diff(pts)
            median = np.median(intervals)
            nominal_fps = 1 / median if median else None
            actual_fps = len(pts) / (pts[-1] - pts[0]) if pts[-1] > pts[0] else None
            threshold_interval = 1 / round(nominal_fps) if nominal_fps else None
            abnormal = (intervals > threshold_interval * 1.5) | (intervals < threshold_interval * 0.5) if threshold_interval else []
            abnormal_count = int(np.sum(abnormal))
            abnormal_ratio = abnormal_count / len(intervals) * 100
            diff_pct = abs(nominal_fps - actual_fps) / nominal_fps * 100 if nominal_fps and actual_fps else None
            judge = "OK" if diff_pct is not None and diff_pct <= 3 else "警告：帧率略不稳定" if diff_pct is not None and diff_pct <= 8 else "NG：疑似掉帧"
            return {
                "nom_fps": round(nominal_fps, 2) if nominal_fps else "N/A",
                "act_fps": round(actual_fps, 2) if actual_fps else "N/A",
                "diff_pct": round(diff_pct, 1) if diff_pct is not None else "N/A",
                "max_interval_ms": round(float(np.max(intervals)) * 1000, 1),
                "abnormal_count": abnormal_count,
                "abnormal_ratio": round(abnormal_ratio, 1),
                "judge": judge,
            }
        except Exception:
            return None

    def detect_black_white(self, path, abnormal_root):
        frames = self._sample_frames(path)
        if not frames:
            return None
        black_count = white_count = shots = 0
        black_segments = []
        white_segments = []
        black_start = white_start = None
        for second, frame in frames:
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            mean = float(np.mean(gray))
            black_pix = float(np.sum(gray < 15) / gray.size)
            white_pix = float(np.sum(gray > 240) / gray.size)
            is_black = (mean < 15 and black_pix > BLACK_PIX_RATIO) or (mean < 30 and black_pix > SUS_BLACK_RATIO)
            is_white = (mean > 240 and white_pix > WHITE_PIX_RATIO) or (mean > 225 and white_pix > SUS_WHITE_RATIO)
            if is_black:
                black_count += 1
                shots += 1
                self.save_abnormal_frame(frame, path.name, second, "黑屏", abnormal_root)
                black_start = second if black_start is None else black_start
            elif black_start is not None:
                black_segments.append((black_start, second))
                black_start = None
            if is_white:
                white_count += 1
                shots += 1
                self.save_abnormal_frame(frame, path.name, second, "白屏", abnormal_root)
                white_start = second if white_start is None else white_start
            elif white_start is not None:
                white_segments.append((white_start, second))
                white_start = None
        last_second = frames[-1][0]
        if black_start is not None:
            black_segments.append((black_start, last_second))
        if white_start is not None:
            white_segments.append((white_start, last_second))
        longest_black = max([e - s for s, e in black_segments], default=0)
        longest_white = max([e - s for s, e in white_segments], default=0)
        issues = []
        if longest_black >= 3:
            issues.append("NG：黑屏≥3秒")
        elif longest_black >= 1:
            issues.append("警告：存在短暂黑屏")
        if longest_white >= 3:
            issues.append("NG：白屏≥3秒")
        elif longest_white >= 1:
            issues.append("警告：存在短暂白屏")
        return {
            "b_segs": [f"{format_time(s)}-{format_time(e)}" for s, e in black_segments[:3]],
            "w_segs": [f"{format_time(s)}-{format_time(e)}" for s, e in white_segments[:3]],
            "long_b": round(longest_black, 2),
            "long_w": round(longest_white, 2),
            "b_ratio": round(black_count / len(frames) * 100, 1),
            "w_ratio": round(white_count / len(frames) * 100, 1),
            "judge": "；".join(issues) if issues else "正常",
            "abnormal_screenshots": shots,
        }

    def detect_frozen(self, path, abnormal_root):
        frames = self._sample_frames(path)
        if len(frames) < 2:
            return None
        prev = None
        start = None
        segments = []
        shots = 0
        for second, frame in frames:
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            if prev is not None:
                mean_diff = float(np.mean(cv2.absdiff(prev, gray)))
                if mean_diff < FREEZE_LOW_DIFF:
                    if start is None:
                        start = second
                        shots += 1
                        self.save_abnormal_frame(frame, path.name, second, "冻结", abnormal_root)
                elif start is not None:
                    segments.append((start, second, second - start))
                    start = None
            prev = gray
        if start is not None:
            end = frames[-1][0]
            segments.append((start, end, end - start))
        if not segments:
            return {"segs": [], "longest": 0, "count": 0, "total_dur": 0, "ratio": 0, "judge": "正常", "abnormal_screenshots": 0}
        duration = self.get_duration(path) or max(frames[-1][0], 1)
        longest = max(s[2] for s in segments)
        total_dur = sum(s[2] for s in segments)
        judge = "严重 NG" if longest > 3 else "NG：存在明显画面冻结" if longest > 1 else "警告" if longest > 0.5 else "正常"
        return {
            "segs": [f"{format_time(s)}-{format_time(e)}" for s, e, _ in segments],
            "longest": round(longest, 2),
            "count": len(segments),
            "total_dur": round(total_dur, 2),
            "ratio": round(total_dur / duration * 100, 1) if duration else 0,
            "judge": judge,
            "abnormal_screenshots": shots,
        }

    def analyze_audio(self, path, audio_info):
        if not audio_info:
            return {"has_audio": False, "judge": "无音频"}
        mean_db = None
        if Path(self.ffmpeg_path).exists():
            try:
                cmd = [self.ffmpeg_path, "-i", str(path), "-af", "volumedetect", "-f", "null", "-"]
                result = subprocess.run(cmd, capture_output=True, text=True, stderr=subprocess.STDOUT, timeout=120)
                for line in (result.stdout + result.stderr).splitlines():
                    if "mean_volume:" in line:
                        mean_db = float(line.split("mean_volume:")[1].strip().split()[0])
                        break
            except Exception:
                mean_db = None
        issues = []
        bitrate = audio_info.get("bitrate")
        if bitrate:
            if bitrate < 64000:
                issues.append("NG：音频码率过低")
            elif bitrate < 96000:
                issues.append("警告：音质可能偏低")
        sample_rate = str(audio_info.get("sample_rate", ""))
        if sample_rate.isdigit():
            sr = int(sample_rate)
            if sr < 22050:
                issues.append("NG：采样率过低")
            elif sr < 32000:
                issues.append("警告：采样率略低")
        if mean_db is not None:
            if mean_db < SILENCE_DBFS:
                issues.append("静音/接近无声")
            elif mean_db < LOW_VOL_DBFS:
                issues.append("音量偏低")
            elif mean_db > HIGH_VOL_DBFS:
                issues.append("可能爆音/削波")
        return {
            "has_audio": True,
            "codec": audio_info.get("codec", "N/A"),
            "sample_rate": sample_rate or "N/A",
            "channels": audio_info.get("channels", "N/A"),
            "bitrate_kbps": round(bitrate / 1000, 1) if bitrate else "N/A",
            "mean_volume_db": round(mean_db, 1) if mean_db is not None else "N/A",
            "silence_segs": [],
            "judge": "；".join(issues) if issues else "OK",
        }

    def generate_preview(self, path, output_path):
        cap = cv2.VideoCapture(str(path))
        if not cap.isOpened():
            return None
        total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = cap.get(cv2.CAP_PROP_FPS)
        duration = total_frames / fps if fps else 0
        if duration <= 0:
            cap.release()
            return None
        times = [duration / (PREVIEW_ROWS * PREVIEW_COLS + 1) * (i + 1) for i in range(PREVIEW_ROWS * PREVIEW_COLS)]
        thumbs = []
        for second in times:
            cap.set(cv2.CAP_PROP_POS_MSEC, second * 1000)
            ok, frame = cap.read()
            if ok and frame is not None:
                thumbs.append(cv2.resize(frame, PREVIEW_SIZE))
            else:
                thumbs.append(np.zeros((PREVIEW_SIZE[1], PREVIEW_SIZE[0], 3), dtype=np.uint8))
        cap.release()
        preview = np.vstack([np.hstack(thumbs[r * PREVIEW_COLS:(r + 1) * PREVIEW_COLS]) for r in range(PREVIEW_ROWS)])
        return str(output_path) if imwrite_chinese(output_path, preview) else None

    def compute_scores(self, bitrate_judge, quality):
        scores = {
            "码率": score_from_judge(bitrate_judge),
            "清晰度": score_from_judge((quality.get("sharpness") or {}).get("judge")),
            "曝光": score_from_judge((quality.get("exposure") or {}).get("judge")),
            "稳定性": (quality.get("stability") or {}).get("score", 50),
            "掉帧": score_from_judge((quality.get("frame_drops") or {}).get("judge")),
            "音频": 100 if not (quality.get("audio") or {}).get("has_audio") else score_from_judge((quality.get("audio") or {}).get("judge")),
        }
        bw_judge = (quality.get("black_white") or {}).get("judge", "正常")
        frozen_judge = (quality.get("frozen") or {}).get("judge", "正常")
        scores["黑屏/白屏/冻结"] = 0 if "NG" in bw_judge or "NG" in frozen_judge else 60 if "警告" in bw_judge or "警告" in frozen_judge else 100
        weights = {"码率": 0.20, "清晰度": 0.25, "曝光": 0.15, "稳定性": 0.15, "掉帧": 0.10, "黑屏/白屏/冻结": 0.10, "音频": 0.05}
        total = sum(scores[key] * weights[key] for key in weights)
        scores["总分"] = round(total, 1)
        scores["等级"] = "优秀" if total >= 90 else "良好" if total >= 75 else "一般" if total >= 60 else "较差"
        return scores

    def run(self):
        if self.mode not in STANDARDS_CONFIG:
            raise ValueError("视频综合检测模式无效。")
        if not self.input_dir.is_dir():
            raise FileNotFoundError(f"视频输入目录不存在：{self.input_dir}")
        self.report_dir.mkdir(parents=True, exist_ok=True)
        videos = [p for p in self.input_dir.iterdir() if p.is_file() and p.suffix.lower() in VIDEO_EXTENSIONS]
        if not videos:
            raise RuntimeError(f"在 {self.input_dir} 中未找到支持的视频文件。")
        abnormal_root = self.report_dir / "异常截图"
        preview_dir = self.report_dir / "预览图"
        preview_dir.mkdir(parents=True, exist_ok=True)

        results = []
        quality_data = {}
        scores = {}
        preview_paths = {}
        for index, video_path in enumerate(videos, 1):
            self.log(f"综合分析视频：{video_path.name}")
            try:
                info = self.extract_video_info(video_path)
                quality = {
                    "sharpness": self.detect_sharpness(video_path, abnormal_root),
                    "exposure": self.detect_exposure(video_path, abnormal_root),
                    "stability": self.detect_stability(video_path),
                    "frame_drops": self.detect_frame_drops(video_path),
                    "black_white": self.detect_black_white(video_path, abnormal_root),
                    "frozen": self.detect_frozen(video_path, abnormal_root),
                    "audio": self.analyze_audio(video_path, info.get("audio_info")),
                }
                preview_path = self.generate_preview(video_path, preview_dir / f"{video_path.stem}_preview.jpg")
                if preview_path:
                    preview_paths[video_path.name] = preview_path
                score = self.compute_scores(info["bitrate_judge"], quality)
                quality_data[video_path.name] = quality
                scores[video_path.name] = score
                results.append(self._result_record(video_path.name, info, quality, score, preview_path))
            except Exception as exc:
                self.log(f"综合分析失败：{video_path.name}，原因：{exc}")
                results.append({"文件名": video_path.name, "综合等级": "ERROR", "综合评分": 0, "错误信息": str(exc)})
            self.progress(index, len(videos))
        return {"results": results, "quality_data": quality_data, "scores": scores, "preview_paths": preview_paths, "mode_config": self.mode_config}

    def _result_record(self, name, info, quality, score, preview_path):
        return {
            "文件名": name,
            "分辨率": f"{info.get('width')}x{info.get('height')}" if info.get("width") else "N/A",
            "总码率": format_bps(info.get("total_bitrate")),
            "视频码率": format_bps(info.get("video_bitrate")),
            "音频码率": format_bps(info.get("audio_bitrate")),
            "帧率": round(info["fps"], 2) if info.get("fps") else "",
            "总帧数": info.get("total_frames") or "",
            "阈值Mbps": info.get("threshold_mbps") or "",
            "码率判断": info.get("bitrate_judge", ""),
            "清晰度判断": (quality.get("sharpness") or {}).get("judge", "N/A"),
            "曝光判断": (quality.get("exposure") or {}).get("judge", "N/A"),
            "稳定性判断": (quality.get("stability") or {}).get("judge", "N/A"),
            "掉帧判断": (quality.get("frame_drops") or {}).get("judge", "N/A"),
            "黑白屏判断": (quality.get("black_white") or {}).get("judge", "N/A"),
            "冻结判断": (quality.get("frozen") or {}).get("judge", "N/A"),
            "音频判断": (quality.get("audio") or {}).get("judge", "N/A"),
            "综合评分": score["总分"],
            "综合等级": score["等级"],
            "预览图": preview_path or "",
        }

    @staticmethod
    def _display_width(value):
        return sum(2 if ord(ch) > 255 else 1 for ch in str(value or ""))

    def _style_sheet(self, ws, header_row=1, max_col=None):
        border = Border(left=Side("thin", color="D9E2F3"), right=Side("thin", color="D9E2F3"), top=Side("thin", color="D9E2F3"), bottom=Side("thin", color="D9E2F3"))
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        max_col = max_col or ws.max_column
        for col_idx in range(1, max_col + 1):
            letter = get_column_letter(col_idx)
            max_width = max(self._display_width(ws.cell(row, col_idx).value) for row in range(1, ws.max_row + 1))
            ws.column_dimensions[letter].width = min(max(max_width + 2, 12), 42)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"A{header_row + 1}"

    def write_report(self, payload):
        self.report_dir.mkdir(parents=True, exist_ok=True)
        report_path = self.report_dir / f"video_scoring_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb = Workbook()
        self._add_summary_sheet(wb, payload["results"])
        self._add_standard_sheet(wb)
        self._add_quality_sheets(wb, payload["quality_data"])
        self._add_scoring_standard_sheet(wb)
        self._add_scores_sheet(wb, payload["scores"])
        self._add_preview_sheet(wb, payload["preview_paths"])
        self._add_charts_sheet(wb, payload["results"])
        wb.save(report_path)
        return report_path

    def _add_summary_sheet(self, wb, results):
        ws = wb.active
        ws.title = "Summary"
        headers = ["文件名", "分辨率", "视频码率", "帧率", "码率判断", "清晰度", "曝光", "稳定性", "掉帧", "黑白屏", "冻结", "音频", "综合评分", "等级"]
        ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=len(headers))
        total = len(results)
        ok = sum(1 for r in results if r.get("综合等级") in ("优秀", "良好"))
        ng = total - ok
        ws["A1"] = f"本次共检测 {total} 个视频。综合结果 OK {ok} 个，NG/需关注 {ng} 个，NG/需关注占比 {(ng / total * 100 if total else 0):.1f}%。"
        ws["A1"].font = Font(bold=True, size=12, color="1F4E78")
        ws["A1"].fill = PatternFill(start_color="EAF4FF", end_color="EAF4FF", fill_type="solid")
        for idx, header in enumerate(headers, 1):
            ws.cell(5, idx, header)
        for row_idx, record in enumerate(results, 6):
            values = [
                record.get("文件名"), record.get("分辨率"), record.get("视频码率"), record.get("帧率"), record.get("码率判断"),
                record.get("清晰度判断"), record.get("曝光判断"), record.get("稳定性判断"), record.get("掉帧判断"),
                record.get("黑白屏判断"), record.get("冻结判断"), record.get("音频判断"), record.get("综合评分"), record.get("综合等级"),
            ]
            for col_idx, value in enumerate(values, 1):
                ws.cell(row_idx, col_idx, value)
        self._style_sheet(ws, header_row=5)

    def _add_standard_sheet(self, wb):
        ws = wb.create_sheet("码率评判标准")
        ws.merge_cells("A1:E1")
        ws["A1"] = f"评判标准 - {self.mode_config['name']}"
        headers = ["拍摄模式/分辨率", "分辨率", "帧率", "建议码率", "检测判断"]
        for idx, header in enumerate(headers, 1):
            ws.cell(3, idx, header)
        for row_idx, row in enumerate(self.mode_config["table"], 4):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row_idx, col_idx, value)
        self._style_sheet(ws, header_row=3)

    def _add_quality_sheets(self, wb, quality_data):
        sheet_defs = [
            ("清晰度检测", ["视频文件", "平均清晰度", "最低清晰度", "最高清晰度", "模糊帧数量", "抽帧总数", "模糊帧占比(%)", "判断", "异常截图数"], "sharpness", ["avg", "min", "max", "blurry_count", "total", "blurry_ratio", "judge", "abnormal_screenshots"]),
            ("曝光检测", ["视频文件", "平均亮度", "最低亮度", "最高亮度", "欠曝帧占比(%)", "过曝帧占比(%)", "正常帧占比(%)", "判断", "异常截图数"], "exposure", ["avg_b", "min_b", "max_b", "under_ratio", "over_ratio", "normal_ratio", "judge", "abnormal_screenshots"]),
            ("稳定性检测", ["视频文件", "平均位移量(px)", "抖动帧占比(%)", "稳定性评分", "判断"], "stability", ["avg_disp", "shake_ratio", "score", "judge"]),
            ("掉帧检测", ["视频文件", "标称fps", "实际fps", "帧率差值(%)", "最大帧间隔(ms)", "异常帧数量", "异常帧占比(%)", "判断"], "frame_drops", ["nom_fps", "act_fps", "diff_pct", "max_interval_ms", "abnormal_count", "abnormal_ratio", "judge"]),
            ("黑屏白屏检测", ["视频文件", "黑屏片段", "白屏片段", "最长黑屏(s)", "最长白屏(s)", "黑屏帧占比(%)", "白屏帧占比(%)", "判断", "异常截图数"], "black_white", ["b_segs", "w_segs", "long_b", "long_w", "b_ratio", "w_ratio", "judge", "abnormal_screenshots"]),
            ("画面冻结检测", ["视频文件", "冻结片段", "最长冻结(s)", "冻结次数", "冻结总时长(s)", "冻结占比(%)", "判断", "异常截图数"], "frozen", ["segs", "longest", "count", "total_dur", "ratio", "judge", "abnormal_screenshots"]),
            ("音频基础检测", ["视频文件", "是否有音频", "编码", "采样率(Hz)", "声道", "码率(kbps)", "平均音量(dBFS)", "静音片段", "判断"], "audio", ["has_audio", "codec", "sample_rate", "channels", "bitrate_kbps", "mean_volume_db", "silence_segs", "judge"]),
        ]
        for sheet_name, headers, key, fields in sheet_defs:
            ws = wb.create_sheet(sheet_name)
            for col_idx, header in enumerate(headers, 1):
                ws.cell(1, col_idx, header)
            for row_idx, (video_name, data) in enumerate(quality_data.items(), 2):
                d = data.get(key) or {}
                ws.cell(row_idx, 1, video_name)
                for col_idx, field in enumerate(fields, 2):
                    value = d.get(field, "N/A")
                    if isinstance(value, list):
                        value = ", ".join(value) or "无"
                    if field == "has_audio":
                        value = "有" if value else "无"
                    ws.cell(row_idx, col_idx, value)
            self._style_sheet(ws)

    def _add_scoring_standard_sheet(self, wb):
        ws = wb.create_sheet("评分标准")
        rows = [
            ["模块", "权重"], ["码率/分辨率/帧率基础参数", "20%"], ["视频清晰度", "25%"], ["视频曝光", "15%"],
            ["视频稳定性", "15%"], ["卡顿/掉帧", "10%"], ["黑屏/白屏/冻结", "10%"], ["音频基础检测", "5%"],
            [], ["综合评分", "等级", "判断"], ["90-100", "优秀", "OK"], ["75-89", "良好", "OK"],
            ["60-74", "一般", "警告"], ["0-59", "较差", "NG"],
        ]
        for row in rows:
            ws.append(row)
        self._style_sheet(ws)

    def _add_scores_sheet(self, wb, scores):
        ws = wb.create_sheet("综合评分")
        headers = ["视频文件", "码率(20%)", "清晰度(25%)", "曝光(15%)", "稳定性(15%)", "掉帧(10%)", "黑屏/白屏/冻结(10%)", "音频(5%)", "加权总分", "等级"]
        ws.append(headers)
        for video_name, score in scores.items():
            ws.append([video_name, score["码率"], score["清晰度"], score["曝光"], score["稳定性"], score["掉帧"], score["黑屏/白屏/冻结"], score["音频"], score["总分"], score["等级"]])
        self._style_sheet(ws)

    def _add_preview_sheet(self, wb, preview_paths):
        ws = wb.create_sheet("视频预览图")
        ws["A1"] = "视频文件"
        ws["B1"] = "预览图"
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 44
        for row_idx, (video_name, path) in enumerate(preview_paths.items(), 2):
            ws.cell(row_idx, 1, video_name)
            if path and Path(path).exists():
                try:
                    img = XLImage(path)
                    img.width = 300
                    img.height = 170
                    ws.add_image(img, f"B{row_idx}")
                    ws.row_dimensions[row_idx].height = 132
                except Exception as exc:
                    ws.cell(row_idx, 2, f"预览图加载失败：{exc}")
        self._style_sheet(ws)

    def _add_charts_sheet(self, wb, results):
        ws = wb.create_sheet("统计图表")
        ws.append(["视频文件", "综合评分"])
        for record in results:
            ws.append([record.get("文件名"), record.get("综合评分", 0)])
        level_counts = {}
        for record in results:
            level = record.get("综合等级", "未知")
            level_counts[level] = level_counts.get(level, 0) + 1
        start = max(ws.max_row + 3, 6)
        ws.cell(start, 4, "等级")
        ws.cell(start, 5, "数量")
        for idx, (level, count) in enumerate(level_counts.items(), start + 1):
            ws.cell(idx, 4, level)
            ws.cell(idx, 5, count)
        if len(results) > 0:
            bar = BarChart()
            bar.title = "综合评分分布"
            bar.y_axis.title = "分数"
            data = Reference(ws, min_col=2, min_row=1, max_row=len(results) + 1)
            cats = Reference(ws, min_col=1, min_row=2, max_row=len(results) + 1)
            bar.add_data(data, titles_from_data=True)
            bar.set_categories(cats)
            ws.add_chart(bar, "G2")

            pie = PieChart()
            pie.title = "综合等级占比"
            pie_data = Reference(ws, min_col=5, min_row=start, max_row=start + len(level_counts))
            pie_cats = Reference(ws, min_col=4, min_row=start + 1, max_row=start + len(level_counts))
            pie.add_data(pie_data, titles_from_data=True)
            pie.set_categories(pie_cats)
            pie.dataLabels = DataLabelList()
            pie.dataLabels.showPercent = True
            ws.add_chart(pie, "G18")
        self._style_sheet(ws)
