#!/usr/bin/env python3
"""
视频码率提取工具（批量处理文件夹 → 单 Excel 报告）
用于图像质量检测：提取每个视频的总码率、视频流码率、音频流码率、帧率、分辨率、总帧数，以及可选逐帧数据。
所有结果汇总到一个 Excel 文件中，输出到视频文件夹。
依赖 FFmpeg（ffprobe）和 openpyxl。

使用方法：
1. 确保 FFPROBE_PATH 指向你的 ffprobe.exe 实际位置。
2. 修改下方配置区域（视频文件夹等）。
3. 运行脚本：python "Bitrate Extraction.py"
"""

import subprocess
import json
import sys
import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple

# 检查 openpyxl 依赖
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl 库。请执行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# ==================== 配置区域 ====================
FFPROBE_PATH = r"D:\码率工具FFepeg\ffmpeg-8.1.1-full_build\bin\ffprobe.exe"

INPUT_FOLDER = r"D:\Codex\码率测试\input"                     # 包含视频的文件夹路径
OUTPUT_EXCEL_NAME = "bitrate_report"         # 报告基础名（不含扩展名，会自动追加时间戳，避免覆盖）
INCLUDE_FRAMES = True                        # 是否包含逐帧大小数据
# 支持的视频文件扩展名（小写）
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".mpg", ".mpeg"}
# =================================================


def run_ffprobe(args: List[str]) -> dict:
    """调用 ffprobe 并返回 JSON 解析结果"""
    cmd = [FFPROBE_PATH, "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams"] + args
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return json.loads(result.stdout)
    except subprocess.CalledProcessError as e:
        print(f"ffprobe 执行失败: {e.stderr}", file=sys.stderr)
        return {}
    except FileNotFoundError:
        print(f"错误: 未找到 ffprobe，请检查 FFPROBE_PATH 是否正确。当前路径：{FFPROBE_PATH}", file=sys.stderr)
        sys.exit(1)


def extract_bitrates_and_info(file_path: str) -> Dict:
    """提取码率、帧率、分辨率、总帧数等"""
    info = run_ffprobe([file_path])
    if not info:
        return {}

    fmt = info.get("format", {})
    total_bitrate = fmt.get("bit_rate")
    if total_bitrate is not None:
        total_bitrate = int(total_bitrate)

    video_bitrate = None
    audio_bitrate = None
    fps = None
    video_total_frames = None
    width = None
    height = None

    for stream in info.get("streams", []):
        codec_type = stream.get("codec_type")
        if codec_type == "video":
            if video_bitrate is None:
                if "bit_rate" in stream:
                    video_bitrate = int(stream["bit_rate"])
                else:
                    duration = float(stream.get("duration", 0))
                    size = stream.get("size")
                    if duration > 0 and size:
                        video_bitrate = int(float(size) * 8 / duration)
            if fps is None:
                r_frame_rate = stream.get("r_frame_rate")
                if r_frame_rate:
                    try:
                        if '/' in r_frame_rate:
                            num, den = r_frame_rate.split('/')
                            fps = float(num) / float(den)
                        else:
                            fps = float(r_frame_rate)
                    except (ValueError, ZeroDivisionError):
                        fps = None
            if video_total_frames is None:
                nb_frames = stream.get("nb_frames")
                if nb_frames is not None:
                    video_total_frames = int(nb_frames)
                else:
                    duration = float(stream.get("duration", 0))
                    if duration > 0 and fps is not None:
                        video_total_frames = int(duration * fps)
            # 提取分辨率
            if width is None or height is None:
                w = stream.get("width")
                h = stream.get("height")
                if w is not None and h is not None:
                    width = int(w)
                    height = int(h)
        elif codec_type == "audio" and audio_bitrate is None:
            if "bit_rate" in stream:
                audio_bitrate = int(stream["bit_rate"])
            else:
                duration = float(stream.get("duration", 0))
                size = stream.get("size")
                if duration > 0 and size:
                    audio_bitrate = int(float(size) * 8 / duration)

    return {
        "total_bitrate_bps": total_bitrate,
        "video_bitrate_bps": video_bitrate,
        "audio_bitrate_bps": audio_bitrate,
        "fps": fps,
        "width": width,
        "height": height,
        "video_total_frames": video_total_frames,
    }


def get_frame_sizes(file_path: str, stream_index: int = 0) -> List[int]:
    """获取视频流中每一帧的 packet 大小（字节）"""
    cmd = [
        FFPROBE_PATH,
        "-v", "quiet",
        "-select_streams", f"v:{stream_index}",
        "-show_packets",
        "-show_entries", "packet=size",
        "-of", "json",
        file_path
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        data = json.loads(result.stdout)
        packets = data.get("packets", [])
        return [int(pkt["size"]) for pkt in packets]
    except subprocess.CalledProcessError as e:
        print(f"获取帧大小失败: {e.stderr}", file=sys.stderr)
        return []
    except FileNotFoundError:
        print(f"错误: 未找到 ffprobe。", file=sys.stderr)
        return []


def format_bps(bps: Optional[int]) -> str:
    """将 bps 转换为可读字符串"""
    if bps is None:
        return "N/A"
    if bps >= 1_000_000:
        return f"{bps / 1_000_000:.2f} Mbps"
    elif bps >= 1_000:
        return f"{bps / 1_000:.2f} kbps"
    else:
        return f"{bps} bps"


def sanitize_sheet_name(name: str) -> str:
    """清理工作表名称"""
    name = re.sub(r'[\\/*?:\[\]]', '_', name)
    return name[:31]


def apply_wrap_and_width(ws, headers_count: int):
    """给所有数据单元格设置自动换行，并根据内容调整列宽"""
    wrap_alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=headers_count):
        for cell in row:
            cell.alignment = wrap_alignment

    for col_idx in range(1, headers_count + 1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        length = 0
                        for char in line:
                            if ord(char) > 127:
                                length += 2
                            else:
                                length += 1
                        max_length = max(max_length, length)
        ws.column_dimensions[col_letter].width = max_length + 2


def create_summary_sheet(wb: Workbook,
                         all_results: List[Tuple[str, Dict, Optional[List[int]]]],
                         include_frames: bool):
    """创建 Summary 工作表"""
    ws = wb.active
    ws.title = "Summary"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # 标题行（合并单元格，扩展至 K 列）
    ws.merge_cells("A1:K1")
    ws["A1"] = "视频码率分析报告"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:K2")
    ws["A2"] = f"扫描文件夹: {INPUT_FOLDER}"
    ws["A2"].font = Font(italic=True)

    start_row = 4
    # 基础表头：视频文件、总码率、视频码率、音频码率、帧率、分辨率、视频总帧数
    headers = ["视频文件", "总码率", "视频码率", "音频码率", "帧率 (fps)", "分辨率", "视频总帧数"]
    if include_frames:
        headers.extend(["提取包数", "平均帧大小(字节)", "最小帧大小(字节)", "最大帧大小(字节)"])

    # 写表头
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 写数据
    for row_idx, (video_name, info, frame_sizes) in enumerate(all_results, start_row + 1):
        ws.cell(row=row_idx, column=1, value=video_name)
        ws.cell(row=row_idx, column=2, value=format_bps(info.get("total_bitrate_bps")))
        ws.cell(row=row_idx, column=3, value=format_bps(info.get("video_bitrate_bps")))
        ws.cell(row=row_idx, column=4, value=format_bps(info.get("audio_bitrate_bps")))

        fps_val = info.get("fps")
        ws.cell(row=row_idx, column=5, value=f"{fps_val:.2f}" if fps_val is not None else "N/A")

        # 分辨率
        w = info.get("width")
        h = info.get("height")
        if w is not None and h is not None:
            ws.cell(row=row_idx, column=6, value=f"{w}x{h}")
        else:
            ws.cell(row=row_idx, column=6, value="N/A")

        total_frames_video = info.get("video_total_frames")
        ws.cell(row=row_idx, column=7, value=total_frames_video if total_frames_video is not None else "N/A")

        if include_frames:
            if frame_sizes:
                packet_count = len(frame_sizes)
                avg_size = sum(frame_sizes) / packet_count if packet_count > 0 else 0
                min_size = min(frame_sizes) if packet_count > 0 else 0
                max_size = max(frame_sizes) if packet_count > 0 else 0
                ws.cell(row=row_idx, column=8, value=packet_count)
                ws.cell(row=row_idx, column=9, value=round(avg_size, 2))
                ws.cell(row=row_idx, column=10, value=min_size)
                ws.cell(row=row_idx, column=11, value=max_size)
            else:
                ws.cell(row=row_idx, column=8, value="N/A")
                ws.cell(row=row_idx, column=9, value="N/A")
                ws.cell(row=row_idx, column=10, value="N/A")
                ws.cell(row=row_idx, column=11, value="N/A")

    apply_wrap_and_width(ws, len(headers))
    ws.freeze_panes = f"A{start_row + 1}"


def add_frame_sheet(wb: Workbook, video_name: str, frame_sizes: List[int]):
    """为单个视频添加逐帧数据工作表"""
    sheet_name = sanitize_sheet_name(os.path.splitext(video_name)[0])
    original_sheet_name = sheet_name
    counter = 1
    while sheet_name in wb.sheetnames:
        sheet_name = f"{original_sheet_name}_{counter}"
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:31]
        counter += 1

    ws = wb.create_sheet(title=sheet_name)
    bold_font = Font(bold=True)

    ws.append(["帧索引", "帧大小 (字节)"])
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["B1"].alignment = Alignment(horizontal="center")

    for idx, size in enumerate(frame_sizes, 1):
        ws.append([idx, size])

    apply_wrap_and_width(ws, 2)
    ws.freeze_panes = "A2"


def main():
    if not os.path.isdir(INPUT_FOLDER):
        print(f"错误: 文件夹不存在 - {INPUT_FOLDER}", file=sys.stderr)
        sys.exit(1)

    video_files = []
    for fname in os.listdir(INPUT_FOLDER):
        file_path = os.path.join(INPUT_FOLDER, fname)
        if os.path.isfile(file_path):
            ext = os.path.splitext(fname)[1].lower()
            if ext in VIDEO_EXTENSIONS:
                video_files.append(file_path)

    if not video_files:
        print(f"在 {INPUT_FOLDER} 中未找到支持的视频文件。")
        return

    print(f"找到 {len(video_files)} 个视频文件，开始分析...\n")

    all_results = []
    for video_path in video_files:
        video_name = os.path.basename(video_path)
        print(f"分析: {video_name}")

        info = extract_bitrates_and_info(video_path)
        if not info:
            print(f"  无法提取媒体信息，跳过。")
            continue

        frame_sizes = None
        if INCLUDE_FRAMES:
            frame_sizes = get_frame_sizes(video_path)
            if not frame_sizes:
                print("  警告: 未能提取逐帧包数据。")

        all_results.append((video_name, info, frame_sizes))
        # 打印时也显示分辨率
        res_str = f"{info.get('width')}x{info.get('height')}" if (info.get('width') and info.get('height')) else "N/A"
        print(f"  总码率: {format_bps(info.get('total_bitrate_bps'))}, "
              f"视频码率: {format_bps(info.get('video_bitrate_bps'))}, "
              f"音频码率: {format_bps(info.get('audio_bitrate_bps'))}, "
              f"帧率: {f'{info.get('fps'):.2f}' if info.get('fps') else 'N/A'}, "
              f"分辨率: {res_str}, "
              f"总帧数: {info.get('video_total_frames')}")

    if not all_results:
        print("没有成功分析任何视频，未生成报告。")
        return

    wb = Workbook()
    create_summary_sheet(wb, all_results, INCLUDE_FRAMES)

    if INCLUDE_FRAMES:
        for video_name, _, frame_sizes in all_results:
            if frame_sizes:
                add_frame_sheet(wb, video_name, frame_sizes)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_filename = f"{OUTPUT_EXCEL_NAME}_{timestamp}.xlsx"
    output_path = os.path.join(INPUT_FOLDER, output_filename)
    wb.save(output_path)
    print(f"\n报告已生成: {output_path}")


if __name__ == "__main__":
    main()