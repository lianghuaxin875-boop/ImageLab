#!/usr/bin/env python3
"""
视频码率提取与质量检测工具（交互式选择拍摄模式）
提取总码率、视频码率、音频码率、分辨率、帧率、总帧数，并依据所选模式的标准判断码率是否达标。
报告包含所选模式的码率评判标准工作表。
依赖 FFmpeg（ffprobe）和 openpyxl。
"""

import subprocess
import json
import sys
import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl 库。请执行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# ==================== 配置区域 ====================
FFPROBE_PATH = r"D:\码率工具FFepeg\ffmpeg-8.1.1-full_build\bin\ffprobe.exe"
INPUT_FOLDER = r"D:\Codex\码率测试\input"                     # 视频文件夹
OUTPUT_EXCEL_NAME = "bitrate_report"         # 报告基础名（自动加时间戳）
INCLUDE_FRAMES = True                        # 是否包含逐帧大小数据
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".flv", ".wmv", ".webm", ".m4v", ".mpg", ".mpeg"}
# =================================================

# ---------- 四种模式的码率标准定义 ----------
# 每项: (最大像素数, 高帧率阈值fps, 普通最低码率Mbps, 高帧率最低码率Mbps)
# 如果某项没有高帧率标准，high_min 设为 None，则会统一用 normal_min

STANDARDS_CONFIG = {
    1: {   # 普通手机拍摄模式
        "name": "普通手机拍摄模式",
        "standards": [
            (921600,   30, 4.0,  None),    # 720P (1280x720)
            (2073600,  30, 8.0,  12.0),    # 1080P (1920x1080)
            (8294400,  30, 35.0, 50.0),    # 4K (3840x2160)
            (float("inf"), 30, 100.0, 150.0),  # 6K以上兜底
        ],
        "table": [
            ["720P 30fps", "1280×720", "30", "4 - 6 Mbps", "低于 4 Mbps 偏低"],
            ["1080P 30fps", "1920×1080", "30", "8 - 12 Mbps", "低于 8 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 30fps", "3840×2160", "30", "35 - 50 Mbps", "低于 35 Mbps 偏低"],
            ["4K 60fps", "3840×2160", "60", "50 - 80 Mbps", "低于 50 Mbps 偏低"],
        ]
    },
    2: {   # 短视频平台模式
        "name": "短视频平台拍摄/上传模式",
        "standards": [
            (921600,   30, 3.0,  None),    # 720P 竖屏 (720x1280)
            (2073600,  30, 6.0,  10.0),    # 1080P 竖屏 (1080x1920)
            (3686400,  30, 12.0, 16.0),    # 2K 竖屏 (1440x2560)
            (8294400,  30, 30.0, 45.0),    # 4K 竖屏 (2160x3840)
            (float("inf"), 30, 30.0, 45.0), # 兜底
        ],
        "table": [
            ["720P 竖屏", "720×1280", "30fps", "3 - 6 Mbps", "低于 3 Mbps 偏低"],
            ["1080P 竖屏", "1080×1920", "30fps", "6 - 10 Mbps", "低于 6 Mbps 偏低"],
            ["1080P 竖屏高帧率", "1080×1920", "60fps", "10 - 16 Mbps", "低于 10 Mbps 偏低"],
            ["2K 竖屏", "1440×2560", "30fps", "12 - 20 Mbps", "低于 12 Mbps 偏低"],
            ["4K 竖屏", "2160×3840", "30fps", "30 - 45 Mbps", "低于 30 Mbps 偏低"],
        ]
    },
    3: {   # 监控/低码率存储模式
        "name": "监控/低码率存储模式",
        "standards": [
            (921600,   30, 1.0,  None),    # 720P 监控 (15/25fps)
            (2073600,  30, 2.0,  None),    # 1080P 监控
            (3686400,  30, 4.0,  None),    # 2K 监控
            (8294400,  30, 8.0,  None),    # 4K 监控
            (float("inf"), 30, 8.0,  None),# 兜底
        ],
        "table": [
            ["720P 监控", "1280×720", "15/25fps", "1 - 3 Mbps", "低于 1 Mbps 偏低"],
            ["1080P 监控", "1920×1080", "15/25fps", "2 - 6 Mbps", "低于 2 Mbps 偏低"],
            ["2K 监控", "2560×1440", "15/25fps", "4 - 10 Mbps", "低于 4 Mbps 偏低"],
            ["4K 监控", "3840×2160", "15/25fps", "8 - 20 Mbps", "低于 8 Mbps 偏低"],
        ]
    },
    4: {   # 高质量拍摄/相机素材模式
        "name": "高质量拍摄/相机素材模式",
        "standards": [
            (2073600,  30, 20.0, 30.0),    # 1080P
            (8294400,  30, 60.0, 100.0),   # 4K
            (18000000, 30, 200.0, 200.0),  # 6K (约6000x3000)
            (float("inf"), 30, 200.0, 200.0),# 8K及以上
        ],
        "table": [
            ["1080P 30fps", "1920×1080", "30", "20 - 50 Mbps", "低于 20 Mbps 偏低"],
            ["1080P 60fps", "1920×1080", "60", "30 - 80 Mbps", "低于 30 Mbps 偏低"],
            ["4K 24/25/30fps", "3840×2160", "24/25/30", "60 - 100 Mbps", "低于 60 Mbps 偏低"],
            ["4K 50/60fps", "3840×2160", "50/60", "100 - 200 Mbps", "低于 100 Mbps 偏低"],
            ["4K 10bit/LOG", "3840×2160", "24/25/30", "100 - 400 Mbps", "视编码而定"],
            ["6K/8K 高质量", "6K/8K", "24/30/60", "200 Mbps 以上", "低于 200 Mbps 偏低"],
        ]
    }
}

# ---------- 码率判断逻辑 ----------
def get_bitrate_threshold_mbps(standards, pixel_count: int, fps: float) -> Optional[float]:
    if pixel_count is None or fps is None:
        return None
    for max_pixels, fps_limit, normal_min, high_min in standards:
        if pixel_count <= max_pixels:
            if high_min is not None and fps > fps_limit:
                return high_min
            else:
                return normal_min
    return None

def judge_bitrate(video_bitrate_bps: Optional[int], threshold_mbps: Optional[float]) -> str:
    if video_bitrate_bps is None or threshold_mbps is None:
        return "N/A"
    threshold_bps = threshold_mbps * 1_000_000
    return "OK" if video_bitrate_bps >= threshold_bps else "NG"

# ---------- ffprobe 相关函数 ----------
def run_ffprobe(args: List[str]) -> dict:
    cmd = [FFPROBE_PATH, "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams"] + args
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return json.loads(result.stdout)
    except subprocess.CalledProcessError as e:
        print(f"ffprobe 执行失败: {e.stderr}", file=sys.stderr)
        return {}
    except FileNotFoundError:
        print(f"错误: 未找到 ffprobe，请检查路径。当前路径：{FFPROBE_PATH}", file=sys.stderr)
        sys.exit(1)

def extract_video_info(file_path: str) -> Dict:
    info = run_ffprobe([file_path])
    if not info:
        return {}

    fmt = info.get("format", {})
    total_bitrate = fmt.get("bit_rate")
    if total_bitrate is not None:
        total_bitrate = int(total_bitrate)

    video_bitrate = None
    audio_bitrate = None
    width = None
    height = None
    fps = None
    video_total_frames = None

    for stream in info.get("streams", []):
        codec_type = stream.get("codec_type")
        if codec_type == "video":
            if width is None:
                width = stream.get("width")
                height = stream.get("height")
            if video_bitrate is None:
                if "bit_rate" in stream:
                    video_bitrate = int(stream["bit_rate"])
                else:
                    duration = float(stream.get("duration", 0))
                    size = stream.get("size")
                    if duration > 0 and size:
                        video_bitrate = int(float(size) * 8 / duration)
            if fps is None:
                r_frame_rate = stream.get("r_frame_rate")
                if r_frame_rate:
                    try:
                        if '/' in r_frame_rate:
                            num, den = r_frame_rate.split('/')
                            fps = float(num) / float(den)
                        else:
                            fps = float(r_frame_rate)
                    except (ValueError, ZeroDivisionError):
                        fps = None
            if video_total_frames is None:
                nb_frames = stream.get("nb_frames")
                if nb_frames is not None:
                    video_total_frames = int(nb_frames)
                else:
                    duration = float(stream.get("duration", 0))
                    if duration > 0 and fps is not None:
                        video_total_frames = int(duration * fps)
        elif codec_type == "audio" and audio_bitrate is None:
            if "bit_rate" in stream:
                audio_bitrate = int(stream["bit_rate"])
            else:
                duration = float(stream.get("duration", 0))
                size = stream.get("size")
                if duration > 0 and size:
                    audio_bitrate = int(float(size) * 8 / duration)

    pixel_count = width * height if width and height else None
    return {
        "total_bitrate_bps": total_bitrate,
        "video_bitrate_bps": video_bitrate,
        "audio_bitrate_bps": audio_bitrate,
        "width": width,
        "height": height,
        "pixel_count": pixel_count,
        "fps": fps,
        "video_total_frames": video_total_frames,
    }

def get_frame_sizes(file_path: str, stream_index: int = 0) -> List[int]:
    cmd = [
        FFPROBE_PATH, "-v", "quiet",
        "-select_streams", f"v:{stream_index}",
        "-show_packets", "-show_entries", "packet=size",
        "-of", "json", file_path
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, check=True)
        packets = json.loads(result.stdout).get("packets", [])
        return [int(p["size"]) for p in packets]
    except Exception as e:
        print(f"获取帧大小失败: {e}", file=sys.stderr)
        return []

def format_bps(bps: Optional[int]) -> str:
    if bps is None: return "N/A"
    if bps >= 1_000_000: return f"{bps/1_000_000:.2f} Mbps"
    elif bps >= 1_000: return f"{bps/1_000:.2f} kbps"
    return f"{bps} bps"

def sanitize_sheet_name(name: str) -> str:
    name = re.sub(r'[\\/*?:\[\]]', '_', name)
    return name[:31]

# ---------- Excel 生成 ----------
def apply_wrap_and_width(ws, headers_count: int):
    wrap_align = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=headers_count):
        for cell in row:
            cell.alignment = wrap_align
    for col_idx in range(1, headers_count + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    for line in str(cell.value).split('\n'):
                        length = sum(2 if ord(c) > 127 else 1 for c in line)
                        max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max_len + 2

def create_summary_sheet(wb: Workbook, all_results: list, mode_config: dict, include_frames: bool):
    ws = wb.active
    ws.title = "Summary"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A1:K1")
    ws["A1"] = f"视频码率分析报告（{mode_config['name']}）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:K2")
    ws["A2"] = f"扫描文件夹: {INPUT_FOLDER}"
    ws["A2"].font = Font(italic=True)

    start_row = 4
    headers = ["视频文件", "分辨率", "总码率", "视频码率", "音频码率",
               "帧率 (fps)", "视频总帧数", "码率判断"]
    if include_frames:
        headers.extend(["提取包数", "平均帧大小(字节)", "最小帧大小(字节)", "最大帧大小(字节)"])

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # 取出当前模式的标准
    standards = mode_config["standards"]

    for row_idx, (video_name, info, frame_sizes) in enumerate(all_results, start_row + 1):
        ws.cell(row=row_idx, column=1, value=video_name)
        width = info.get("width")
        height = info.get("height")
        resolution_str = f"{width}x{height}" if width and height else "N/A"
        ws.cell(row=row_idx, column=2, value=resolution_str)

        ws.cell(row=row_idx, column=3, value=format_bps(info.get("total_bitrate_bps")))
        ws.cell(row=row_idx, column=4, value=format_bps(info.get("video_bitrate_bps")))
        ws.cell(row=row_idx, column=5, value=format_bps(info.get("audio_bitrate_bps")))

        fps_val = info.get("fps")
        ws.cell(row=row_idx, column=6, value=f"{fps_val:.2f}" if fps_val else "N/A")

        total_frames = info.get("video_total_frames")
        ws.cell(row=row_idx, column=7, value=total_frames if total_frames else "N/A")

        pixel_count = info.get("pixel_count")
        threshold = get_bitrate_threshold_mbps(standards, pixel_count, fps_val) if pixel_count and fps_val else None
        judgement = judge_bitrate(info.get("video_bitrate_bps"), threshold)
        judge_cell = ws.cell(row=row_idx, column=8, value=judgement)
        judge_cell.alignment = center_align
        if judgement == "OK":
            judge_cell.fill = green_fill
        elif judgement == "NG":
            judge_cell.fill = red_fill
        else:
            judge_cell.fill = gray_fill

        if include_frames:
            base_col = 9
            if frame_sizes:
                packet_count = len(frame_sizes)
                avg_size = sum(frame_sizes) / packet_count if packet_count > 0 else 0
                min_s = min(frame_sizes) if packet_count > 0 else 0
                max_s = max(frame_sizes) if packet_count > 0 else 0
                ws.cell(row=row_idx, column=base_col, value=packet_count)
                ws.cell(row=row_idx, column=base_col+1, value=round(avg_size, 2))
                ws.cell(row=row_idx, column=base_col+2, value=min_s)
                ws.cell(row=row_idx, column=base_col+3, value=max_s)
            else:
                for i in range(4):
                    ws.cell(row=row_idx, column=base_col+i, value="N/A")

    apply_wrap_and_width(ws, len(headers))
    ws.freeze_panes = f"A{start_row + 1}"

def add_standard_sheet(wb: Workbook, mode_config: dict):
    ws = wb.create_sheet("码率评判标准", 1)
    mode_name = mode_config["name"]
    ws.merge_cells("A1:E1")
    ws["A1"] = f"评判标准 - {mode_name}"
    ws["A1"].font = Font(bold=True, size=13)

    headers = ["拍摄模式/分辨率", "分辨率", "帧率", "建议码率", "检测判断"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for row_idx, row_data in enumerate(mode_config["table"], 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    length = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                    max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = max_len + 3

    ws.freeze_panes = "A4"

def add_frame_sheet(wb, video_name, frame_sizes):
    sheet_name = sanitize_sheet_name(os.path.splitext(video_name)[0])
    original = sheet_name
    counter = 1
    while sheet_name in wb.sheetnames:
        sheet_name = f"{original}_{counter}"[:31]
        counter += 1
    ws = wb.create_sheet(title=sheet_name)
    bold_font = Font(bold=True)
    ws.append(["帧索引", "帧大小 (字节)"])
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["B1"].alignment = Alignment(horizontal="center")
    for idx, size in enumerate(frame_sizes, 1):
        ws.append([idx, size])
    apply_wrap_and_width(ws, 2)
    ws.freeze_panes = "A2"

def select_mode() -> int:
    """交互式选择拍摄模式，返回1-4的有效模式编号"""
    print("请选择拍摄模式：")
    print("1 - 普通手机拍摄模式")
    print("2 - 短视频平台模式")
    print("3 - 监控/低码率存储模式")
    print("4 - 高质量拍摄/相机素材模式")
    while True:
        try:
            choice = input("请输入数字 (1-4): ").strip()
            mode = int(choice)
            if mode in STANDARDS_CONFIG:
                return mode
            else:
                print("输入无效，请输入 1、2、3 或 4。")
        except ValueError:
            print("输入无效，请输入数字。")

def main():
    # 第一步：选择模式
    selected_mode = select_mode()
    mode_config = STANDARDS_CONFIG[selected_mode]
    print(f"当前模式：{mode_config['name']}\n")

    if not os.path.isdir(INPUT_FOLDER):
        print(f"错误: 文件夹不存在 - {INPUT_FOLDER}")
        sys.exit(1)

    video_files = []
    for fname in os.listdir(INPUT_FOLDER):
        fp = os.path.join(INPUT_FOLDER, fname)
        if os.path.isfile(fp) and os.path.splitext(fname)[1].lower() in VIDEO_EXTENSIONS:
            video_files.append(fp)

    if not video_files:
        print(f"在 {INPUT_FOLDER} 中未找到支持的视频文件。")
        return

    print(f"找到 {len(video_files)} 个视频文件，开始分析...\n")
    all_results = []

    for vpath in video_files:
        vname = os.path.basename(vpath)
        print(f"分析: {vname}")
        info = extract_video_info(vpath)
        if not info:
            print("  无法提取媒体信息，跳过。")
            continue
        frame_sizes = None
        if INCLUDE_FRAMES:
            frame_sizes = get_frame_sizes(vpath)
            if not frame_sizes:
                print("  警告: 未能提取逐帧数据。")
        all_results.append((vname, info, frame_sizes))
        pixel = info.get('pixel_count')
        fps = info.get('fps')
        threshold = get_bitrate_threshold_mbps(mode_config["standards"], pixel, fps) if pixel and fps else None
        judge = judge_bitrate(info.get('video_bitrate_bps'), threshold)
        print(f"  分辨率: {info.get('width')}x{info.get('height')}, "
              f"帧率: {fps}, "
              f"视频码率: {format_bps(info.get('video_bitrate_bps'))}, "
              f"判断: {judge} (阈值: {threshold} Mbps)")

    if not all_results:
        print("没有成功分析任何视频。")
        return

    wb = Workbook()
    create_summary_sheet(wb, all_results, mode_config, INCLUDE_FRAMES)
    add_standard_sheet(wb, mode_config)
    if INCLUDE_FRAMES:
        for vname, _, frames in all_results:
            if frames:
                add_frame_sheet(wb, vname, frames)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"{OUTPUT_EXCEL_NAME}_{timestamp}.xlsx"
    out_path = os.path.join(INPUT_FOLDER, out_name)
    wb.save(out_path)
    print(f"\n报告已生成: {out_path}")

if __name__ == "__main__":
    main()