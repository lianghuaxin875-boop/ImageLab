# -*- coding: utf-8 -*-
"""Excel and CSV report generation for rule-based image enhancement."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPORT_COLUMNS = [
    "检测时间",
    "文件名",
    "原图路径",
    "增强图路径",
    "对比图路径",
    "原图综合分",
    "增强后综合分",
    "分数变化",
    "增强效果判断",
    "原图质量等级",
    "增强后质量等级",
    "原图问题原因",
    "执行增强方法",
    "增强后问题原因",
    "处理建议",
    "原图清晰度分",
    "增强后清晰度分",
    "原图曝光分",
    "增强后曝光分",
    "原图对比度分",
    "增强后对比度分",
    "原图噪声分",
    "增强后噪声分",
    "原图压缩伪影分",
    "增强后压缩伪影分",
]


def write_reports(rows: list[dict[str, Any]], report_dir: str | Path) -> tuple[Path, Path]:
    report_dir = Path(report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = report_dir / f"图像增强前后对比报告_{timestamp}.xlsx"
    csv_path = report_dir / f"图像增强前后对比报告_{timestamp}.csv"

    write_excel(rows, xlsx_path)
    write_csv(rows, csv_path)
    return xlsx_path, csv_path


def write_csv(rows: list[dict[str, Any]], csv_path: str | Path) -> None:
    csv_path = Path(csv_path)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_excel(rows: list[dict[str, Any]], xlsx_path: str | Path) -> None:
    xlsx_path = Path(xlsx_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "增强对比"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    ng_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin = Side("thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(REPORT_COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        ws.append([row.get(column, "") for column in REPORT_COLUMNS])

    effect_col = REPORT_COLUMNS.index("增强效果判断") + 1
    for row_idx in range(2, ws.max_row + 1):
        effect = str(ws.cell(row_idx, effect_col).value or "")
        fill = ok_fill if "有效" in effect and "无效" not in effect else warn_fill if "不明显" in effect else ng_fill if "过度" in effect or "无效" in effect else None
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    ws.freeze_panes = "A2"
    widths = {
        "A": 20,
        "B": 24,
        "C": 34,
        "D": 34,
        "E": 34,
        "L": 28,
        "M": 34,
        "N": 28,
        "O": 36,
    }
    for idx, column in enumerate(REPORT_COLUMNS, 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(letter, 16 if "分" in column or "等级" in column else 20)

    summary = wb.create_sheet("汇总")
    total = len(rows)
    effective = sum(1 for row in rows if row.get("增强效果判断") == "增强有效")
    unclear = sum(1 for row in rows if row.get("增强效果判断") == "效果不明显")
    worse = sum(1 for row in rows if "过度" in str(row.get("增强效果判断", "")) or "无效" in str(row.get("增强效果判断", "")))
    avg_delta = round(sum(float(row.get("分数变化", 0) or 0) for row in rows) / total, 2) if total else 0
    summary_rows = [
        ("报告生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("处理图片数量", total),
        ("增强有效数量", effective),
        ("效果不明显数量", unclear),
        ("可能过度或无效数量", worse),
        ("平均分数变化", avg_delta),
    ]
    for item in summary_rows:
        summary.append(item)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 28

    wb.save(xlsx_path)

